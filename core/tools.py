"""Headless document tools shared by the modern UI and tests."""

from __future__ import annotations

import csv
import html
import os
import re
import secrets
import tempfile
import xml.etree.ElementTree as ET
from collections.abc import Callable, Iterable
from pathlib import Path

import fitz

from .capabilities import CapabilityId, detect_capabilities
from .platform_service import PlatformService

ProgressCallback = Callable[[int, int, str], None]


class ToolError(RuntimeError):
    pass


def _progress(callback: ProgressCallback | None, current: int, total: int, message: str) -> None:
    if callback:
        callback(current, total, message)


def _temporary_pdf_path(target: Path) -> Path:
    """Return a non-existent temporary PDF beside *target* for atomic replacement."""
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".pdf", dir=target.parent
    )
    os.close(descriptor)
    os.unlink(name)
    return Path(name)


def _pdf_files(folder: Path, recursive: bool) -> list[Path]:
    """Find PDFs without relying on case-sensitive ``*.pdf`` globs."""
    iterator = folder.rglob("*") if recursive else folder.iterdir()
    return sorted(
        path for path in iterator if path.is_file() and path.suffix.casefold() == ".pdf"
    )


def _check_distinct_targets(jobs: Iterable[tuple[Path, Path]]) -> None:
    seen: dict[str, Path] = {}
    for source, target in jobs:
        key = str(target).casefold() if PlatformService.WINDOWS or PlatformService.MACOS else str(target)
        previous = seen.get(key)
        if previous is not None and previous != source:
            raise ToolError(
                f"Multiple source files would create the same output: {target.name}"
            )
        seen[key] = source


def _check_no_inplace(jobs: Iterable[tuple[Path, Path]]) -> None:
    """Reject jobs whose output would overwrite the source in place.

    An empty suffix with overwrite disabled silently produced a target
    equal to the source, rewriting the original file without consent.
    """
    for source, target in jobs:
        if target == source:
            raise ToolError(
                "The output would overwrite the source file. Add a filename "
                "suffix or enable overwrite."
            )


def merge_pdfs(
    paths: Iterable[str | os.PathLike[str]],
    output_path: str | os.PathLike[str],
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> Path:
    sources = [Path(path).expanduser().resolve() for path in paths]
    if not sources:
        raise ToolError("Select at least one PDF to merge.")
    target = Path(output_path).expanduser().resolve()
    missing = next((source for source in sources if not source.is_file()), None)
    if missing:
        raise ToolError(f"Source file not found: {missing}")
    temporary = _temporary_pdf_path(target)
    try:
        with fitz.open() as output:
            total = len(sources)
            first_metadata: dict | None = None
            for index, source_path in enumerate(sources):
                if is_cancelled and is_cancelled():
                    break
                _progress(progress, index, total, source_path.name)
                try:
                    with fitz.open(source_path) as source:
                        if first_metadata is None:
                            first_metadata = source.metadata or {}
                        output.insert_pdf(source)
                except Exception as exc:
                    raise ToolError(f"Cannot read {source_path.name}: {exc}") from exc
                _progress(progress, index + 1, total, source_path.name)
            if is_cancelled and is_cancelled():
                raise ToolError("The merge was cancelled.")
            if first_metadata:
                output.set_metadata(first_metadata)
            output.save(temporary, garbage=4, deflate=True)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def overlay_pdf(
    template_path: str | os.PathLike[str],
    target_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
) -> Path:
    target_output = Path(output_path).expanduser().resolve()
    temporary = _temporary_pdf_path(target_output)
    try:
        try:
            template = fitz.open(template_path)
        except Exception as exc:
            raise ToolError(f"Cannot read the overlay PDF: {exc}") from exc
        with template:
            try:
                target = fitz.open(target_path)
            except Exception as exc:
                raise ToolError(f"Cannot read the target PDF: {exc}") from exc
            with target:
                if template.page_count == 0:
                    raise ToolError("The overlay template has no pages.")
                if target.page_count == 0:
                    raise ToolError("The target PDF has no pages.")
                for page_number in range(target.page_count):
                    page = target.load_page(page_number)
                    template_page = min(page_number, template.page_count - 1)
                    page.show_pdf_page(
                        page.rect,
                        template,
                        template_page,
                        overlay=True,
                        keep_proportion=True,
                    )
                target.save(temporary, garbage=4, deflate=True)
        os.replace(temporary, target_output)
    finally:
        temporary.unlink(missing_ok=True)
    return target_output


def compress_pdf(
    source_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
    *,
    garbage: int = 4,
    clean: bool = True,
    deflate: bool = True,
    deflate_images: bool = True,
    deflate_fonts: bool = True,
    linear: bool = False,
) -> Path:
    """Rewrite one PDF using the selected PyMuPDF optimization options."""
    source = Path(source_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    temporary = _temporary_pdf_path(target)
    try:
        try:
            with fitz.open(source) as doc:
                doc.save(
                    temporary,
                    garbage=max(0, min(4, int(garbage))),
                    clean=clean,
                    deflate=deflate,
                    deflate_images=deflate_images,
                    deflate_fonts=deflate_fonts,
                    linear=linear,
                    encryption=fitz.PDF_ENCRYPT_KEEP,
                )
        except ToolError:
            raise
        except Exception as exc:
            raise ToolError(f"Cannot read {source.name}: {exc}") from exc
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def compress_pdfs(
    paths: Iterable[str | os.PathLike[str]],
    output_folder: str | os.PathLike[str] | None = None,
    suffix: str = "_compressed",
    overwrite: bool = False,
    garbage: int = 4,
    clean: bool = True,
    deflate: bool = True,
    deflate_images: bool = True,
    deflate_fonts: bool = True,
    linear: bool = False,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> list[Path]:
    sources = [Path(value).expanduser().resolve() for value in paths]
    destination = Path(output_folder).expanduser().resolve() if output_folder else None
    jobs: list[tuple[Path, Path]] = []
    for source in sources:
        target_folder = destination or source.parent
        target = source if overwrite and not suffix else target_folder / f"{source.stem}{suffix}.pdf"
        jobs.append((source, target))
    _check_distinct_targets(jobs)
    if not overwrite:
        _check_no_inplace(jobs)
    existing = next(
        (target for source, target in jobs if target.exists() and target != source),
        None,
    )
    if existing and not overwrite:
        raise ToolError(f"Output already exists: {existing.name}")

    outputs: list[Path] = []
    for index, (source, target) in enumerate(jobs, 1):
        if is_cancelled and is_cancelled():
            break
        _progress(progress, index - 1, len(sources), source.name)
        outputs.append(
            compress_pdf(
                source,
                target,
                garbage=garbage,
                clean=clean,
                deflate=deflate,
                deflate_images=deflate_images,
                deflate_fonts=deflate_fonts,
                linear=linear,
            )
        )
        _progress(progress, index, len(sources), source.name)
    return outputs


def overlay_pdfs(
    template_path: str | os.PathLike[str],
    targets: Iterable[str | os.PathLike[str]],
    output_folder: str | os.PathLike[str],
    suffix: str = "_overlay",
    overwrite: bool = False,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> list[Path]:
    template = Path(template_path).expanduser().resolve()
    folder = Path(output_folder).expanduser().resolve()
    sources = [Path(value).expanduser().resolve() for value in targets]
    folder.mkdir(parents=True, exist_ok=True)
    jobs = [
        (
            source,
            source if overwrite and not suffix else folder / f"{source.stem}{suffix}.pdf",
        )
        for source in sources
    ]
    _check_distinct_targets(jobs)
    if not overwrite:
        _check_no_inplace(jobs)
    existing = next(
        (target for source, target in jobs if target.exists() and target != source),
        None,
    )
    if existing and not overwrite:
        raise ToolError(f"Output already exists: {existing.name}")

    outputs: list[Path] = []
    for index, (source, target) in enumerate(jobs, 1):
        if is_cancelled and is_cancelled():
            break
        _progress(progress, index - 1, len(sources), source.name)
        overlay_pdf(template, source, target)
        outputs.append(target)
        _progress(progress, index, len(sources), source.name)
    return outputs


def convert_pdf_to_word(source_path: str | os.PathLike[str], output_path: str | os.PathLike[str]) -> Path:
    capability = detect_capabilities()[CapabilityId.PDF_TO_WORD]
    if not capability.available:
        raise ToolError(capability.reason)
    from pdf2docx import Converter

    source = Path(source_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".docx", dir=target.parent
    )
    os.close(handle)
    try:
        try:
            converter = Converter(str(source))
        except Exception as exc:
            raise ToolError(f"Cannot read {source.name}: {exc}") from exc
        try:
            # Write to a temp file first so a failed conversion never leaves
            # a partial .docx behind.
            converter.convert(temp_name, start=0, end=None)
        finally:
            converter.close()
        os.replace(temp_name, target)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    return target


def convert_office_files(
    paths: Iterable[str | os.PathLike[str]],
    output_dir: str | os.PathLike[str],
    overwrite: bool = False,
    source_root: str | os.PathLike[str] | None = None,
    keep_structure: bool = False,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> list[Path]:
    sources = [Path(path).expanduser().resolve() for path in paths]
    folder = Path(output_dir).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    capability = detect_capabilities()[CapabilityId.OFFICE_TO_PDF]
    if not capability.available:
        raise ToolError(capability.reason)

    root = Path(source_root).expanduser().resolve() if source_root else None
    supported = {".xlsx", ".xls", ".docx", ".doc", ".pptx", ".ppt"}
    unsupported = sorted(
        {source.suffix.lower() for source in sources if source.suffix.lower() not in supported}
    )
    if unsupported:
        raise ToolError(
            "Unsupported file type(s) for Office conversion: " + ", ".join(unsupported)
        )
    outputs: list[Path] = []

    def destination_for(source: Path) -> tuple[Path, Path]:
        destination = folder
        if keep_structure and root:
            try:
                destination = folder / source.parent.relative_to(root)
            except ValueError:
                destination = folder
        destination.mkdir(parents=True, exist_ok=True)
        return destination, destination / f"{source.stem}.pdf"

    jobs = [(source, *destination_for(source)) for source in sources]
    _check_distinct_targets((source, target) for source, _folder, target in jobs)

    if capability.backend == "LibreOffice":
        for index, (source, destination, target) in enumerate(jobs, 1):
            if is_cancelled and is_cancelled():
                break
            _progress(progress, index - 1, len(sources), source.name)
            if target.exists() and not overwrite:
                _progress(progress, index, len(sources), f"Skipped {source.name}")
                continue
            result = PlatformService.run(
                [
                    capability.path,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(destination),
                    str(source),
                ],
                timeout=300,
            )
            if result.returncode != 0 or not target.exists():
                raise ToolError(result.stderr.strip() or f"Failed to convert {source.name}.")
            outputs.append(target)
            _progress(progress, index, len(sources), source.name)
        return outputs

    # Import Windows COM only inside the Windows-only backend.
    import comtypes
    import comtypes.client

    comtypes.CoInitialize()
    try:
        for index, (source, _destination, target) in enumerate(jobs, 1):
            if is_cancelled and is_cancelled():
                break
            _progress(progress, index - 1, len(sources), source.name)
            suffix = source.suffix.lower()
            app = document = None
            if target.exists() and not overwrite:
                _progress(progress, index, len(sources), f"Skipped {source.name}")
                continue
            try:
                if suffix in {".xlsx", ".xls"}:
                    app = comtypes.client.CreateObject("Excel.Application")
                    app.Visible = False
                    app.DisplayAlerts = False
                    document = app.Workbooks.Open(str(source))
                    document.ExportAsFixedFormat(0, str(target))
                elif suffix in {".docx", ".doc"}:
                    app = comtypes.client.CreateObject("Word.Application")
                    app.Visible = False
                    app.DisplayAlerts = 0
                    document = app.Documents.Open(str(source))
                    document.SaveAs(str(target), FileFormat=17)
                elif suffix in {".pptx", ".ppt"}:
                    app = comtypes.client.CreateObject("PowerPoint.Application")
                    document = app.Presentations.Open(str(source), WithWindow=False)
                    document.SaveAs(str(target), FileFormat=32)
                else:
                    continue
                outputs.append(target)
            except Exception as exc:
                raise ToolError(f"Failed to convert {source.name}: {exc}") from exc
            finally:
                if document is not None:
                    try:
                        document.Close(SaveChanges=False)
                    except Exception:
                        pass
                    document = None
                if app is not None:
                    try:
                        app.Quit()
                    except Exception:
                        pass
                    try:
                        app.Release()
                    except Exception:
                        pass
                    app = None
            _progress(progress, index, len(sources), source.name)
    finally:
        comtypes.CoUninitialize()
    return outputs


def encrypt_pdf_file(
    source_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
    user_password: str,
    *,
    owner_password: str = "",
    encryption: int | None = None,
    permissions: int | None = None,
) -> Path:
    """Encrypt a PDF file (safe to run in a background task)."""
    source = Path(source_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".pdf", dir=target.parent
    )
    os.close(handle)
    try:
        with fitz.open(source) as doc:
            doc.save(
                temp_name,
                garbage=4,
                deflate=True,
                encryption=encryption or fitz.PDF_ENCRYPT_AES_256,
                owner_pw=owner_password or secrets.token_hex(20),
                user_pw=user_password,
                permissions=int(
                    permissions
                    if permissions is not None
                    else fitz.PDF_PERM_ACCESSIBILITY | fitz.PDF_PERM_PRINT
                ),
            )
        os.replace(temp_name, target)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    return target


def decrypt_pdf_file(
    source_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
    password: str | None = None,
) -> Path:
    """Remove encryption from a PDF file (safe to run in a background task)."""
    source = Path(source_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".pdf", dir=target.parent
    )
    os.close(handle)
    try:
        with fitz.open(source) as doc:
            if doc.needs_pass:
                if not password or not doc.authenticate(password):
                    raise ToolError("The password is not valid.")
            try:
                doc.save(
                    temp_name,
                    garbage=4,
                    deflate=True,
                    encryption=fitz.PDF_ENCRYPT_NONE,
                )
            except Exception:
                Path(temp_name).unlink(missing_ok=True)
                with fitz.open() as output:
                    if doc.page_count:
                        output.insert_pdf(
                            doc, from_page=0, to_page=doc.page_count - 1
                        )
                    output.set_metadata(doc.metadata or {})
                    output.set_toc(doc.get_toc() or [])
                    output.save(
                        temp_name,
                        garbage=4,
                        deflate=True,
                        encryption=fitz.PDF_ENCRYPT_NONE,
                    )
        os.replace(temp_name, target)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    return target


def convert_postscript(
    source_path: str | os.PathLike[str], output_path: str | os.PathLike[str]
) -> Path:
    capability = detect_capabilities()[CapabilityId.POSTSCRIPT]
    if not capability.available:
        raise ToolError(capability.reason)
    source = Path(source_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    temporary = _temporary_pdf_path(target)
    try:
        result = PlatformService.run(
            [
                capability.path,
                "-dNOPAUSE",
                "-dBATCH",
                "-dSAFER",
                "-sDEVICE=pdfwrite",
                "-dCompatibilityLevel=1.7",
                f"-sOutputFile={temporary}",
                str(source),
            ],
            timeout=300,
        )
        if result.returncode != 0 or not temporary.exists():
            raise ToolError(
                result.stderr.strip() or "Ghostscript did not create an output file."
            )
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def convert_postscript_files(
    paths: Iterable[str | os.PathLike[str]],
    output_folder: str | os.PathLike[str],
    prefix: str = "",
    overwrite: bool = False,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> list[Path]:
    sources = [Path(value).expanduser().resolve() for value in paths]
    folder = Path(output_folder).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    jobs = [(source, folder / f"{prefix}{source.stem}.pdf") for source in sources]
    _check_distinct_targets(jobs)
    existing = next((target for _source, target in jobs if target.exists()), None)
    if existing and not overwrite:
        raise ToolError(f"Output already exists: {existing.name}")
    outputs: list[Path] = []
    for index, (source, target) in enumerate(jobs, 1):
        if is_cancelled and is_cancelled():
            break
        _progress(progress, index - 1, len(sources), source.name)
        outputs.append(convert_postscript(source, target))
        _progress(progress, index, len(sources), source.name)
    return outputs


def _read_text_file(source: Path, encoding_choice: int = 0) -> str:
    raw = source.read_bytes()
    choices = {
        1: ("utf-8-sig", "utf-8"),
        2: ("big5",),
        3: ("gb18030",),
    }
    encodings = choices.get(encoding_choice, ("utf-8-sig", "utf-8", "big5", "gb18030", "latin-1"))
    for encoding in encodings:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode(encodings[0], errors="replace")


def text_files_to_pdf(
    paths: Iterable[str | os.PathLike[str]],
    output_path: str | os.PathLike[str],
    encoding_choice: int = 0,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> Path:
    sources = [Path(path).expanduser().resolve() for path in paths]
    target = Path(output_path).expanduser().resolve()
    if not sources:
        raise ToolError("Select at least one text file.")
    temporary = _temporary_pdf_path(target)
    writer = fitz.DocumentWriter(temporary)
    wrote_page = False
    try:
        media_box = fitz.paper_rect("a4")
        text_box = fitz.Rect(54, 54, 541, 788)
        css = (
            ".document { font-size: 10pt; line-height: 1.45; "
            "white-space: pre-wrap; overflow-wrap: anywhere; }"
        )
        for index, source in enumerate(sources, 1):
            if is_cancelled and is_cancelled():
                break
            _progress(progress, index - 1, len(sources), source.name)
            content = _read_text_file(source, encoding_choice)
            safe = html.escape(content) or "&#8203;"
            story = fitz.Story(f'<div class="document">{safe}</div>', user_css=css)
            more = True
            while more:
                if is_cancelled and is_cancelled():
                    break
                device = writer.begin_page(media_box)
                more, _filled = story.place(text_box)
                story.draw(device)
                writer.end_page()
                wrote_page = True
            _progress(progress, index, len(sources), source.name)
        writer.close()
        # PyMuPDF keeps the output file open until the writer object is
        # garbage-collected. Drop the reference so the handle is released
        # before the rename below (required on Windows).
        writer = None
        if not wrote_page:
            raise ToolError("Text conversion was cancelled before any pages were created.")
        if is_cancelled and is_cancelled():
            raise ToolError("The text conversion was cancelled.")
        os.replace(temporary, target)
    except Exception:
        try:
            if writer is not None:
                writer.close()
        except Exception:
            pass
        raise
    finally:
        temporary.unlink(missing_ok=True)
    return target


def text_files_to_pdfs(
    paths: Iterable[str | os.PathLike[str]],
    output_folder: str | os.PathLike[str],
    filename: str = "text-files.pdf",
    single: bool = True,
    encoding_choice: int = 0,
    overwrite: bool = False,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> list[Path]:
    sources = [Path(value).expanduser().resolve() for value in paths]
    folder = Path(output_folder).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    if single:
        target = folder / filename
        if target.exists() and not overwrite:
            raise ToolError(f"Output already exists: {target.name}")
        _progress(progress, 0, len(sources), "Preparing text files")
        output = text_files_to_pdf(
            sources,
            target,
            encoding_choice,
            progress=progress,
            is_cancelled=is_cancelled,
        )
        _progress(progress, len(sources), len(sources), output.name)
        return [output]
    jobs = [(source, folder / f"{source.stem}.pdf") for source in sources]
    _check_distinct_targets(jobs)
    existing = next((target for _source, target in jobs if target.exists()), None)
    if existing and not overwrite:
        raise ToolError(f"Output already exists: {existing.name}")
    outputs: list[Path] = []
    for index, (source, target) in enumerate(jobs, 1):
        if is_cancelled and is_cancelled():
            break
        _progress(progress, index - 1, len(sources), source.name)
        outputs.append(text_files_to_pdf([source], target, encoding_choice))
        _progress(progress, index, len(sources), source.name)
    return outputs


def create_page_count_report(
    folder_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str],
    recursive: bool = True,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> Path:
    capability = detect_capabilities()[CapabilityId.SPREADSHEET]
    if not capability.available:
        raise ToolError(capability.reason)
    from openpyxl import Workbook
    from openpyxl.styles import Font

    folder = Path(folder_path).expanduser().resolve()
    rows = []
    paths = _pdf_files(folder, recursive)
    for index, path in enumerate(paths, 1):
        if is_cancelled and is_cancelled():
            break
        _progress(progress, index - 1, len(paths), path.name)
        try:
            with fitz.open(path) as doc:
                first_size = doc.load_page(0).rect if doc.page_count else None
                paper = (
                    f"{first_size.width * 25.4 / 72:.1f} × "
                    f"{first_size.height * 25.4 / 72:.1f} mm"
                    if first_size
                    else "—"
                )
                rows.append((path.name, doc.page_count, path.stat().st_size, paper, str(path)))
        except Exception:
            rows.append((path.name, "Error", path.stat().st_size, "—", str(path)))
        _progress(progress, index, len(paths), path.name)
    target = Path(output_path).expanduser().resolve()
    if is_cancelled and is_cancelled():
        raise ToolError("The page count report was cancelled.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PDF Report"
    sheet.append(["Filename", "Pages", "Size (bytes)", "First page size", "Path"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 24
    sheet.column_dimensions["E"].width = 72
    workbook.save(target)
    return target


def deep_search(
    folder_path: str | os.PathLike[str],
    query: str,
    include_subfolders: bool = True,
    search_barcodes: bool = False,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> list[dict[str, object]]:
    folder = Path(folder_path).expanduser().resolve()
    if not folder.is_dir():
        raise ToolError("Choose a valid folder to search.")
    keywords = [keyword.strip().casefold() for keyword in query.split(",") if keyword.strip()]
    if not keywords:
        raise ToolError("Enter text to search for.")
    results: list[dict[str, object]] = []
    files = _pdf_files(folder, include_subfolders)
    for index, path in enumerate(files):
        if is_cancelled and is_cancelled():
            break
        _progress(progress, index, len(files), path.name)
        try:
            with fitz.open(path) as doc:
                pages: list[int] = []
                snippets: list[str] = []
                for page_index, page in enumerate(doc):
                    if is_cancelled and is_cancelled():
                        break
                    snippet = _page_search_match(page, keywords, search_barcodes)
                    if snippet:
                        pages.append(page_index + 1)
                        snippets.append(snippet)
                if pages:
                    results.append(
                        {"path": str(path), "filename": path.name, "pages": pages, "snippets": snippets}
                    )
        except Exception as exc:
            results.append(
                {
                    "path": str(path),
                    "filename": path.name,
                    "pages": [],
                    "snippets": [],
                    "error": str(exc) or exc.__class__.__name__,
                }
            )
        _progress(progress, index + 1, len(files), path.name)
    return results


def _page_search_match(page, keywords: list[str], search_barcodes: bool) -> str:
    """Return a match snippet for the first keyword found on the page.

    Keyword matching covers the extracted text and, when enabled, the decoded
    barcode/QR content of the rendered page.
    """
    text = page.get_text()
    for keyword in keywords:
        # casefold can change string length (e.g. ß -> ss), which would shift
        # the snippet window; regex spans index the ORIGINAL text instead.
        match = re.search(re.escape(keyword), text, re.IGNORECASE)
        if match:
            return " ".join(text[max(0, match.start() - 60): match.end() + 90].split())
    if not search_barcodes:
        return ""
    data = _page_barcode_data(page)
    if not data:
        return ""
    for keyword in keywords:
        match = re.search(re.escape(keyword), data, re.IGNORECASE)
        if match:
            return "[barcode] " + " ".join(
                data[max(0, match.start() - 60): match.end() + 90].split()
            )
    return ""


def _page_barcode_data(page) -> str:
    """Decode barcodes/QR codes rendered on a page, joined into one string."""
    try:
        from PIL import Image
        from pyzbar.pyzbar import decode
    except ImportError:
        return ""
    try:
        pixmap = page.get_pixmap(dpi=150, alpha=False)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        parts: list[str] = []
        for value in decode(image):
            try:
                parts.append(value.data.decode("utf-8", errors="replace"))
            except Exception:
                continue
        return " ".join(parts)
    except Exception:
        return ""


def export_search_results_csv(results: list[dict[str, object]], output_path: str | os.PathLike[str]) -> Path:
    target = Path(output_path).expanduser().resolve()
    with target.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Filename", "Pages", "Context", "Path"])
        for result in results:
            writer.writerow([
                result["filename"],
                ", ".join(map(str, result["pages"])),
                str(result.get("error") or " | ".join(map(str, result["snippets"]))),
                result["path"],
            ])
    return target


def extract_region_text(
    source_path: str | os.PathLike[str],
    pages: Iterable[int],
    rect: tuple[float, float, float, float],
    output_path: str | os.PathLike[str],
    excel: bool = True,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> Path:
    selected = list(pages)
    target = Path(output_path).expanduser().resolve()
    extracted: list[tuple[int, list[str]]] = []
    with fitz.open(source_path) as document:
        for index, page_number in enumerate(selected, 1):
            if is_cancelled and is_cancelled():
                break
            _progress(progress, index - 1, len(selected), f"Page {page_number + 1}")
            text = document.load_page(page_number).get_text(
                "text", clip=fitz.Rect(rect), sort=True
            ).strip()
            extracted.append(
                (page_number + 1, [line.strip() for line in text.splitlines() if line.strip()])
            )
            _progress(progress, index, len(selected), f"Page {page_number + 1}")
    if is_cancelled and is_cancelled():
        return target
    if excel:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise ToolError("Excel export requires openpyxl.") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Extracted Text"
        maximum = max((len(lines) for _page, lines in extracted), default=0)
        sheet.append(["Page"] + [f"Field {index}" for index in range(1, maximum + 1)])
        for page_number, lines in extracted:
            sheet.append([page_number, *lines])
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, sheet.max_row + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(
                80, max(10, max(map(len, values), default=8) + 2)
            )
        workbook.save(target)
    else:
        chunks = [
            f"--- Page {page_number} ---\n" + "\n".join(lines)
            for page_number, lines in extracted
        ]
        target.write_text("\n\n".join(chunks), encoding="utf-8")
    return target


def merge_spreadsheets(
    paths: Iterable[str | os.PathLike[str]],
    output_path: str | os.PathLike[str],
    skip_rows: int = 0,
    exclude_keywords: Iterable[str] | None = None,
    first_cell_only: bool = True,
    encoding: str = "auto",
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
) -> Path:
    capability = detect_capabilities()[CapabilityId.SPREADSHEET]
    if not capability.available:
        raise ToolError(capability.reason)
    from openpyxl import Workbook, load_workbook

    def csv_value(value: str) -> object:
        text = value.strip()
        if not text:
            return None
        # Match pandas' useful numeric inference from the legacy tool while
        # preserving identifiers such as 00123 as text.
        if re.fullmatch(r"[+-]?(?:0|[1-9][0-9]*)", text):
            unsigned = text.lstrip("+-")
            if len(unsigned) == 1 or not unsigned.startswith("0"):
                try:
                    return int(text)
                except ValueError:
                    pass
        if re.fullmatch(
            r"[+-]?(?:[0-9]+[.][0-9]*|[.][0-9]+|"
            r"[0-9]+[eE][+-]?[0-9]+|"
            r"[0-9]+[.][0-9]*[eE][+-]?[0-9]+)",
            text,
        ):
            try:
                return float(text)
            except ValueError:
                pass
        return value

    def blank_row(row: Iterable[object]) -> bool:
        return all(value is None or str(value).strip() == "" for value in row)

    def read_xls(path: Path) -> list[list[object]]:
        try:
            import xlrd
        except ImportError as exc:
            raise ToolError(
                "Legacy .xls support requires the xlrd package."
            ) from exc
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheet = workbook.sheet_by_index(0)
            rows: list[list[object]] = []
            for row_index in range(sheet.nrows):
                values: list[object] = []
                for column_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, column_index)
                    value: object = cell.value
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        number = float(cell.value)
                        value = int(number) if number.is_integer() else number
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate.xldate_as_datetime(
                            cell.value, workbook.datemode
                        )
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = None
                    values.append(value)
                rows.append(values)
            return rows
        finally:
            workbook.release_resources()

    def unique_headers(values: Iterable[object]) -> list[str]:
        result: list[str] = []
        counts: dict[str, int] = {}
        for column, value in enumerate(values, 1):
            base = str(value).strip() if value is not None else ""
            base = base or f"Column {column}"
            counts[base] = counts.get(base, 0) + 1
            result.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
        return result

    tables: list[tuple[list[object], list[list[object]]]] = []
    source_values = list(paths)
    keywords = [value.casefold() for value in (exclude_keywords or []) if value]
    for file_index, value in enumerate(source_values, 1):
        if is_cancelled and is_cancelled():
            break
        path = Path(value).expanduser().resolve()
        _progress(progress, file_index - 1, len(source_values), path.name)
        if path.suffix.lower() == ".csv":
            raw = path.read_bytes()
            decoded = None
            encodings = (
                ("utf-8-sig", "utf-8", "big5", "gb18030", "latin-1")
                if encoding == "auto"
                else (encoding,)
            )
            for candidate in encodings:
                try:
                    decoded = raw.decode(candidate)
                    break
                except UnicodeDecodeError:
                    continue
            rows = [
                [csv_value(cell) for cell in row]
                for row in csv.reader((decoded or "").splitlines())
            ]
        elif path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
            finally:
                workbook.close()
        elif path.suffix.lower() == ".xls":
            rows = read_xls(path)
        else:
            raise ToolError(f"Unsupported spreadsheet format: {path.suffix or path.name}")
        rows = rows[max(0, int(skip_rows)) :]
        rows = [list(row) for row in rows if not blank_row(row)]
        if rows:
            data_rows: list[list[object]] = []
            for row in rows[1:]:
                cells = (
                    [str(row[0] if row and row[0] is not None else "")]
                    if first_cell_only
                    else [str(cell) if cell is not None else "" for cell in row]
                )
                if keywords and any(
                    cell.strip().casefold().startswith(keyword)
                    for cell in cells
                    for keyword in keywords
                ):
                    continue
                data_rows.append(list(row))
            tables.append((unique_headers(rows[0]), data_rows))
        _progress(progress, file_index, len(source_values), path.name)
    target = Path(output_path).expanduser().resolve()
    if is_cancelled and is_cancelled():
        return target
    if not tables:
        raise ToolError("Select at least one CSV or Excel file.")

    headers: list[object] = []
    for source_headers, _rows in tables:
        for header in source_headers:
            if header not in headers:
                headers.append(header)
    combined: list[list[object]] = []
    for source_headers, rows in tables:
        positions = {header: index for index, header in enumerate(source_headers)}
        for row in rows:
            merged_row = [
                row[position] if (position := positions.get(header)) is not None and position < len(row) else None
                for header in headers
            ]
            if not blank_row(merged_row):
                combined.append(merged_row)

    if target.suffix.lower() == ".csv":
        with target.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(combined)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet1"
        sheet.append(headers)
        for row in combined:
            sheet.append(row)
        workbook.save(target)
    return target


def _barcode_type_key(value: object) -> str:
    """Normalize pyzbar and zbarimg names (for example QR-Code/QRCODE)."""
    return "".join(character for character in str(value).upper() if character.isalnum())


def scan_barcodes(
    source_path: str | os.PathLike[str],
    pages: Iterable[int] | None = None,
    dpi: int = 180,
    is_cancelled: Callable[[], bool] | None = None,
    barcode_types: Iterable[str] | None = None,
) -> list[dict[str, object]]:
    capability = detect_capabilities()[CapabilityId.BARCODE]
    if not capability.available:
        raise ToolError(capability.reason)

    allowed_types = (
        None if barcode_types is None else {_barcode_type_key(value) for value in barcode_types}
    )
    results: list[dict[str, object]] = []
    with fitz.open(source_path) as doc:
        selected = list(pages) if pages is not None else list(range(doc.page_count))
        if capability.backend == "zbarimg CLI":
            with tempfile.TemporaryDirectory(prefix="pdfdocuedit-zbar-") as folder:
                for index in selected:
                    if is_cancelled and is_cancelled():
                        break
                    if not 0 <= index < doc.page_count:
                        continue
                    page = doc.load_page(index)
                    pixmap = page.get_pixmap(dpi=max(72, min(600, dpi)), alpha=False)
                    image_path = Path(folder) / f"page-{index + 1}.png"
                    pixmap.save(image_path)
                    process = PlatformService.run(
                        [capability.path, "--xml", "--quiet", str(image_path)],
                        timeout=60,
                    )
                    if process.returncode not in {0, 4}:
                        raise ToolError(
                            process.stderr.strip() or "zbarimg could not scan the page."
                        )
                    if not process.stdout.strip():
                        continue  # no symbols on this page
                    try:
                        root = ET.fromstring(process.stdout)
                    except ET.ParseError as exc:
                        raise ToolError("zbarimg returned invalid barcode data.") from exc
                    for symbol in root.findall(".//{*}symbol"):
                        type_name = symbol.get("type", "Unknown")
                        if allowed_types is not None and _barcode_type_key(type_name) not in allowed_types:
                            continue
                        data = symbol.find("{*}data")
                        results.append(
                            {
                                "page": index + 1,
                                "type": type_name,
                                "data": "" if data is None else "".join(data.itertext()),
                            }
                        )
            return results

        from PIL import Image
        from pyzbar.pyzbar import decode

        for index in selected:
            if is_cancelled and is_cancelled():
                break
            if not 0 <= index < doc.page_count:
                continue
            page = doc.load_page(index)
            pixmap = page.get_pixmap(dpi=max(72, min(600, dpi)), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            for value in decode(image):
                if allowed_types is not None and _barcode_type_key(value.type) not in allowed_types:
                    continue
                results.append(
                    {
                        "page": index + 1,
                        "type": value.type,
                        "data": value.data.decode("utf-8", errors="replace"),
                    }
                )
    return results


def scan_barcodes_batch(
    paths: Iterable[str | os.PathLike[str]],
    page_range: str = "",
    dpi: int = 180,
    progress: ProgressCallback | None = None,
    is_cancelled: Callable[[], bool] | None = None,
    barcode_types: Iterable[str] | None = None,
) -> list[dict[str, object]]:
    from .pdf_engine import parse_page_range

    sources = [Path(value).expanduser().resolve() for value in paths]
    combined: list[dict[str, object]] = []
    matched_pages = not page_range.strip()
    for index, source in enumerate(sources, 1):
        if is_cancelled and is_cancelled():
            break
        _progress(progress, index - 1, len(sources), source.name)
        with fitz.open(source) as document:
            pages = (
                parse_page_range(page_range, document.page_count)
                if page_range.strip()
                else list(range(document.page_count))
            )
        matched_pages = matched_pages or bool(pages)
        for result in scan_barcodes(source, pages, dpi, is_cancelled, barcode_types):
            result["file"] = source.name
            result["path"] = str(source)
            combined.append(result)
        _progress(progress, index, len(sources), source.name)
    if page_range.strip() and not matched_pages:
        raise ToolError("The page range does not match any page in the selected PDFs.")
    return combined
