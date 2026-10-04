"""Read Excel worksheets without Excel automation, macros or formula evaluation."""
from __future__ import annotations

import math
import re
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from composition.template.model import CompositionError

EXCEL_SUFFIXES = frozenset({".xlsx", ".xls"})


def is_excel(path):
    return Path(path).suffix.lower() in EXCEL_SUFFIXES


def cell_text(value, number_format="", preserve_zeros=True):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        has_time_format = bool(re.search(r"[hs]", re.sub(r'"[^"]*"|\\.', "", number_format.lower())))
        return value.date().isoformat() if value.time() == time() and not has_time_format else value.isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, int | float):
        if not math.isfinite(value):
            raise CompositionError("Non-finite numeric Excel value.")
        # Keep plain decimal strings for existing numeric rules, without locale/currency notation.
        text = format(Decimal(str(value)), "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        if preserve_zeros and re.fullmatch(r"0{1,32}", number_format) and value == int(value):
            digits = str(abs(int(value))).zfill(len(number_format))
            text = ("-" if value < 0 else "") + digits
        return text
    raise CompositionError(f"Unsupported Excel cell value: {type(value).__name__}")


def workbook_info(path):
    source = Path(path)
    try:
        if source.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            try:
                sheets = [{"name": sheet.title, "hidden": sheet.sheet_state != "visible"}
                          for sheet in book.worksheets]
            finally:
                book.close()
        elif source.suffix.lower() == ".xls":
            import xlrd
            book = xlrd.open_workbook(source, on_demand=True)
            try:
                sheets = [{"name": name, "hidden": None} for name in book.sheet_names()]
            finally:
                book.release_resources()
        else:
            raise CompositionError("Supported Excel formats: .xlsx and .xls.")
    except Exception as exc:
        raise CompositionError(f"Unable to open Excel workbook: {exc}") from exc
    if not sheets:
        raise CompositionError("The workbook contains no worksheets.")
    default = next((s["name"] for s in sheets if not s["hidden"]), sheets[0]["name"])
    return {"sheets": sheets, "default_sheet": default}


@contextmanager
def excel_rows(config, diagnostics):
    """Yield physical row ordinal and text cells; all readers closed on error/cancel."""
    if config.excel_formulas not in ("reject", "cached") or type(config.preserve_zeros) is not bool:
        raise CompositionError("Invalid Excel formula/number-format configuration.")
    if not isinstance(config.sheet, str) or len(config.sheet) > 255:
        raise CompositionError("Invalid Excel worksheet name.")
    source = Path(config.path)
    if source.suffix.lower() == ".xlsx":
        with _xlsx_rows(config, diagnostics) as rows:
            yield rows
    elif source.suffix.lower() == ".xls":
        with _xls_rows(config, diagnostics) as rows:
            yield rows
    else:
        raise CompositionError("Supported Excel formats: .xlsx and .xls.")


@contextmanager
def _xlsx_rows(config, diagnostics):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    books = []
    try:
        book = load_workbook(config.path, read_only=True, data_only=False, keep_links=False)
        books.append(book)
        name = config.sheet or next((s.title for s in book.worksheets if s.sheet_state == "visible"), "")
        if name not in [s.title for s in book.worksheets]:
            raise CompositionError(f"Worksheet not found: {name}. Select an available worksheet.")
        sheet = book[name]
        # Real XML rows are authoritative; stale exported dimension metadata can be incorrect.
        sheet.reset_dimensions()
        cached_rows = None
        if config.excel_formulas == "cached":
            cache = load_workbook(config.path, read_only=True, data_only=True, keep_links=False)
            books.append(cache)
            cache[name].reset_dimensions()
            cached_rows = iter(cache[name].iter_rows(min_row=config.header_row))
            diagnostics["warnings"].append("Formula values use saved Excel results; their freshness cannot be verified. Recalculate and save the workbook before import.")
        diagnostics["sheet"] = name
        def rows():
            for row_number, cells in enumerate(sheet.iter_rows(min_row=config.header_row), config.header_row):
                cached = next(cached_rows, ()) if cached_rows is not None else ()
                if len(cells) > 1000:
                    raise CompositionError(f"Worksheet {name}, row {row_number}: more than 1,000 columns. Remove unused columns or export a smaller range.")
                values = []
                for column, cell in enumerate(cells):
                    location = f"Worksheet {name}, cell {get_column_letter(column+1)}{row_number}"
                    value = cell.value
                    if cell.data_type == "f":
                        if config.excel_formulas == "reject":
                            raise CompositionError(location + ": formula found. Enable Use saved formula results or replace formulas with values.")
                        if column >= len(cached) or cached[column].value is None:
                            raise CompositionError(location + ": formula has no saved result. Recalculate and save in Excel, or replace it with a value.")
                        if cached[column].data_type == "e":
                            raise CompositionError(location + f": saved formula error {cached[column].value}.")
                        value = cached[column].value
                    elif cell.data_type == "e":
                        raise CompositionError(location + f": Excel error {value}.")
                    try:
                        values.append(cell_text(value, cell.number_format, config.preserve_zeros))
                    except CompositionError as exc:
                        raise CompositionError(location + ": " + str(exc)) from exc
                while values and values[-1] == "":
                    values.pop()
                yield row_number, values
        yield rows()
    except CompositionError:
        raise
    except Exception as exc:
        raise CompositionError(f"Unable to read Excel worksheet: {exc}") from exc
    finally:
        for book in books:
            book.close()


@contextmanager
def _xls_rows(config, diagnostics):
    import xlrd
    if config.excel_formulas != "cached":
        raise CompositionError("Legacy .xls exposes saved cell values only. Enable Use saved formula results, or save as .xlsx for formula detection.")
    book = None
    try:
        book = xlrd.open_workbook(config.path, on_demand=True, formatting_info=True)
        name = config.sheet or book.sheet_names()[0]
        if name not in book.sheet_names():
            raise CompositionError(f"Worksheet not found: {name}. Select an available worksheet.")
        sheet = book.sheet_by_name(name)
        if sheet.ncols > 1000:
            raise CompositionError("The worksheet contains more than 1,000 columns.")
        diagnostics["sheet"] = name
        diagnostics["warnings"].append("Legacy .xls uses saved cell values. Formula identity, missing caches and freshness cannot be audited; use .xlsx for stricter formula checks.")
        def rows():
            for row_index in range(config.header_row-1, sheet.nrows):
                values = []
                for column in range(sheet.ncols):
                    cell = sheet.cell(row_index, column)
                    location = f"Worksheet {name}, row {row_index+1}, column {column+1}"
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        raise CompositionError(location + f": Excel error code {cell.value}.")
                    value = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        converted = xlrd.xldate_as_datetime(value, book.datemode)
                        value = converted.time() if 0 <= cell.value < 1 else converted
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    values.append(cell_text(value, fmt, config.preserve_zeros))
                while values and values[-1] == "":
                    values.pop()
                yield row_index+1, values
        yield rows()
    except CompositionError:
        raise
    except Exception as exc:
        raise CompositionError(f"Unable to read Excel worksheet: {exc}") from exc
    finally:
        if book is not None:
            book.release_resources()
