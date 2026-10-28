"""Quick extraction must preserve selected-page data without production approval."""
import csv
from dataclasses import replace

import fitz
import pytest
from openpyxl import load_workbook

from core.cancellation import TaskCancelled
from workflow.extraction import ExtractionSpec, Region


def fixture_pdf(path):
    with fitz.open() as pdf:
        for text in ("0000123\nSecond line", "=1+1\nThird line", "Last page"):
            page = pdf.new_page(width=300, height=400)
            page.insert_text((30, 40), text)
        pdf.new_page(width=400, height=300)
        pdf.new_page(width=300, height=400)
        pdf.save(path)
    return ExtractionSpec([Region(name="Code", x_mm=0, y_mm=0, width_mm=100,
                                 height_mm=100, page_width_mm=300*25.4/72,
                                 page_height_mm=400*25.4/72)])


def api():
    from core import text_extraction
    return text_extraction


def test_selected_pages_keep_multiline_in_one_field_and_raw(tmp_path):
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    result = service.scan(source, spec, [0, 2], tmp_path / "job.sqlite")
    with service.QuickStore(result.database) as store:
        rows = store.rows(0, 100)
        assert [r["page"] for r in rows] == [1, 3]
        assert rows[0]["values"] == {"Code": "0000123 Second line"}
        assert store.details(1)[0]["raw"] == "0000123\nSecond line"
        assert store.metadata()["accepted"] is False
    spec.regions[0] = replace(spec.regions[0], join_lines=False)
    result = service.scan(source, spec, [0], tmp_path / "lines.sqlite")
    with service.QuickStore(result.database) as store:
        assert store.rows()[0]["values"]["Code"] == "0000123\nSecond line"


def test_mismatch_empty_and_invalid_values_keep_rows(tmp_path):
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    spec.regions[0].format = "digits"
    result = service.scan(source, spec, [0, 3, 4], tmp_path / "job.sqlite")
    with service.QuickStore(result.database) as store:
        rows = store.rows()
        assert len(rows) == 3
        assert "digits" in rows[0]["issues"]
        assert "Geometry Mismatch" in rows[1]["issues"]
        assert "No extractable text" in rows[2]["issues"]
        assert store.rows(1, 1)[0]["page"] == 4


def test_export_string_types_raw_issues_and_formula_protection(tmp_path):
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    result = service.scan(source, spec, [0, 1, 3], tmp_path / "job.sqlite")
    output = tmp_path / "result.xlsx"
    service.export(result.database, source, spec, [0, 1, 3], output)
    book = load_workbook(output)
    assert book.sheetnames == ["Data", "Details", "Issues"]
    assert book["Data"]["D2"].value == "0000123 Second line"
    assert book["Data"]["D3"].value == "=1+1 Third line"
    assert book["Data"]["D3"].data_type == "s"
    assert book["Details"]["C2"].value == "0000123\nSecond line"
    assert book["Issues"].max_row == 2
    book.close()
    service.export(result.database, source, spec, [0, 1, 3], tmp_path / "result.csv")
    with (tmp_path / "result.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert rows[1]["Code"] == "'=1+1 Third line"
    assert rows[0]["Code"] == "0000123 Second line"
    assert rows[0]["Raw: Code"] == "0000123\nSecond line"


def test_cancel_preserves_previous_database_and_export(tmp_path):
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    db = tmp_path / "job.sqlite"
    db.write_bytes(b"previous")
    with pytest.raises(TaskCancelled):
        service.scan(source, spec, [0, 1], db, is_cancelled=lambda: True)
    assert db.read_bytes() == b"previous"
    service.scan(source, spec, [0], db)
    target = tmp_path / "result.xlsx"
    target.write_bytes(b"old output")
    with pytest.raises(TaskCancelled):
        service.export(db, source, spec, [0], target, overwrite=True, is_cancelled=lambda: True)
    assert target.read_bytes() == b"old output"


def test_stale_source_config_selection_and_protected_outputs(tmp_path):
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    db = tmp_path / "job.sqlite"
    service.scan(source, spec, [0], db)
    changed = ExtractionSpec([replace(spec.regions[0], join_lines=False)])
    for config, pages, target in ((changed, [0], tmp_path/"a.csv"),
                                  (spec, [1], tmp_path/"b.csv"),
                                  (spec, [0], source), (spec, [0], db)):
        with pytest.raises(ValueError):
            service.export(db, source, config, pages, target)
    with source.open("ab") as stream:
        stream.write(b"\n% changed")
    with pytest.raises(ValueError, match="changed|stale"):
        service.export(db, source, spec, [0], tmp_path/"c.csv")


def test_xlsx_refuses_overlong_or_illegal_text_without_truncation(tmp_path):
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    db = tmp_path / "job.sqlite"
    service.scan(source, spec, [0], db)
    import sqlite3
    for value in ("a"*32768, "text\x01"):
        with sqlite3.connect(db) as conn:
            conn.execute("UPDATE cells SET raw=?,value=?", (value, value))
        target = tmp_path / "result.xlsx"
        target.write_bytes(b"old")
        with pytest.raises(ValueError, match="XLSX"):
            service.export(db, source, spec, [0], target, overwrite=True)
        assert target.read_bytes() == b"old"
        service.export(db, source, spec, [0], tmp_path/"result.csv", overwrite=True)
        with (tmp_path/"result.csv").open(encoding="utf-8-sig", newline="") as stream:
            assert list(csv.DictReader(stream))[0]["Code"] == value


def test_publication_receipt_survives_late_cancel(tmp_path):
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    db = tmp_path / "job.sqlite"
    service.scan(source, spec, [0], db)
    receipts = []
    target = tmp_path/"result.txt"
    service.export(db, source, spec, [0], target, on_committed=receipts.append,
                   is_cancelled=lambda: target.exists())
    assert target.exists()
    assert receipts == [str(target)]


def test_worker_dispatch_and_workflow_selection_preserve_provenance(tmp_path):
    from composition.worker import dispatch
    service = api()
    source = tmp_path / "source.pdf"
    spec = fixture_pdf(source)
    result = dispatch({"task": "quick_extract", "operation": "scan", "source": str(source),
                       "extraction": spec.to_dict(), "pages": [0, 2], "target": str(tmp_path/"data.sqlite")})
    assert result["pages"] == 2
    rows = dispatch({"task": "quick_extract", "operation": "rows", "database": result["database"], "offset": 0})
    assert [r["page"] for r in rows["rows"]] == [1, 3]
    workflow = service.workflow_spec(source, spec, [0, 2])
    workflow.validate()
    assert workflow.node("input").params["paths"] == [str(source)]
    assert workflow.node("extract").params == spec.to_dict()
    assert workflow.node("merge").params["pages"] == {str(source): "1,3"}
    assert workflow.node("output").params == {}


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_rotation_cropbox_uses_visible_region_and_exports_exact_value(tmp_path, rotation):
    service = api()
    source = tmp_path/"rotated.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page(width=400, height=500)
        page.insert_text((90, 110), "0000123")
        page.set_cropbox(fitz.Rect(20, 30, 380, 470))
        page.set_rotation(rotation)
        area = page.search_for("0000123")[0]*page.rotation_matrix
        region = Region(name="Number", x_mm=area.x0*25.4/72, y_mm=area.y0*25.4/72,
                        width_mm=area.width*25.4/72, height_mm=area.height*25.4/72,
                        page_width_mm=page.rect.width*25.4/72, page_height_mm=page.rect.height*25.4/72)
        pdf.save(source)
    spec = ExtractionSpec([region])
    result = service.scan(source, spec, [0], tmp_path/"job.sqlite")
    with service.QuickStore(result.database) as store:
        assert store.rows()[0]["values"] == {"Number": "0000123"}


def test_cancel_during_xlsx_write_preserves_old_file_and_cleans_staging(tmp_path):
    service = api()
    source = tmp_path/"source.pdf"
    spec = fixture_pdf(source)
    db = tmp_path/"job.sqlite"
    service.scan(source, spec, [0, 1, 2], db)
    target = tmp_path/"result.xlsx"
    target.write_bytes(b"previous report")
    stop = False
    def progress(*_):
        nonlocal stop
        stop = True
    with pytest.raises(TaskCancelled):
        service.export(db, source, spec, [0, 1, 2], target, overwrite=True,
                       progress=progress, is_cancelled=lambda: stop)
    assert target.read_bytes() == b"previous report"
    assert not list(tmp_path.glob(".result-*.xlsx"))


def test_existing_destination_requires_explicit_overwrite_and_csv_whitespace_formula(tmp_path):
    service = api()
    source = tmp_path/"source.pdf"
    spec = fixture_pdf(source)
    db = tmp_path/"job.sqlite"
    service.scan(source, spec, [0], db)
    target = tmp_path/"result.csv"
    target.write_bytes(b"previous report")
    with pytest.raises(FileExistsError):
        service.export(db, source, spec, [0], target)
    assert target.read_bytes() == b"previous report"
    import sqlite3
    with sqlite3.connect(db) as conn:
        conn.execute("UPDATE cells SET raw=?,value=?", ("\t=1+1", "  @SUM(1)"))
    service.export(db, source, spec, [0], target, overwrite=True)
    with target.open(encoding="utf-8-sig", newline="") as stream:
        row = next(csv.DictReader(stream))
    assert row["Code"] == "'  @SUM(1)"
    assert row["Raw: Code"] == "'\t=1+1"


def test_home_alias_cannot_bypass_source_overwrite_guard(tmp_path, monkeypatch):
    from pathlib import Path
    service = api()
    source = tmp_path/"source.pdf"
    spec = fixture_pdf(source)
    original = source.read_bytes()
    real_expanduser = Path.expanduser
    def isolated_home(path):
        if path.parts and path.parts[0] == "~":
            return tmp_path.joinpath(*path.parts[1:])
        return real_expanduser(path)
    # Resolve the alias inside the fixture, without changing HOME/USERPROFILE
    # or accessing any actual user document. Atomic output expands this alias.
    monkeypatch.setattr(Path, "expanduser", isolated_home)
    with pytest.raises(ValueError, match="overwrite"):
        service.scan(source, spec, [0], Path("~/source.pdf"))
    assert source.read_bytes() == original
