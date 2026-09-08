from __future__ import annotations

from pathlib import Path
from threading import Event, Thread

import fitz
import pytest

import core.pdf_engine as pdf_engine_module
import core.tools as tools_module
from core.annotations import AnnotationOp, apply_annotation
from core.capabilities import Capability, CapabilityId
from core.pdf_engine import DOCUMENT_LOCK, PdfEngine, PdfEngineError, parse_page_range
from core.pdf_io import PdfValidationError
from core.platform_service import ProcessResult
from core.tools import (
    ToolError,
    compress_pdf,
    compress_pdfs,
    create_page_count_report,
    deep_search,
    extract_region_text,
    merge_pdfs,
    merge_spreadsheets,
    overlay_pdf,
    overlay_pdfs,
    scan_barcodes,
    text_files_to_pdf,
    text_files_to_pdfs,
)


def test_batch_tools_refuse_in_place_overwrite(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "inplace.pdf")
    with pytest.raises(ToolError):
        compress_pdfs([source], tmp_path, suffix="", overwrite=False)
    with pytest.raises(ToolError):
        overlay_pdfs(source, [source], tmp_path, suffix="", overwrite=False)
    assert make_pdf(tmp_path / "inplace.pdf").exists()


def test_parse_page_range_keywords() -> None:
    assert parse_page_range("all", 5) == [0, 1, 2, 3, 4]
    assert parse_page_range("odd", 5) == [0, 2, 4]
    assert parse_page_range("even", 5) == [1, 3]
    assert parse_page_range("ALL", 3) == [0, 1, 2]
    assert parse_page_range("1,even", 5) == [0, 1, 3]
    with pytest.raises(ValueError):
        parse_page_range("odd")  # needs a known page count
    assert parse_page_range("12", 5) == []  # out of range pages are dropped


def make_pdf(path: Path, pages: int = 3, prefix: str = "Page") -> Path:
    with fitz.open() as document:
        for index in range(pages):
            page = document.new_page(width=595, height=842)
            page.insert_text((72, 96), f"{prefix} {index + 1} searchable text")
        document.save(path)
    return path


def make_malformed_metadata_pdf(path: Path, pages: int = 1) -> Path:
    """Create the invalid scanner metadata that PyMuPDF 1.26 exposes as surrogates."""

    make_pdf(path, pages=pages, prefix="Scanner")
    with fitz.open(path) as document:
        document.set_metadata({"producer": "placeholder"})
        document.saveIncr()
    with fitz.open(path) as document:
        kind, reference = document.xref_get_key(-1, "Info")
        assert kind == "xref"
        info_xref = int(reference.split()[0])
        producer = b"Adobe PSL 1.3e for Canon" + bytes((0xC0, 0x80))
        document.xref_set_key(info_xref, "Producer", f"<{producer.hex()}>")
        document.saveIncr()
    return path


def test_live_document_reader_waits_for_global_lock(tmp_path: Path) -> None:
    engine = PdfEngine()
    engine.open(make_pdf(tmp_path / "locked-reader.pdf", pages=1))
    started = Event()
    finished = Event()
    extracted: list[str] = []

    def read_document() -> None:
        started.set()
        extracted.append(engine.extract_text(0))
        finished.set()

    worker = Thread(target=read_document, daemon=True)
    with DOCUMENT_LOCK:
        worker.start()
        assert started.wait(1)
        assert not finished.wait(0.1)

    assert finished.wait(2)
    worker.join(timeout=1)
    assert "Page 1 searchable text" in extracted[0]
    engine.close()


def test_parse_page_range() -> None:
    assert parse_page_range("1, 3, 5-7", 6) == [0, 2, 4, 5]
    assert parse_page_range("4-2") == [1, 2, 3]
    assert parse_page_range("0, 99", 5) == []
    with pytest.raises(ValueError):
        parse_page_range("first-last")
    assert parse_page_range("1-999999999", 3) == [0, 1, 2]
    with pytest.raises(ValueError, match="too large"):
        parse_page_range("1-999999999")


def test_annotation_save_preserves_original_content_streams(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "annotation-safe.pdf", pages=1, prefix="QR-ABC123")
    original_file = source.read_bytes()
    with fitz.open(source) as document:
        original_streams = [
            document.xref_stream_raw(xref)
            for xref in document.load_page(0).get_contents()
        ]

    engine = PdfEngine()
    engine.open(source)
    assert engine.document is not None
    apply_annotation(
        engine.document,
        AnnotationOp(kind="rect", page=0, rects=(fitz.Rect(20, 20, 50, 50),)),
    )
    engine.mark_modified()
    engine.save()
    engine.close()

    assert source.read_bytes().startswith(original_file)
    with fitz.open(source) as document:
        page = document.load_page(0)
        assert [
            document.xref_stream_raw(xref) for xref in page.get_contents()
        ] == original_streams


def test_engine_edit_save_extract_split_and_search(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "source.pdf")
    engine = PdfEngine()
    engine.open(source)

    assert engine.page_count == 3
    assert len(engine.search_text("searchable")) == 3
    assert "Page 2" in engine.extract_text(1)

    engine.rotate_page(0, 90)
    engine.reorder_pages([2, 1, 0])
    engine.delete_page(1)
    assert engine.is_modified
    assert engine.page_count == 2

    extracted = engine.extract_pages([0], tmp_path / "extracted.pdf")
    parts = engine.split_pdf([(0, 0), (1, 1)], tmp_path / "parts")
    saved = engine.save()
    assert saved == source.resolve()
    assert not engine.is_modified
    assert len(parts) == 2
    with fitz.open(extracted) as document:
        assert document.page_count == 1
    with fitz.open(source) as document:
        assert document.page_count == 2
        assert document.load_page(0).rotation == 0
        assert "Page 3" in document.load_page(0).get_text()
    engine.close()


def test_engine_insert_encrypt_and_decrypt(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "base.pdf", 1, "Base")
    inserted = make_pdf(tmp_path / "insert.pdf", 2, "Insert")
    engine = PdfEngine()
    engine.open(source)
    engine.insert_pages(str(inserted), [0, 1], 1)
    assert engine.page_count == 3

    encrypted = engine.encrypt("secret", tmp_path / "encrypted.pdf")
    engine.close()
    with fitz.open(encrypted) as document:
        assert document.needs_pass
        assert document.authenticate("secret")

    engine.open(encrypted, "secret")
    decrypted = engine.decrypt(tmp_path / "decrypted.pdf")
    engine.close()
    with fitz.open(decrypted) as document:
        assert not document.needs_pass
        assert document.page_count == 3
        # Content must survive decryption intact (the old direct-save path
        # produced files whose streams failed to decrypt). The base page is
        # first; the inserted pages follow it.
        assert "Base 1" in document.load_page(0).get_text()
        assert "Insert 1" in document.load_page(1).get_text()


def test_engine_rejects_deleting_every_page(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "single.pdf", 1)
    engine = PdfEngine()
    engine.open(source)
    with pytest.raises(PdfEngineError):
        engine.delete_pages([0])
    engine.close()


def test_rotation_full_save_cleans_scanner_metadata_and_reopens(tmp_path: Path) -> None:
    source = make_malformed_metadata_pdf(tmp_path / "scanner.pdf", pages=2)
    original_bytes = source.read_bytes()
    engine = PdfEngine()
    engine.open(source)
    engine.rotate_page(0, 90)
    engine.save()
    engine.close()

    # Page-tree changes must be fully rewritten, not appended to the scanner
    # file as an incremental update.
    assert not source.read_bytes().startswith(original_bytes)
    with fitz.open(source) as document:
        assert document.page_count == 2
        assert document.load_page(0).rotation == 90
        assert document.metadata["producer"] == "Adobe PSL 1.3e for Canon"
        for page in document:
            page.get_pixmap(matrix=fitz.Matrix(0.2, 0.2), alpha=False)


def test_save_validation_failure_does_not_replace_original(
    tmp_path: Path, monkeypatch
) -> None:
    source = make_pdf(tmp_path / "protected.pdf", pages=1)
    original_bytes = source.read_bytes()
    engine = PdfEngine()
    engine.open(source)
    engine.rotate_page(0, 90)

    def reject_output(*_args, **_kwargs) -> None:
        raise PdfValidationError("simulated invalid output")

    monkeypatch.setattr(pdf_engine_module, "validate_pdf_file", reject_output)
    with pytest.raises(PdfValidationError, match="simulated invalid output"):
        engine.save()
    assert source.read_bytes() == original_bytes
    assert engine.is_modified
    engine.close()


def test_digital_signature_reference_is_detected() -> None:
    class SignatureWidget:
        field_type = fitz.PDF_WIDGET_TYPE_SIGNATURE
        field_value = ""
        xref = 17

    class SignaturePage:
        @staticmethod
        def widgets():
            return [SignatureWidget()]

    class SignatureDocument:
        page_count = 1

        @staticmethod
        def load_page(_page_number):
            return SignaturePage()

        @staticmethod
        def xref_object(_xref):
            return "<< /FT /Sig /V 23 0 R >>"

    engine = PdfEngine()
    engine._doc = SignatureDocument()
    assert engine.has_digital_signatures()
    engine._doc = None


def test_visual_organizer_transaction_and_repeat_insert(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "organize.pdf", 3, "Original")
    stamp = make_pdf(tmp_path / "stamp.pdf", 1, "Stamp")
    engine = PdfEngine()
    engine.open(source)
    engine.organize_pages([2, 0], {2: 90})
    assert engine.page_count == 2
    assert "Original 3" in engine.extract_text(0)
    assert engine.document is not None
    assert engine.document.load_page(0).rotation == 90
    engine.repeat_insert_pages(str(stamp), [0], 1)
    assert engine.page_count == 4
    engine.close()


def test_engine_rejects_noop_inserts_and_protects_split_outputs(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "source.pdf", 2)
    inserted = make_pdf(tmp_path / "insert.pdf", 1)
    engine = PdfEngine()
    engine.open(source)
    with pytest.raises(PdfEngineError, match="out of range"):
        engine.insert_pages(str(inserted), [99], 0)
    assert not engine.is_modified
    with pytest.raises(PdfEngineError, match="insertion points"):
        engine.repeat_insert_pages(str(inserted), [0], 99)
    assert not engine.is_modified

    outputs = engine.split_pdf([(0, 0), (1, 1)], tmp_path / "parts")
    assert len(outputs) == 2
    with pytest.raises(PdfEngineError, match="already exists"):
        engine.split_pdf([(0, 0), (1, 1)], tmp_path / "parts")
    replaced = engine.split_pdf([(0, 0), (1, 1)], tmp_path / "parts", overwrite=True)
    assert replaced == outputs
    engine.close()


def test_engine_snapshot_contains_current_unsaved_state(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "source.pdf", 1)
    inserted = make_pdf(tmp_path / "insert.pdf", 1, "Inserted")
    engine = PdfEngine()
    engine.open(source)
    engine.insert_pages(str(inserted), [0], 1)
    snapshot = engine.snapshot(tmp_path / "working-copy.pdf")
    with fitz.open(snapshot) as document:
        assert document.page_count == 2
        assert "Inserted" in document.load_page(1).get_text()
    assert engine.is_modified
    engine.close()


def test_document_tools(tmp_path: Path) -> None:
    first = make_pdf(tmp_path / "first.pdf", 1, "First")
    second = make_pdf(tmp_path / "second.pdf", 2, "Second")
    merged = merge_pdfs([first, second], tmp_path / "merged.pdf")
    compressed = compress_pdf(merged, tmp_path / "compressed.pdf")
    overlay = overlay_pdf(first, second, tmp_path / "overlay.pdf")

    with fitz.open(merged) as document:
        assert document.page_count == 3
    with fitz.open(compressed) as document:
        assert document.page_count == 3
    with fitz.open(overlay) as document:
        assert document.page_count == 2

    same_source = make_pdf(tmp_path / "same-source.pdf", 1)
    outputs = compress_pdfs([same_source], suffix="", overwrite=True)
    assert outputs == [same_source.resolve()]
    with fitz.open(same_source) as document:
        assert document.page_count == 1


def test_merge_and_extract_sanitize_invalid_scanner_metadata(tmp_path: Path) -> None:
    first = make_malformed_metadata_pdf(tmp_path / "scanner.pdf", pages=2)
    second = make_pdf(tmp_path / "normal.pdf", pages=1)

    merged = merge_pdfs([first, second], tmp_path / "merged-scans.pdf")
    with fitz.open(merged) as document:
        assert document.page_count == 3
        assert document.metadata["producer"] == "Adobe PSL 1.3e for Canon"

    engine = PdfEngine()
    engine.open(first)
    extracted = engine.extract_pages([0], tmp_path / "extracted-scan.pdf")
    engine.close()
    with fitz.open(extracted) as document:
        assert document.page_count == 1
        assert document.metadata["producer"] == "Adobe PSL 1.3e for Canon"


def test_text_to_pdf_supports_unicode(tmp_path: Path) -> None:
    source = tmp_path / "notes.txt"
    source.write_text(
        "PDFDocuEdit Pro\n中文內容\nProfessional output", encoding="utf-8"
    )
    output = text_files_to_pdf([source], tmp_path / "notes.pdf")
    with fitz.open(output) as document:
        assert document.page_count >= 1
        text = "".join(page.get_text() for page in document)
        assert "PDFDocuEdit Pro" in text

    second = tmp_path / "second.txt"
    second.write_text("Second file", encoding="utf-8")
    outputs = text_files_to_pdfs(
        [source, second], tmp_path, single=False, overwrite=True
    )
    assert [path.name for path in outputs] == ["notes.pdf", "second.pdf"]


def test_text_to_pdf_paginates_long_content_without_blank_tail(tmp_path: Path) -> None:
    source = tmp_path / "long.txt"
    source.write_text("START " + "word " * 8000 + "END", encoding="utf-8")
    output = text_files_to_pdf([source], tmp_path / "long.pdf")
    with fitz.open(output) as document:
        text = "".join(page.get_text() for page in document)
        assert document.page_count > 2
        assert all(page.get_text().strip() for page in document)
        assert "START" in text
        assert "END" in text


def test_text_batch_preflights_colliding_filenames(tmp_path: Path) -> None:
    first = tmp_path / "one" / "same.txt"
    second = tmp_path / "two" / "same.txt"
    first.parent.mkdir()
    second.parent.mkdir()
    first.write_text("First", encoding="utf-8")
    second.write_text("Second", encoding="utf-8")
    with pytest.raises(ToolError, match="same output"):
        text_files_to_pdfs([first, second], tmp_path, single=False, overwrite=True)


def test_merge_csv_and_xlsx_without_optional_data_stack(tmp_path: Path) -> None:
    from openpyxl import Workbook, load_workbook

    csv_path = tmp_path / "first.csv"
    csv_path.write_text("Name,Pages\nAlpha,2\n", encoding="utf-8")
    xlsx_path = tmp_path / "second.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Status"])
    sheet.append(["Beta", "Ready"])
    workbook.save(xlsx_path)

    output = merge_spreadsheets([csv_path, xlsx_path], tmp_path / "merged.xlsx")
    merged = load_workbook(output, read_only=True, data_only=True)
    try:
        assert merged.active.title == "sheet1"
        rows = list(merged.active.iter_rows(values_only=True))
    finally:
        merged.close()
    assert rows[0] == ("Name", "Pages", "Status")
    assert rows[1] == ("Alpha", 2, None)
    assert rows[2] == ("Beta", None, "Ready")


def test_merge_spreadsheets_preserves_numbers_ids_and_removes_blank_rows(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "numbers.csv"
    source.write_text(
        "Name,Count,Rate,Identifier\nAlpha,2,3.5,00123\n,,,\nBeta,-4,1e3,A-01\n",
        encoding="utf-8",
    )
    output = merge_spreadsheets([source], tmp_path / "numbers.xlsx")
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert workbook.active.title == "sheet1"
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()

    assert rows == [
        ("Name", "Count", "Rate", "Identifier"),
        ("Alpha", 2, 3.5, "00123"),
        ("Beta", -4, 1000.0, "A-01"),
    ]


def test_merge_spreadsheets_reads_legacy_xls_backend(
    tmp_path: Path, monkeypatch
) -> None:
    import sys
    from types import SimpleNamespace

    class Cell:
        def __init__(self, value, cell_type):
            self.value = value
            self.ctype = cell_type

    values = [
        [Cell("Name", 1), Cell("Count", 1)],
        [Cell("Legacy", 1), Cell(7.0, 2)],
        [Cell("", 0), Cell("", 0)],
    ]

    class Sheet:
        nrows = len(values)
        ncols = len(values[0])

        @staticmethod
        def cell(row, column):
            return values[row][column]

    class Workbook:
        datemode = 0

        @staticmethod
        def sheet_by_index(index):
            assert index == 0
            return Sheet()

        @staticmethod
        def release_resources():
            return None

    fake_xlrd = SimpleNamespace(
        XL_CELL_EMPTY=0,
        XL_CELL_TEXT=1,
        XL_CELL_NUMBER=2,
        XL_CELL_DATE=3,
        XL_CELL_BOOLEAN=4,
        XL_CELL_ERROR=5,
        XL_CELL_BLANK=6,
        open_workbook=lambda path, on_demand: Workbook(),
        xldate=SimpleNamespace(xldate_as_datetime=lambda value, datemode: value),
    )
    monkeypatch.setitem(sys.modules, "xlrd", fake_xlrd)
    source = tmp_path / "legacy.xls"
    source.write_bytes(b"legacy-test-double")

    output = merge_spreadsheets([source], tmp_path / "legacy-output.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    assert rows == [("Name", "Count"), ("Legacy", 7)]


def test_merge_rules_and_region_text_export(tmp_path: Path) -> None:
    csv_path = tmp_path / "rules.csv"
    csv_path.write_text(
        "Ignore this row\nName,Value\nAlpha,1\nTotal,1\nBeta,2\n",
        encoding="utf-8",
    )
    merged = merge_spreadsheets(
        [csv_path],
        tmp_path / "rules-output.csv",
        skip_rows=1,
        exclude_keywords=["Total"],
    )
    assert "Total" not in merged.read_text(encoding="utf-8-sig")
    assert "Alpha" in merged.read_text(encoding="utf-8-sig")

    source = make_pdf(tmp_path / "region.pdf", 2, "Region")
    output = extract_region_text(
        source,
        [0, 1],
        (0, 0, 595, 200),
        tmp_path / "region.txt",
        excel=False,
    )
    text = output.read_text(encoding="utf-8")
    assert "Region 1" in text
    assert "Region 2" in text


def test_uppercase_pdf_discovery_for_report_and_deep_search(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    source = make_pdf(tmp_path / "UPPER.PDF", 2, "Uppercase")
    report = create_page_count_report(tmp_path, tmp_path / "report.xlsx")
    workbook = load_workbook(report, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    source_row = next(row for row in rows[1:] if row[0] == source.name)
    assert source_row[1] == 2
    assert source_row[3].endswith(" mm")
    results = deep_search(tmp_path, "searchable")
    assert [result["filename"] for result in results] == ["UPPER.PDF"]

    corrupt = tmp_path / "corrupt.PDF"
    corrupt.write_bytes(b"not a pdf")
    results = deep_search(tmp_path, "searchable")
    failed = next(result for result in results if result["filename"] == corrupt.name)
    assert failed["error"]


def test_merge_spreadsheets_preserves_duplicate_and_blank_columns(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "duplicate.csv"
    source.write_text("Name,Name,\nA,B,C\n", encoding="utf-8")
    output = merge_spreadsheets([source], tmp_path / "duplicate.xlsx")
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    assert rows[0] == ("Name", "Name (2)", "Column 3")
    assert rows[1] == ("A", "B", "C")


def test_barcode_cli_fallback_parses_zbar_xml(tmp_path: Path, monkeypatch) -> None:
    source = make_pdf(tmp_path / "barcode.pdf", 1)
    capability = Capability(
        CapabilityId.BARCODE,
        "Barcode / QR Code",
        True,
        "zbarimg CLI",
        "/fake/zbarimg",
    )
    monkeypatch.setattr(
        tools_module,
        "detect_capabilities",
        lambda: {CapabilityId.BARCODE: capability},
    )
    xml = """<barcodes xmlns='http://zbar.sourceforge.net/2008/barcode'>
    <source><index num='0'><symbol type='QR-Code'><data>PDFDocuEdit</data></symbol></index></source>
    </barcodes>"""
    monkeypatch.setattr(
        tools_module.PlatformService,
        "run",
        lambda *_args, **_kwargs: ProcessResult(0, xml, ""),
    )
    assert scan_barcodes(source) == [
        {"page": 1, "type": "QR-Code", "data": "PDFDocuEdit"}
    ]
    assert scan_barcodes(source, barcode_types=["QRCODE"])[0]["data"] == "PDFDocuEdit"
    assert scan_barcodes(source, barcode_types=["CODE128"]) == []


@pytest.mark.parametrize(
    ("returncode", "raises_error"),
    ((4, False), (1, True)),
)
def test_barcode_cli_distinguishes_no_symbol_from_processing_error(
    tmp_path: Path, monkeypatch, returncode: int, raises_error: bool
) -> None:
    source = make_pdf(tmp_path / f"barcode-exit-{returncode}.pdf", 1)
    capability = Capability(
        CapabilityId.BARCODE,
        "Barcode / QR Code",
        True,
        "zbarimg CLI",
        "/fake/zbarimg",
    )
    monkeypatch.setattr(
        tools_module,
        "detect_capabilities",
        lambda: {CapabilityId.BARCODE: capability},
    )
    monkeypatch.setattr(
        tools_module.PlatformService,
        "run",
        lambda *_args, **_kwargs: ProcessResult(returncode, "", "decode failed"),
    )

    if raises_error:
        with pytest.raises(ToolError, match="decode failed"):
            scan_barcodes(source)
    else:
        assert scan_barcodes(source) == []


# --- encryption round trips ------------------------------------------------


def test_encrypt_open_round_trip_with_correct_password(tmp_path: Path) -> None:
    source = make_pdf(tmp_path / "plain.pdf", 2, "Secret")
    engine = PdfEngine()
    engine.open(str(source))
    output = engine.encrypt("secret123", str(tmp_path / "encrypted.pdf"))

    direct = fitz.open(output)
    assert direct.needs_pass
    assert direct.authenticate("secret123")
    direct.close()

    reopened = PdfEngine()
    reopened.open(str(output), "secret123")
    assert reopened.page_count == 2
    assert "Secret 1" in reopened.document.load_page(0).get_text()
    assert reopened.is_encrypted()
    assert reopened.password == "secret123"


def test_encrypt_open_rejects_wrong_password(tmp_path: Path) -> None:
    from core.pdf_engine import PdfInvalidPassword

    source = make_pdf(tmp_path / "plain.pdf")
    engine = PdfEngine()
    engine.open(str(source))
    output = engine.encrypt("secret123", str(tmp_path / "encrypted.pdf"))
    with pytest.raises(PdfInvalidPassword):
        PdfEngine().open(str(output), "wrong-password")


def test_open_password_failure_keeps_current_document(tmp_path: Path) -> None:
    """A cancelled or wrong password must not close the open document."""
    from core.pdf_engine import PdfInvalidPassword, PdfPasswordRequired

    current = make_pdf(tmp_path / "current.pdf", 1, "Still here")
    locked = make_pdf(tmp_path / "locked.pdf", 1, "Locked")
    encrypting = PdfEngine()
    encrypting.open(str(locked))
    encrypted = encrypting.encrypt("secret123", str(tmp_path / "locked_enc.pdf"))

    engine = PdfEngine()
    engine.open(str(current))

    with pytest.raises(PdfPasswordRequired):
        engine.open(str(encrypted), None)
    assert engine.is_loaded()
    assert "Still here" in engine.document.load_page(0).get_text()

    with pytest.raises(PdfInvalidPassword):
        engine.open(str(encrypted), "wrong")
    assert engine.is_loaded()
    assert "Still here" in engine.document.load_page(0).get_text()


@pytest.mark.parametrize("operation", ["plan", "insert"])
@pytest.mark.parametrize("restore_fails", [False, True])
def test_mutation_rollback_integrity(tmp_path, monkeypatch, operation, restore_fails):
    from core.pdf_engine import PagePlanEntry

    source = make_pdf(tmp_path / "source.pdf", 6, "Source")
    engine = PdfEngine()
    engine.open(make_pdf(tmp_path / "destination.pdf", 3, "Original"))
    engine.rotate_pages([0], 90)
    before = [(p.get_text(), p.rotation) for p in engine.document]
    context = (engine.original_path, engine.password, engine.is_modified,
               engine.document_id, engine.revision, engine._saved_permissions,
               engine._reencrypt_on_save, engine._requires_full_save)
    real_insert = fitz.Document.insert_pdf
    real_open = fitz.open
    calls = 0

    def fail_second_insert(doc, *args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("injected mutation failure")
        return real_insert(doc, *args, **kwargs)

    def fail_restore(*args, **kwargs):
        if restore_fails and calls == 2 and "stream" in kwargs:
            raise RuntimeError("injected restore failure")
        return real_open(*args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(fitz.Document, "insert_pdf", fail_second_insert)
        patch.setattr(fitz, "open", fail_restore)
        message = "broken state; reopen" if restore_fails else "state was restored"
        with pytest.raises(PdfEngineError, match=message) as error:
            if operation == "insert":
                engine.insert_pages(str(source), [5, 2], 1)
            else:
                engine.apply_page_plan([
                    PagePlanEntry("original", "current", 0),
                    PagePlanEntry("one", "external", 5, str(source)),
                    PagePlanEntry("two", "external", 2, str(source)),
                ])
        assert isinstance(error.value.__cause__, RuntimeError)
    if restore_fails:
        assert not engine.is_loaded()
        engine.open(source)
        assert engine.page_count == 6
    else:
        assert engine.page_count == 3
        assert [(p.get_text(), p.rotation) for p in engine.document] == before
        assert context == (engine.original_path, engine.password, engine.is_modified,
                           engine.document_id, engine.revision, engine._saved_permissions,
                           engine._reencrypt_on_save, engine._requires_full_save)
        output = engine.save_as(tmp_path / "restored.pdf")
        with fitz.open(output) as saved:
            assert [(p.get_text(), p.rotation) for p in saved] == before
        engine.insert_pages(str(source), [0], 0)
        assert engine.page_count == 4
    engine.close()


def test_insert_pages_preserves_caller_order_and_duplicates(tmp_path):
    source = make_pdf(tmp_path / "source.pdf", 6, "Source")
    engine = PdfEngine()
    engine.open(make_pdf(tmp_path / "destination.pdf", 1, "Original"))
    engine.insert_pages(str(source), [5, 2, 5], 0)
    assert [page.get_text().splitlines()[0] for page in engine.document] == [
        "Source 6 searchable text", "Source 3 searchable text",
        "Source 6 searchable text", "Original 1 searchable text",
    ]
    engine.close()
