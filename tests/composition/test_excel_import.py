from __future__ import annotations

import copy
import json
import sqlite3
import zipfile
from contextlib import closing
from datetime import date, datetime, time

import fitz
import pytest
from openpyxl import Workbook
from PyQt6.QtWidgets import QApplication, QDialogButtonBox

from composition.data.excel_source import cell_text, workbook_info
from composition.data.source import import_records, sample_records, suggest_import
from composition.designer.data_dialog import DataDialog
from composition.designer.workspace import CompositionWindow
from composition.production.generator import generate
from composition.production.model import ProductionJob
from composition.template.model import CompositionError, DataConfig, Element, FontSpec, SequenceSpec, Template
from composition.template.serializer import load_project, save_project
from scripts.composition_excel_fixture import make_xls, make_xlsx
from tests.composition.test_designer_controls import cleanup, wait


def write_rows(path, rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


def set_formula_cache(path, value=None, error=False):
    """Create saved OOXML calculation results without an Excel/eval dependency."""
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    for cell in root.findall(".//s:c", ns):
        if cell.find("s:f", ns) is not None:
            cache = cell.find("s:v", ns)
            cache.text = value
            if error:
                cell.set("t", "e")
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, raw in entries.items():
            archive.writestr(name, raw)


@pytest.mark.parametrize("value,fmt,preserve,expected", [
    ("000123", "", True, "000123"), (123, "000000", True, "000123"),
    (123, "000000", False, "123"), (-12, "0000", True, "-0012"),
    (0.25, "0.00", True, "0.25"), (1e-7, "General", True, "0.0000001"),
    (True, "", True, "TRUE"), (False, "", True, "FALSE"), (None, "", True, ""),
    (date(2026, 10, 2), "", True, "2026-10-02"),
    (datetime(2026, 10, 2, 9, 30), "", True, "2026-10-02 09:30:00"),
    (time(9, 30), "", True, "09:30:00"),
    (datetime(2026, 10, 2), "yyyy-mm-dd hh:mm", True, "2026-10-02 00:00:00"),
])
def test_explicit_excel_text_conversion(value, fmt, preserve, expected):
    assert cell_text(value, fmt, preserve) == expected


def test_xlsx_sheets_header_mapping_zeros_dates_empty_and_source_rows(tmp_path):
    path = make_xlsx(tmp_path/"來源資料.xlsx", 3)
    suggested = suggest_import(path)
    assert suggested.sheet == "Summary"
    assert [s["name"] for s in workbook_info(path)["sheets"]] == ["Summary", "Statements"]
    config = DataConfig(str(path), sheet="Statements", header_row=2, mapping={"Customer Name": "Name"})
    before = path.read_bytes()
    preview = sample_records(config)
    assert preview["originals"][0] == "Customer Name"
    store = import_records(config, tmp_path/"records.db")
    assert store.count == 3 and store.fields[0] == "Name"
    assert store.record(1) == {"Name": "陳小明", "Account": "000001", "Balance": "10000.25",
                               "Date": "2026-10-02", "Flag": "TRUE", "Note": ""}
    assert store.metadata["source"]["worksheet"] == "Statements"
    assert store.metadata["config"]["header_row"] == 2
    assert store.metadata["config"]["mapping"] == {"Customer Name": "Name"}
    with closing(sqlite3.connect(store.path)) as connection:
        assert connection.execute("SELECT source_row FROM records ORDER BY ordinal").fetchall() == [(3,), (4,), (5,)]
    assert path.read_bytes() == before


def test_actual_ole_biff_xls_exact_formats_unicode_dates_and_warning(tmp_path):
    source = make_xls(tmp_path/"legacy.xls", 3)
    config = suggest_import(source)
    assert config.excel_formulas == "cached" and config.sheet == "Statements"
    store = import_records(config, tmp_path/"legacy.db")
    assert store.count == 3
    assert store.record(3) == {"Name": "陳小明", "Account": "000003", "Date": "2026-10-02"}
    assert "cannot be audited" in store.metadata["warnings"][0]
    config.excel_formulas = "reject"
    with pytest.raises(CompositionError, match="saved cell values"):
        import_records(config, tmp_path/"reject.db")
    assert not (tmp_path/"reject.db").exists()


def test_formula_default_reject_cached_missing_and_saved_errors(tmp_path):
    source = write_rows(tmp_path/"formula.xlsx", [["Balance"], ["=10+20"]])
    config = DataConfig(str(source))
    with pytest.raises(CompositionError, match="A2: formula found"):
        import_records(config, tmp_path/"rejected.db")
    config.excel_formulas = "cached"
    with pytest.raises(CompositionError, match="A2: formula has no saved result"):
        import_records(config, tmp_path/"missing.db")
    set_formula_cache(source, "30")
    store = import_records(config, tmp_path/"cached.db")
    assert store.record(1)["Balance"] == "30" and store.metadata["warnings"]
    set_formula_cache(source, "#DIV/0!", True)
    with pytest.raises(CompositionError, match="A2: saved formula error"):
        import_records(config, tmp_path/"error.db")
    assert all(not (tmp_path/name).exists() for name in ("rejected.db", "missing.db", "error.db"))


def test_blank_rows_padding_and_non_header_mode(tmp_path):
    path = write_rows(tmp_path/"blank.xlsx", [["Name", "Note"], ["Alice", None], [None, None], ["Bob", "Text"]])
    store = import_records(DataConfig(str(path)), tmp_path/"blank.db")
    assert store.count == 2 and store.metadata["skipped_blank_rows"] == 1
    assert store.record(1) == {"Name": "Alice", "Note": ""}
    source = write_rows(tmp_path/"no-header.xlsx", [["001", "Two"], ["002", "Three"]])
    store = import_records(DataConfig(str(source), header=False), tmp_path/"no-header.db")
    assert store.record(1) == {"Field_1": "001", "Field_2": "Two"}


@pytest.mark.parametrize("case", ["duplicate_header", "empty_header", "no_data", "extra_column", "error_cell", "bad_book", "missing_sheet"])
def test_invalid_workbooks_fail_with_cleanup(tmp_path, case):
    path = tmp_path/"invalid.xlsx"
    rows = {"duplicate_header": [["Name", "Name"], ["a", "b"]],
            "empty_header": [["Name", None, "Code"], ["a", "b", "c"]],
            "no_data": [["Name"]], "extra_column": [["Name"], ["a", "extra"]],
            "error_cell": [["Value"], ["#N/A"]],
            "missing_sheet": [["Name"], ["a"]], "bad_book": [["Name"]]}[case]
    write_rows(path, rows)
    config = DataConfig(str(path), sheet="Missing" if case == "missing_sheet" else "")
    if case == "bad_book":
        path.write_bytes(b"not a workbook")
    with pytest.raises(CompositionError):
        import_records(config, tmp_path/"bad.db")
    assert not (tmp_path/"bad.db").exists()


def test_cancel_close_source_and_header_normalization_collision(tmp_path):
    path = write_rows(tmp_path/"cancel.xlsx", [["Account No", "Account-No"], ["01", "02"]])
    with pytest.raises(CompositionError, match="unique"):
        import_records(DataConfig(str(path)), tmp_path/"bad.db")
    config = DataConfig(str(path), mapping={"Account-No": "Other"})
    with pytest.raises(CompositionError, match="cancelled"):
        import_records(config, tmp_path/"cancel.db", is_cancelled=lambda: True)
    assert not (tmp_path/"cancel.db").exists()
    renamed = path.with_name("released.xlsx")
    path.rename(renamed)
    config.path = str(renamed)
    store = import_records(config, tmp_path/"good.db")
    assert store.record(1) == {"Account_No": "01", "Other": "02"}


def test_excel_data_sequence_pdf_and_schema5_migration(tmp_path):
    source = make_xlsx(tmp_path/"input.xlsx", 3)
    config = DataConfig(str(source), sheet="Statements", header_row=2, mapping={"Customer Name": "Name"})
    store = import_records(config, tmp_path/"records.db")
    template = Template(data=config, sequences=[SequenceSpec("Seq", padding=3)],
                        elements=[Element(value="{{Name}} {{Account}} {{Date}} {{Seq}}",
                                          font=FontSpec(family="Noto Sans CJK HK"), width_mm=170, height_mm=20)])
    saved = save_project(template, tmp_path/"project.pdcx")
    assert load_project(saved).data == config
    result = generate(ProductionJob(template.to_dict(), str(store.path), str(tmp_path/"out")))
    assert result.status == "completed", result.error
    with fitz.open(result.output_pdf) as pdf:
        assert pdf.page_count == 3 and "000003 2026-10-02 003" in pdf[2].get_text()
    log = json.loads((tmp_path/"out"/result.job_id/"job.json").read_text(encoding="utf-8"))
    assert log["source"]["worksheet"] == "Statements"
    assert log["import_configuration"]["header_row"] == 2
    assert "worksheet + import_configuration" in log["record_identity"]
    raw = template.to_dict()
    raw["template_version"] = 5
    for key in ("sheet", "excel_formulas", "preserve_zeros"):
        raw["data"].pop(key)
    before = copy.deepcopy(raw)
    migrated = Template.from_dict(raw)
    assert migrated.template_version == 8 and raw == before and migrated.sequences == template.sequences
    assert migrated.data.sheet == "" and migrated.data.excel_formulas == "reject"


def test_large_xlsx_snapshot(tmp_path):
    path = tmp_path/"large.xlsx"
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Data")
    sheet.append(["Account"])
    for i in range(1000):
        sheet.append([f"{i:08}"])
    book.save(path)
    book.close()
    store = import_records(DataConfig(str(path), sheet="Data"), tmp_path/"large.db")
    assert store.count == 1000 and store.record(1000)["Account"] == "00000999"


@pytest.fixture(scope="session")
def app():
    return QApplication.instance() or QApplication([])


def test_gui_excel_sheet_header_alias_preview_save_reopen_and_production(app, tmp_path):
    source = make_xlsx(tmp_path/"資料.xlsx", 3)
    dialog = DataDialog(str(source), tmp_path)
    try:
        dialog.show()
        wait(lambda: dialog.sheet.count() == 2 and dialog.sample_valid)
        assert not dialog.encoding.isVisible() and dialog.sheet.isVisible()
        dialog.sheet.setCurrentIndex(dialog.sheet.findData("Statements"))
        dialog.start_row.setValue(2)
        wait(lambda: dialog.sample_valid and dialog.mapping.rowCount() == 6)
        dialog.mapping.item(0, 1).setText("Name")
        assert dialog.config().mapping["Customer Name"] == "Name"
        dialog.mapping.item(1, 1).setText("Name")
        assert not dialog.buttons.button(QDialogButtonBox.StandardButton.Ok).isEnabled()
        dialog.mapping.item(1, 1).setText("Account")
        assert dialog.buttons.button(QDialogButtonBox.StandardButton.Ok).isEnabled()
        dialog.zeros.setChecked(False)
        wait(lambda: dialog.sample_valid and dialog.sample.item(0, 1).text() == "1")
        assert dialog.config().mapping["Customer Name"] == "Name"
        dialog.zeros.setChecked(True)
        wait(lambda: dialog.sample_valid and dialog.sample.item(0, 1).text() == "000001")
        config = dialog.config()
    finally:
        dialog.reject()
        wait(lambda: not dialog.workers)
    w = CompositionWindow()
    try:
        w.show()
        w._start_import(config)
        wait(lambda: w.import_worker is None and w.record_count == 3)
        assert "Statements" in w.source_label.text()
        w.add_field("Account", 20, 30)
        assert w.apply_sequences([SequenceSpec("Seq", padding=3)], "imported", 3)
        w.add_element("text", "{{Seq}}", x=20, y=60)
        w.tabs.setCurrentIndex(2)
        w.record.setValue(3)
        wait(lambda: "Record 3" in w.preview_state.text() and not w.workers)
        saved = save_project(w.template, tmp_path/"excel-project.pdcx")
        w.start_production(str(tmp_path/"production"))
        wait(lambda: not w.production_worker and bool(w.last_output))
        with fitz.open(w.last_output) as pdf:
            assert "000003" in pdf[-1].get_text() and "003" in pdf[-1].get_text()
        w.undo.setClean()
        w.open_project(path=str(saved))
        wait(lambda: not w.import_worker and w.record_count == 3)
        assert w.template.data.sheet == "Statements" and w.template.data.mapping["Customer Name"] == "Name"
    finally:
        cleanup(w)




@pytest.mark.parametrize("delimiter", ["\n", "\r", "\0"])
def test_refactored_csv_rejects_control_delimiters(tmp_path, delimiter):
    source = tmp_path/"input.csv"
    source.write_text("Name\nAlice\n", encoding="utf-8")
    with pytest.raises(CompositionError, match="single-character"):
        import_records(DataConfig(str(source), delimiter=delimiter), tmp_path/"bad.db")
    assert not (tmp_path/"bad.db").exists()


def test_refactored_csv_preserves_ordinary_delimiter_and_rejects_nul(tmp_path):
    source = tmp_path/"input.txt"
    source.write_text("NamerCode\nAlicer01\n", encoding="utf-8")
    assert import_records(DataConfig(str(source), delimiter="r"), tmp_path/"good.db").record(1) == {"Name":"Alice", "Code":"01"}
    source.write_text("Name\nAlice\0Oops\n", encoding="utf-8")
    with pytest.raises(CompositionError, match="NUL"):
        import_records(DataConfig(str(source)), tmp_path/"bad.db")
