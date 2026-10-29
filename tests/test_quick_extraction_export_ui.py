"""Trial and one-action Excel UX must retain complete-scan/export safeguards."""
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import QSettings

from tests.test_quick_extraction_ui import wait_until
from tests.test_quick_extraction_workbench_ui import close_dialog


@pytest.fixture(autouse=True)
def settings(qt_application):
    QSettings().remove("dialogs/extract-region-workbench-v3")
    yield
    QSettings().remove("dialogs/extract-region-workbench-v3")


def make_dialog(tmp_path, **kwargs):
    from dialogs.quick_text_extractor import QuickTextExtractorDialog
    from workflow.extraction import ExtractionSpec, Region
    source = tmp_path/"letters.pdf"
    with fitz.open() as pdf:
        for number in range(15):
            page = pdf.new_page()
            page.insert_text((30, 40), f"{number+1:08d}")
        pdf.save(source)
    spec = ExtractionSpec([Region(name="Code", x_mm=0, y_mm=0, width_mm=100, height_mm=100)])
    dialog = QuickTextExtractorDialog(source, 0, str(tmp_path/"out.xlsx"), spec=spec.to_dict(), **kwargs)
    dialog.show()
    wait_until(lambda: dialog.source_ready and not dialog.host.workers)
    return dialog


def chooser(monkeypatch, target, calls):
    def choose(*args):
        calls.append(args)
        return str(target) if target else "", "Excel workbook (*.xlsx)"
    monkeypatch.setattr("dialogs.quick_text_extractor.QFileDialog.getSaveFileName", choose)


def values(path):
    with path.open("rb") as stream:
        book = load_workbook(stream, read_only=True)
        result = [row[3] for row in book["Data"].iter_rows(min_row=2, values_only=True)]
        book.close()
    return result


@pytest.mark.parametrize("trial", [None, "First 10", "All selected"])
def test_primary_exports_complete_selection_without_mandatory_trial(trial, qt_application, tmp_path, monkeypatch):
    dialog = make_dialog(tmp_path)
    calls, target = [], tmp_path/"all.xlsx"
    chooser(monkeypatch, target, calls)
    try:
        assert dialog.sample_mode.currentText() == "First 10"
        if trial:
            dialog.sample_mode.setCurrentText(trial)
            dialog.sample_button.click()
            wait_until(lambda: dialog.result is not None and not dialog.host.workers)
            assert dialog.result["pages"] == (10 if trial == "First 10" else 15)
            assert not dialog.export_button.isEnabled()
            trial_database = dialog.result["database"]
        dialog.extract_export_button.click()
        dialog.extract_export_button.click()  # A second click cannot create another scan/chooser.
        wait_until(lambda: dialog.last_output is not None and not dialog.host.workers)
        assert len(calls) == 1
        assert values(target) == [f"{i:08d}" for i in range(1, 16)]
        assert dialog.result["pages"] == 15
        if trial:
            assert dialog.result["database"] != trial_database
    finally:
        close_dialog(dialog)


def test_primary_reuses_current_complete_result_but_rescans_stale_rules(qt_application, tmp_path, monkeypatch):
    dialog = make_dialog(tmp_path)
    target, calls = tmp_path/"current.xlsx", []
    chooser(monkeypatch, target, calls)
    try:
        next(action for action, _ in dialog.more_actions if action.text() == "Extract only").trigger()
        wait_until(lambda: dialog.result is not None and not dialog.host.workers)
        database = dialog.result["database"]
        dialog.extract_export_button.click()
        wait_until(lambda: dialog.last_output is not None and not dialog.host.workers)
        assert dialog.result["database"] == database
        dialog.name.setText("Changed")
        dialog.save_controls()
        target = tmp_path/"changed.xlsx"
        chooser(monkeypatch, target, calls)
        dialog.extract_export_button.click()
        wait_until(lambda: dialog.last_output == str(target) and not dialog.host.workers)
        assert dialog.result["database"] != database
        assert len(values(target)) == 15
    finally:
        close_dialog(dialog)


def test_chooser_cancel_keeps_complete_results_and_existing_files(qt_application, tmp_path, monkeypatch):
    dialog = make_dialog(tmp_path)
    target, calls = tmp_path/"existing.xlsx", []
    target.write_bytes(b"previous")
    chooser(monkeypatch, None, calls)
    try:
        dialog.extract_export_button.click()
        wait_until(lambda: len(calls) == 1 and not dialog.host.workers)
        assert target.read_bytes() == b"previous" and dialog.last_output is None
        assert dialog.scan_signature is not None and dialog.result["pages"] == 15
    finally:
        close_dialog(dialog)


@pytest.mark.parametrize("action", ["cancel", "close", "source_change"])
def test_cancel_close_or_stale_source_never_advances_to_chooser(action, qt_application, tmp_path, monkeypatch):
    current = [True]
    dialog = make_dialog(tmp_path, source_guard=lambda: current[0])
    calls = []
    chooser(monkeypatch, tmp_path/"unexpected.xlsx", calls)
    try:
        dialog.extract_export_button.click()
        assert dialog.host.active is not None
        if action == "cancel":
            dialog.host.cancel()
        elif action == "close":
            dialog.reject()
        else:
            current[0] = False
            dialog.check_source()
        wait_until(lambda: not dialog.host.workers)
        assert calls == [] and not (tmp_path/"unexpected.xlsx").exists()
    finally:
        close_dialog(dialog)


def test_scan_failure_never_advances_to_chooser(qt_application, tmp_path, monkeypatch):
    dialog = make_dialog(tmp_path)
    calls = []
    chooser(monkeypatch, tmp_path/"unexpected.xlsx", calls)
    try:
        Path(dialog.source).write_bytes(b"Not a PDF")
        dialog.extract_export_button.click()
        wait_until(lambda: not dialog.host.workers)
        assert calls == [] and dialog.last_output is None
        assert not dialog.export_button.isEnabled()
        assert "failed" in dialog.result_state.text().lower()
    finally:
        close_dialog(dialog)


def test_excel_primary_rejects_non_excel_name_without_overwriting(qt_application, tmp_path, monkeypatch):
    dialog = make_dialog(tmp_path)
    target, calls = tmp_path/"existing.csv", []
    target.write_bytes(b"Keep this file")
    chooser(monkeypatch, target, calls)
    try:
        dialog.extract_export_button.click()
        wait_until(lambda: calls and not dialog.host.workers)
        assert target.read_bytes() == b"Keep this file"
        assert dialog.last_output is None and ".xlsx" in dialog.status.text()
    finally:
        close_dialog(dialog)


def test_source_changed_inside_save_chooser_is_checked_again(qt_application, tmp_path, monkeypatch):
    current = [True]
    dialog = make_dialog(tmp_path, source_guard=lambda: current[0])
    target = tmp_path/"old.xlsx"
    target.write_bytes(b"Previous output")
    def changed(*args):
        current[0] = False
        return str(target), "Excel workbook (*.xlsx)"
    monkeypatch.setattr("dialogs.quick_text_extractor.QFileDialog.getSaveFileName", changed)
    try:
        dialog.extract_export_button.click()
        wait_until(lambda: dialog.source_stale and not dialog.host.workers)
        assert target.read_bytes() == b"Previous output" and dialog.last_output is None
    finally:
        close_dialog(dialog)
