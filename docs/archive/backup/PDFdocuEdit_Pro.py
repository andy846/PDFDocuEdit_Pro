import csv
import sys
import json 
import os
import fitz
import shutil
import tempfile
import pandas as pd
import math
import time
import comtypes.client
import pandas as pd
import logging
import chardet
import openpyxl
import io
from datetime import datetime
from PyQt5.QtWidgets import QApplication, QDesktopWidget, QFormLayout, QFrame, QGraphicsDropShadowEffect, QGraphicsScene, QGraphicsView, QHeaderView, QMainWindow, QPushButton, QVBoxLayout, QHBoxLayout, QWidget, QFileDialog, QLabel, QLineEdit, QGridLayout, QMessageBox, QComboBox, QShortcut, QProgressBar, QDialog, QSplashScreen, QScrollArea, QRadioButton, QListWidget, QAbstractItemView, QDockWidget, QMenu, QAction, QToolBar, QToolButton, QInputDialog, QDialogButtonBox, QTabWidget, QGroupBox, QProgressDialog, QDoubleSpinBox, QSpinBox, QCheckBox, QButtonGroup, QTextEdit, QTableWidget, QTableWidgetItem, QFileDialog
from PyQt5.QtGui import QPixmap, QImage, QKeySequence, QIcon, QTextDocument, QPainter, QPen, QPalette, QColor, QFont, QScreen
from PyQt5.QtCore import QRectF, QSize, QTimer, Qt, QThread, pyqtSignal, QRect
from pdf2docx import Converter
from PyQt5.QtPrintSupport import QPrintDialog, QPrinter, QPrinterInfo
from PyQt5.QtWidgets import QSizePolicy
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
import concurrent.futures



def resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

class DeepSearchDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.parent_window = parent
        self.setWindowTitle('PDF 內容深度搜索')
        
        # 按比例調整視窗大小 - 使用屏幕大小的80%
        screen = QDesktopWidget().screenGeometry()
        self.setMinimumSize(int(screen.width() * 0.8), int(screen.height() * 0.8))
        
        # 設置為非模態對話框，允許用戶與其他窗口互動
        self.setModal(False)
        
        # 設置屬性，使對話框關閉時自動釋放記憶體
        self.setAttribute(Qt.WA_DeleteOnClose)
        
        self.search_thread = None
        self.search_results = []
        self.error_results = []
        self.opened_windows = []  # 追蹤從此視窗打開的PDF視窗
        self.init_ui()

    def init_ui(self):
        main_layout = QVBoxLayout(self)
        
        # 文件夾選擇行
        folder_layout = QHBoxLayout()
        folder_label = QLabel('PDF文件夾:')
        folder_label.setStyleSheet("color: white;")
        self.folder_edit = QLineEdit()
        self.folder_edit.setStyleSheet("""
            QLineEdit {
                background-color: #333333;
                color: white;
                border: 1px solid #555555;
                padding: 5px;
                border-radius: 3px;
            }
        """)
        # 移除 folder_edit 的 returnPressed 連接（如果之前有的話）
        browse_button = QPushButton('瀏覽...')
        browse_button.clicked.connect(self.browse_folder)
        folder_layout.addWidget(folder_label)
        folder_layout.addWidget(self.folder_edit)
        folder_layout.addWidget(browse_button)
        main_layout.addLayout(folder_layout)    
        
        # 選項行
        options_layout = QHBoxLayout()
        
        # 子資料夾選項
        self.subfolder_checkbox = QCheckBox('包含子資料夾')
        self.subfolder_checkbox.setChecked(True)
        self.subfolder_checkbox.setStyleSheet("color: white;")
        options_layout.addWidget(self.subfolder_checkbox)
        
        # 開啟方式選項
        self.open_option_label = QLabel("打開方式:")
        self.open_option_label.setStyleSheet("color: white;")
        options_layout.addWidget(self.open_option_label)
        
        self.open_option_combo = QComboBox()
        self.open_option_combo.addItems(["在新視窗打開", "在當前視窗打開", "使用系統默認應用打開"])
        self.open_option_combo.setCurrentIndex(0)  # 預設為新視窗打開
        self.open_option_combo.setStyleSheet("""
            QComboBox {
                background-color: #333333;
                color: white;
                border: 1px solid #555555;
                padding: 5px;
                border-radius: 3px;
            }
            QComboBox::drop-down {
                border: 0px;
            }
            QComboBox::down-arrow {
                image: url(down_arrow.png);
                width: 12px;
                height: 12px;
            }
            QComboBox QAbstractItemView {
                background-color: #333333;
                color: white;
                selection-background-color: #4a86e8;
            }
        """)
        options_layout.addWidget(self.open_option_combo)
        
        main_layout.addLayout(options_layout)
        
        # 搜索文本行
        search_layout = QHBoxLayout()
        search_label = QLabel('搜索文本 (多個關鍵詞請用逗號分隔):')
        search_label.setStyleSheet("color: white;")
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("輸入要搜索的關鍵詞，多個關鍵詞請以英文逗號分隔")
        self.search_edit.setStyleSheet("""
            QLineEdit {
                background-color: #333333;
                color: white;
                border: 1px solid #555555;
                padding: 5px;
                border-radius: 3px;
            }
        """)
        # 添加 ENTER 鍵綁定到搜索功能
        self.search_edit.returnPressed.connect(self.start_search)
        self.search_button = QPushButton('搜索')
        self.search_button.clicked.connect(self.start_search)
        self.cancel_button = QPushButton('取消')
        self.cancel_button.clicked.connect(self.cancel_search)
        self.cancel_button.setEnabled(False)
        search_layout.addWidget(search_label)
        search_layout.addWidget(self.search_edit)
        search_layout.addWidget(self.search_button)
        search_layout.addWidget(self.cancel_button)
        main_layout.addLayout(search_layout)
        
        # 狀態行
        status_layout = QHBoxLayout()
        self.status_label = QLabel('準備就緒')
        self.status_label.setStyleSheet("color: white;")
        
        # 當前處理文件顯示標籤
        self.current_file_label = QLabel("")
        self.current_file_label.setStyleSheet("color: #64B5F6; font-style: italic;")
        
        # 進度條
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        
        status_layout.addWidget(self.status_label, 1)
        status_layout.addWidget(self.current_file_label, 2)
        status_layout.addWidget(self.progress_bar, 1)
        
        main_layout.addLayout(status_layout)
        
        # 結果標籤頁
        self.tab_widget = QTabWidget()
        self.tab_widget.setStyleSheet("""
            QTabWidget {
                background-color: #222222;
            }
            QTabWidget::pane {
                border: 1px solid #444444;
                background-color: #222222;
            }
            QTabBar::tab {
                background-color: #333333;
                color: #cccccc;
                padding: 8px 12px;
                margin-right: 2px;
                border-top-left-radius: 4px;
                border-top-right-radius: 4px;
            }
            QTabBar::tab:selected {
                background-color: #4a86e8;
                color: white;
            }
            QTabBar::tab:hover:!selected {
                background-color: #444444;
            }
        """)
        
        # 搜索結果表格
        self.result_table = QTableWidget()
        self.setup_result_table()
        self.tab_widget.addTab(self.result_table, "搜索結果")
        
        # 內容預覽
        self.preview_text = QTextEdit()
        self.preview_text.setReadOnly(True)
        self.preview_text.setStyleSheet("""
            QTextEdit {
                background-color: #2a2a2a;
                color: white;
                border: none;
            }
        """)
        self.tab_widget.addTab(self.preview_text, "內容預覽")
        
        # 錯誤表格
        self.error_table = QTableWidget()
        self.setup_error_table()
        self.tab_widget.addTab(self.error_table, "錯誤信息")
        
        main_layout.addWidget(self.tab_widget)
        
        # 表格調整提示
        table_hint_label = QLabel("提示: 可拖動列標題邊緣調整欄寬，雙擊列標題邊緣自動調整至最佳寬度")
        table_hint_label.setStyleSheet("color: #999999; font-style: italic;")
        main_layout.addWidget(table_hint_label)
        
        # 操作按鈕行
        button_layout = QHBoxLayout()
        
        self.clear_all_button = QPushButton('清除所有結果')
        self.clear_all_button.clicked.connect(self.clear_all_results)
        
        self.export_results_button = QPushButton('導出搜索結果')
        self.export_results_button.clicked.connect(self.export_search_results)
        self.export_results_button.setEnabled(False)
        
        self.open_selected_button = QPushButton('打開選中文件')
        self.open_selected_button.clicked.connect(self.open_selected_file)
        self.open_selected_button.setEnabled(False)
        
        self.resize_columns_button = QPushButton('自動調整欄寬')
        self.resize_columns_button.clicked.connect(self.auto_resize_columns)
        
        button_layout.addWidget(self.clear_all_button)
        button_layout.addWidget(self.export_results_button)
        button_layout.addWidget(self.open_selected_button)
        button_layout.addWidget(self.resize_columns_button)
        
        main_layout.addLayout(button_layout)
        
        # 連接表格選擇事件
        self.result_table.itemSelectionChanged.connect(self.update_preview)
        self.result_table.itemSelectionChanged.connect(self.update_button_states)
        
        # 連接表格雙擊事件
        self.result_table.cellDoubleClicked.connect(self.handle_table_double_click)
        
        # 設置按鈕樣式
        self.set_button_styles()
        
        # 設置對話框整體深色背景
        self.setStyleSheet("QDialog { background-color: #222222; }")

    def set_button_styles(self):
        """設置按鈕樣式以提升界面美觀度"""
        button_style = """
            QPushButton {
                background-color: #4a86e8;
                color: white;
                border: none;
                padding: 8px 16px;
                border-radius: 4px;
                min-width: 100px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #3a76d8;
            }
            QPushButton:pressed {
                background-color: #2a66c8;
            }
            QPushButton:disabled {
                background-color: #555555;
                color: #999999;
            }
        """
        
        self.search_button.setStyleSheet(button_style)
        self.cancel_button.setStyleSheet(button_style)
        self.clear_all_button.setStyleSheet(button_style)
        self.export_results_button.setStyleSheet(button_style)
        self.open_selected_button.setStyleSheet(button_style)
        self.resize_columns_button.setStyleSheet(button_style)
        
        # 設置進度條樣式
        progress_style = """
            QProgressBar {
                border: 1px solid #555555;
                border-radius: 3px;
                text-align: center;
                color: white;
                background-color: #333333;
            }
            QProgressBar::chunk {
                background-color: #4a86e8;
                width: 10px;
            }
        """
        
        self.progress_bar.setStyleSheet(progress_style)

    def setup_result_table(self):
        # 設置結果表格
        self.result_table.setColumnCount(4)
        self.result_table.setHorizontalHeaderLabels(["文件名", "頁碼", "匹配數", "上下文摘要"])
        
        # 設置初始列寬 - 但允許用戶調整
        self.result_table.setColumnWidth(0, 250)   # 文件名列
        self.result_table.setColumnWidth(1, 100)   # 頁碼列
        self.result_table.setColumnWidth(2, 80)    # 匹配數列
        self.result_table.setColumnWidth(3, 450)   # 上下文摘要列
        
        # 設置列頭可交互調整大小
        header = self.result_table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Interactive)
        
        # 可以雙擊列分隔線自動調整
        header.setDefaultSectionSize(100)
        header.setMinimumSectionSize(50)
        header.setSectionsClickable(True)
        
        # 允許排序
        self.result_table.setSortingEnabled(True)
        
        # 設置表格樣式
        self.result_table.setSelectionBehavior(QTableWidget.SelectRows)  # 選中整行
        self.result_table.setSelectionMode(QTableWidget.SingleSelection)  # 單選
        self.result_table.setEditTriggers(QTableWidget.NoEditTriggers)  # 不可編輯
        self.result_table.setAlternatingRowColors(True)  # 交替行顏色
        
        # 設置表格深色模式樣式
        self.result_table.setStyleSheet("""
            QTableWidget {
                background-color: #2a2a2a;
                color: white;
                gridline-color: #444444;
                border: none;
                selection-background-color: #4a86e8;
                selection-color: white;
                alternate-background-color: #333333;
            }
            QHeaderView::section {
                background-color: #444444;
                color: white;
                padding: 6px;
                border: 1px solid #555555;
                font-weight: bold;
            }
            QTableWidget::item {
                border-bottom: 1px solid #444444;
                padding: 4px;
            }
            QTableWidget::item:selected {
                background-color: #3a76d8;
                color: white;
            }
            QScrollBar:vertical {
                background-color: #2a2a2a;
                width: 12px;
                margin: 0px;
            }
            QScrollBar::handle:vertical {
                background-color: #666666;
                min-height: 20px;
                border-radius: 6px;
            }
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {
                height: 0px;
            }
            QScrollBar:horizontal {
                background-color: #2a2a2a;
                height: 12px;
                margin: 0px;
            }
            QScrollBar::handle:horizontal {
                background-color: #666666;
                min-width: 20px;
                border-radius: 6px;
            }
            QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal {
                width: 0px;
            }
        """)
        
        # 添加頭部右鍵菜單
        header.setContextMenuPolicy(Qt.CustomContextMenu)
        header.customContextMenuRequested.connect(self.show_header_menu)
    
    def setup_error_table(self):
        # 設置錯誤表格
        self.error_table.setColumnCount(2)
        self.error_table.setHorizontalHeaderLabels(["文件名", "錯誤信息"])
        
        # 設置初始列寬 - 但允許用戶調整
        self.error_table.setColumnWidth(0, 250)  # 文件名列
        self.error_table.setColumnWidth(1, 650)  # 錯誤信息列
        
        # 設置列頭可交互調整大小
        header = self.error_table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Interactive)
        
        # 可以雙擊列分隔線自動調整
        header.setDefaultSectionSize(100)
        header.setMinimumSectionSize(50)
        
        # 允許排序
        self.error_table.setSortingEnabled(True)
        
        # 設置表格樣式
        self.error_table.setSelectionBehavior(QTableWidget.SelectRows)  # 選中整行
        self.error_table.setAlternatingRowColors(True)  # 交替行顏色
        
        # 應用與結果表格相同的深色模式樣式
        self.error_table.setStyleSheet(self.result_table.styleSheet())
        
        # 添加頭部右鍵菜單
        header = self.error_table.horizontalHeader()
        header.setContextMenuPolicy(Qt.CustomContextMenu)
        header.customContextMenuRequested.connect(self.show_error_header_menu)
    
    def show_header_menu(self, pos):
        """為結果表格表頭顯示右鍵菜單"""
        menu = QMenu(self)
        menu.setStyleSheet("""
            QMenu {
                background-color: #333333;
                color: white;
                border: 1px solid #555555;
            }
            QMenu::item {
                padding: 6px 20px;
            }
            QMenu::item:selected {
                background-color: #4a86e8;
                color: white;
            }
        """)
        resize_action = menu.addAction("自動調整此欄寬度")
        resize_all_action = menu.addAction("自動調整所有欄寬度")
        
        action = menu.exec_(self.result_table.horizontalHeader().mapToGlobal(pos))
        
        if action == resize_action:
            column = self.result_table.horizontalHeader().logicalIndexAt(pos)
            self.result_table.resizeColumnToContents(column)
        elif action == resize_all_action:
            self.auto_resize_columns()
    
    def show_error_header_menu(self, pos):
        """為錯誤表格表頭顯示右鍵菜單"""
        menu = QMenu(self)
        menu.setStyleSheet("""
            QMenu {
                background-color: #333333;
                color: white;
                border: 1px solid #555555;
            }
            QMenu::item {
                padding: 6px 20px;
            }
            QMenu::item:selected {
                background-color: #4a86e8;
                color: white;
            }
        """)
        resize_action = menu.addAction("自動調整此欄寬度")
        resize_all_action = menu.addAction("自動調整所有欄寬度")
        
        action = menu.exec_(self.error_table.horizontalHeader().mapToGlobal(pos))
        
        if action == resize_action:
            column = self.error_table.horizontalHeader().logicalIndexAt(pos)
            self.error_table.resizeColumnToContents(column)
        elif action == resize_all_action:
            for col in range(self.error_table.columnCount()):
                self.error_table.resizeColumnToContents(col)
    
    def auto_resize_columns(self):
        """自動調整所有列寬以適應內容"""
        for col in range(self.result_table.columnCount()):
            self.result_table.resizeColumnToContents(col)
            
            # 確保列寬不會太窄，至少有最小寬度
            min_width = 80
            if col == 0:  # 文件名列
                min_width = 200
            elif col == 3:  # 上下文摘要列
                min_width = 300
            
            current_width = self.result_table.columnWidth(col)
            if current_width < min_width:
                self.result_table.setColumnWidth(col, min_width)


    def keyPressEvent(self, event):
        # 如果按下Enter鍵
        if event.key() == Qt.Key_Return or event.key() == Qt.Key_Enter:
            # 檢查哪個輸入框有焦點
            focused_widget = QApplication.focusWidget()
            if focused_widget == self.search_edit:
                # 如果搜索文本框有焦點，則啟動搜索
                self.start_search()
                event.accept()
            elif focused_widget == self.folder_edit:
                # 如果文件夾文本框有焦點，則什麼都不做
                event.accept()
            else:
                # 其他情況，交給父類處理
                super().keyPressEvent(event)
        else:
            # 非Enter鍵，交給父類處理
            super().keyPressEvent(event)
    
    def browse_folder(self):
        folder = QFileDialog.getExistingDirectory(self, '選擇PDF文件夾')
        if folder:
            self.folder_edit.setText(folder)
    
    def start_search(self):
        folder_path = self.folder_edit.text()
        search_text = self.search_edit.text()
        
        if not folder_path:
            self.status_label.setText("請提供文件夾路徑")
            return
            
        if not os.path.isdir(folder_path):
            self.status_label.setText("指定的文件夾不存在")
            return
        
        if not search_text:
            self.status_label.setText("請提供搜索文本")
            return
        
        # 處理多個關鍵詞，分割並去除空白
        keywords = [kw.strip() for kw in search_text.split(',') if kw.strip()]
        
        if not keywords:
            self.status_label.setText("請提供有效的搜索關鍵詞")
            return
        
        self.clear_all_results()
        self.progress_bar.setValue(0)
        self.search_button.setEnabled(False)
        self.cancel_button.setEnabled(True)
        self.search_results = []
        self.error_results = []
        self.export_results_button.setEnabled(False)
        
        # 創建並啟動搜索線程
        include_subfolders = self.subfolder_checkbox.isChecked()
        self.search_thread = PDFSearcher(folder_path, keywords, include_subfolders)
        self.search_thread.progress_signal.connect(self.update_progress)
        self.search_thread.result_signal.connect(self.add_result)
        self.search_thread.error_signal.connect(self.add_error)
        self.search_thread.status_signal.connect(self.update_status)
        self.search_thread.current_file_signal.connect(self.update_current_file)
        self.search_thread.finished_signal.connect(self.search_finished)
        self.search_thread.start()
        
        # 切換到結果標籤頁
        self.tab_widget.setCurrentIndex(0)
    
    def cancel_search(self):
        if self.search_thread and self.search_thread.isRunning():
            self.search_thread.stop()
            self.status_label.setText("正在取消搜索...")
    
    def update_progress(self, value):
        self.progress_bar.setValue(value)
    
    def update_status(self, status):
        self.status_label.setText(status)
    
    def update_current_file(self, file_name):
        self.current_file_label.setText(f"處理中: {file_name}")
    
    def add_result(self, result):
        self.search_results.append(result)
        
        row = self.result_table.rowCount()
        self.result_table.insertRow(row)
        
        # 文件名 - 添加可點擊效果
        file_name_item = QTableWidgetItem(result["file_name"])
        # 設置工具提示，顯示完整路徑
        file_name_item.setToolTip(f"完整路徑: {result['file_path']}\n點擊打開文件")
        # 設置文本顏色為藍色，模擬超連結
        file_name_item.setForeground(QColor(102, 187, 255))  # 使用較亮的藍色，適合深色背景
        # 儲存文件路徑作為用戶數據
        file_name_item.setData(Qt.UserRole, result["file_path"])
        self.result_table.setItem(row, 0, file_name_item)
        
        # 頁碼列表，用逗號分隔
        pages_str = ", ".join(str(p) for p in result["pages"])
        page_item = QTableWidgetItem(pages_str)
        page_item.setToolTip(f"匹配頁碼: {pages_str}")
        self.result_table.setItem(row, 1, page_item)
        
        # 添加匹配數量列
        match_count_item = QTableWidgetItem(str(len(result["pages"])))
        match_count_item.setTextAlignment(Qt.AlignCenter)
        self.result_table.setItem(row, 2, match_count_item)
        
        # 第一個匹配片段的預覽
        if result["snippets"] and len(result["snippets"]) > 0:
            preview = result["snippets"][0]
            # 如果片段太長，截斷它
            if len(preview) > 100:
                preview = preview[:100] + "..."
            context_item = QTableWidgetItem(preview)
            context_item.setToolTip(preview)  # 設置工具提示顯示完整內容
            self.result_table.setItem(row, 3, context_item)
        else:
            self.result_table.setItem(row, 3, QTableWidgetItem("無內容預覽"))
    
    def add_error(self, error):
        self.error_results.append(error)
        
        row = self.error_table.rowCount()
        self.error_table.insertRow(row)
        
        # 文件名
        file_name_item = QTableWidgetItem(error["file_name"])
        file_name_item.setToolTip(f"完整路徑: {error['file_path']}")
        self.error_table.setItem(row, 0, file_name_item)
        
        # 錯誤信息
        error_item = QTableWidgetItem(error["error"])
        error_item.setToolTip(error["error"])
        self.error_table.setItem(row, 1, error_item)
        
        # 切換到錯誤標籤頁以提示用戶
        self.tab_widget.setCurrentIndex(2)
    
    def search_finished(self, summary):
        self.search_button.setEnabled(True)
        self.cancel_button.setEnabled(False)
        self.search_thread = None
        
        # 清除當前處理文件顯示
        self.current_file_label.setText("")
        
        # 更新狀態
        status_msg = (f"搜索完成! 耗時: {summary['total_time']:.2f} 秒, "
                      f"處理: {summary['total_files']} 個文件, "
                      f"找到匹配: {summary['match_files']} 個, "
                      f"錯誤: {summary['error_files']} 個")
        self.status_label.setText(status_msg)
        
        # 如果有結果，啟用導出按鈕
        if summary['match_files'] > 0:
            self.export_results_button.setEnabled(True)
            
            # 搜索完成後自動調整列寬
            QTimer.singleShot(100, self.auto_resize_columns)
        
        # 自動切換到結果標籤頁，如果有錯誤，則選擇錯誤標籤頁
        if summary['match_files'] > 0:
            self.tab_widget.setCurrentIndex(0)
        elif summary['error_files'] > 0:
            self.tab_widget.setCurrentIndex(2)
            # 自動調整錯誤表格列寬
            for col in range(self.error_table.columnCount()):
                self.error_table.resizeColumnToContents(col)
    
    def update_preview(self):
        # 獲取選定的行
        selected_items = self.result_table.selectedItems()
        if not selected_items:
            return
        
        # 獲取所選行
        row = selected_items[0].row()
        
        # 更新按鈕狀態
        self.open_selected_button.setEnabled(True)
        
        # 獲取對應的結果
        if row < len(self.search_results):
            result = self.search_results[row]
            
            # 構建預覽內容 - 適配深色模式
            preview_text = f"<div style='color: white;'>"
            preview_text += f"<h3 style='color: #4a86e8;'>文件: {result['file_name']}</h3>"
            preview_text += f"<p><b>路徑:</b> {result['file_path']}</p>"
            preview_text += f"<p><b>匹配頁碼:</b> {', '.join(str(p) for p in result['pages'])}</p>"
            preview_text += f"<p><b>匹配數量:</b> {len(result['pages'])}</p>"
            preview_text += "<h4>匹配內容:</h4><hr style='border: 1px solid #444444;'>"
            
            # 添加每個匹配片段，突出顯示關鍵詞
            for i, snippet in enumerate(result["snippets"]):
                # HTML格式高亮顯示關鍵詞
                highlighted_snippet = snippet.replace(
                    f"[{self.search_edit.text()}]", 
                    f"<span style='background-color: #ffd700; color: black; font-weight: bold;'>{self.search_edit.text()}</span>"
                )
                preview_text += f"<p><b style='color: #4a86e8;'>第 {result['pages'][i]} 頁:</b><br>{highlighted_snippet}</p>"
                if i < len(result["snippets"]) - 1:
                    preview_text += "<hr style='border: 1px solid #444444;'>"
            
            preview_text += "</div>"
            
            # 設置預覽文本 (使用HTML格式)
            self.preview_text.setHtml(preview_text)
            
            # 不再自動切換到預覽標籤頁
            # 移除: self.tab_widget.setCurrentIndex(1)
    
    def update_button_states(self):
        # 根據當前選擇更新按鈕狀態
        has_selection = len(self.result_table.selectedItems()) > 0
        self.open_selected_button.setEnabled(has_selection)
    
    def handle_table_double_click(self, row, column):
        # 只處理文件名列 (第0列) 的雙擊事件
        if column == 0:
            self.open_selected_file()
    
    def open_selected_file(self):
        selected_items = self.result_table.selectedItems()
        if not selected_items:
            return
        
        row = selected_items[0].row()
        file_path = self.result_table.item(row, 0).data(Qt.UserRole)
        
        # 获取当前选择的打开方式
        open_method = self.open_option_combo.currentIndex()
        
        if open_method == 0:  # 在新視窗打開
            success = self.open_in_new_window(file_path)
            if not success:
                # 如果无法使用新窗口打开，退回到系统默认方式
                self.open_with_system_default(file_path)
        elif open_method == 1:  # 在當前視窗打開
            if self.parent_window and hasattr(self.parent_window, "load_pdf"):
                self.parent_window.load_pdf(file_path)
                # 注意：不再关闭对话框，保留搜索结果
                self.hide()  # 暂时隐藏但不关闭
                self.parent_window.activateWindow()  # 激活父窗口
        else:  # 使用系統默認應用打開
            self.open_with_system_default(file_path)
    
    def open_in_new_window(self, file_path):
        """在新窗口中打开PDF文件"""
        if self.parent_window and hasattr(self.parent_window, "__class__"):
            try:
                # 创建新的PDF查看器窗口
                new_viewer = self.parent_window.__class__()
                new_viewer.show()
                
                # 确保新窗口显示在前台
                new_viewer.raise_()
                new_viewer.activateWindow()
                
                # 异步加载PDF到新窗口，避免在当前方法中阻塞
                QTimer.singleShot(100, lambda: self.load_pdf_to_viewer(new_viewer, file_path))
                
                # 记录打开的窗口
                self.opened_windows.append(new_viewer)
                
                # 如果父窗口有跟踪打开窗口的列表，添加新窗口到列表中
                if hasattr(self.parent_window, "open_windows"):
                    self.parent_window.open_windows.append(new_viewer)
                
                return True
            except Exception as e:
                print(f"在新窗口打开PDF时出错: {str(e)}")
                return False
        else:
            return False
    
    def load_pdf_to_viewer(self, viewer, file_path):
        """异步加载PDF到指定的查看器窗口"""
        try:
            if hasattr(viewer, "load_pdf"):
                viewer.load_pdf(file_path)
            else:
                print("查看器窗口不支持load_pdf方法")
        except Exception as e:
            print(f"加载PDF时出错: {str(e)}")
    
    def open_with_system_default(self, file_path):
        """使用系统默认应用打开PDF文件"""
        try:
            import platform
            system = platform.system()
            if system == "Windows":
                os.startfile(file_path)
            elif system == "Darwin":  # macOS
                import subprocess
                subprocess.call(["open", file_path])
            else:  # Linux and others
                import subprocess
                subprocess.call(["xdg-open", file_path])
        except Exception as e:
            QMessageBox.warning(self, "打開文件失敗", f"無法打開文件: {str(e)}")
    
    def export_search_results(self):
        if not self.search_results:
            QMessageBox.information(self, "導出結果", "沒有可導出的搜索結果")
            return
        
        file_path, _ = QFileDialog.getSaveFileName(
            self, 
            "導出搜索結果", 
            f"PDF搜索結果_{self.search_edit.text()}.html", 
            "HTML 文件 (*.html);;文本文件 (*.txt);;CSV 文件 (*.csv)"
        )
        
        if not file_path:
            return
        
        try:
            if file_path.endswith(".html"):
                self.export_as_html(file_path)
            elif file_path.endswith(".txt"):
                self.export_as_text(file_path)
            elif file_path.endswith(".csv"):
                self.export_as_csv(file_path)
            
            QMessageBox.information(self, "導出成功", f"搜索結果已導出到:\n{file_path}")
        except Exception as e:
            QMessageBox.warning(self, "導出失敗", f"導出結果時發生錯誤: {str(e)}")
    
    def export_as_html(self, file_path):
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write("""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>PDF 搜索結果</title>
                <style>
                    body { font-family: Arial, sans-serif; margin: 20px; }
                    h1 { color: #4a86e8; }
                    .search-info { background-color: #f2f2f2; padding: 10px; border-radius: 5px; }
                    .result { margin: 15px 0; padding: 15px; border: 1px solid #ddd; border-radius: 5px; }
                    .file-name { font-weight: bold; color: #4a86e8; }
                    .file-path { color: #666; font-size: 0.9em; }
                    .pages { margin: 5px 0; }
                    .snippet { background-color: #f9f9f9; padding: 10px; margin: 5px 0; border-left: 3px solid #4a86e8; }
                    .highlight { background-color: yellow; font-weight: bold; }
                </style>
            </head>
            <body>
                <h1>PDF 搜索結果</h1>
                <div class="search-info">
                    <p><strong>搜索關鍵詞:</strong> %s</p>
                    <p><strong>搜索路徑:</strong> %s</p>
                    <p><strong>匹配文件數:</strong> %d</p>
                    <p><strong>搜索時間:</strong> %s</p>
                </div>
            """ % (
                self.search_edit.text(),
                self.folder_edit.text(),
                len(self.search_results),
                time.strftime("%Y-%m-%d %H:%M:%S")
            ))
            
            # 寫入每個搜索結果
            for i, result in enumerate(self.search_results):
                f.write('<div class="result">\n')
                f.write(f'<div class="file-name">文件 {i+1}: {result["file_name"]}</div>\n')
                f.write(f'<div class="file-path">路徑: {result["file_path"]}</div>\n')
                f.write(f'<div class="pages">匹配頁碼: {", ".join(str(p) for p in result["pages"])}</div>\n')
                
                # 寫入每個片段
                for j, snippet in enumerate(result["snippets"]):
                    # 用HTML高亮標記替換關鍵詞
                    highlighted_snippet = snippet.replace(
                        f"[{self.search_edit.text()}]", 
                        f'<span class="highlight">{self.search_edit.text()}</span>'
                    )
                    f.write(f'<div class="snippet">第 {result["pages"][j]} 頁: {highlighted_snippet}</div>\n')
                
                f.write('</div>\n')
            
            f.write("</body>\n</html>")
    
    def export_as_text(self, file_path):
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(f"PDF 搜索結果\n")
            f.write(f"搜索關鍵詞: {self.search_edit.text()}\n")
            f.write(f"搜索路徑: {self.folder_edit.text()}\n")
            f.write(f"匹配文件數: {len(self.search_results)}\n")
            f.write(f"搜索時間: {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
            f.write("-" * 80 + "\n\n")
            
            for i, result in enumerate(self.search_results):
                f.write(f"文件 {i+1}: {result['file_name']}\n")
                f.write(f"路徑: {result['file_path']}\n")
                f.write(f"匹配頁碼: {', '.join(str(p) for p in result['pages'])}\n\n")
                
                for j, snippet in enumerate(result["snippets"]):
                    f.write(f"第 {result['pages'][j]} 頁: {snippet}\n\n")
                
                f.write("-" * 80 + "\n\n")
    
    def export_as_csv(self, file_path):
        import csv
        with open(file_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(["文件名", "文件路徑", "頁碼", "匹配內容"])
            
            for result in self.search_results:
                for j, snippet in enumerate(result["snippets"]):
                    writer.writerow([
                        result["file_name"],
                        result["file_path"],
                        result["pages"][j],
                        snippet
                    ])
    
    def clear_all_results(self):
        # 清除所有結果
        self.result_table.setRowCount(0)
        self.error_table.setRowCount(0)
        self.preview_text.clear()
        self.progress_bar.setValue(0)
        self.status_label.setText("已清除所有結果")
        self.current_file_label.setText("")  # 清除當前處理文件顯示
        self.search_results = []
        self.error_results = []
        self.export_results_button.setEnabled(False)
        self.open_selected_button.setEnabled(False)
    
    def closeEvent(self, event):
        """對話框關閉時清理資源"""
        # 停止搜索线程（如果正在运行）
        if self.search_thread and self.search_thread.isRunning():
            self.search_thread.stop()
            self.search_thread.wait()
        
        # 接受关闭事件
        event.accept()


class PDFSearcher(QThread):
    progress_signal = pyqtSignal(int)
    result_signal = pyqtSignal(dict)
    error_signal = pyqtSignal(dict)
    status_signal = pyqtSignal(str)
    finished_signal = pyqtSignal(dict)
    current_file_signal = pyqtSignal(str)
    
    def __init__(self, folder_path, keywords, include_subfolders):
        super().__init__()
        self.folder_path = folder_path
        self.keywords = keywords  # 現在是關鍵詞列表
        self.include_subfolders = include_subfolders
        self.is_running = True
        
    def run(self):
        start_time = time.time()
        
        # 獲取所有PDF文件
        pdf_files = []
        if self.include_subfolders:
            for root, _, files in os.walk(self.folder_path):
                for file in files:
                    if file.lower().endswith('.pdf'):
                        pdf_files.append(os.path.join(root, file))
        else:
            pdf_files = [os.path.join(self.folder_path, f) for f in os.listdir(self.folder_path) 
                        if f.lower().endswith('.pdf') and os.path.isfile(os.path.join(self.folder_path, f))]
        
        total_files = len(pdf_files)
        if total_files == 0:
            self.status_signal.emit("未找到PDF文件")
            return
            
        self.status_signal.emit(f"找到 {total_files} 個PDF文件，開始搜索...")
        
        # 處理文件
        files_processed = 0
        files_with_errors = 0
        files_with_matches = 0
        
        for pdf_file in pdf_files:
            if not self.is_running:
                self.status_signal.emit("搜索已被取消")
                return
            
            # 發送當前正在處理的檔案名稱
            current_file_name = os.path.basename(pdf_file)
            self.current_file_signal.emit(current_file_name)
            self.status_signal.emit(f"正在搜索: {current_file_name} ({files_processed+1}/{total_files})")
            
            # 對每個關鍵詞進行搜索，合併結果
            all_results = {}
            
            for keyword in self.keywords:
                # 使用搜索函數
                result = search_pdf_with_fitz(pdf_file, keyword)
                
                if result and "error" not in result:
                    # 將結果合併到all_results中
                    if pdf_file not in all_results:
                        all_results[pdf_file] = {
                            "file_path": result["file_path"],
                            "file_name": result["file_name"],
                            "pages": [],
                            "snippets": [],
                            "keywords": []
                        }
                    
                    # 添加頁碼和片段，同時記錄對應的關鍵詞
                    for i, page_num in enumerate(result["pages"]):
                        all_results[pdf_file]["pages"].append(page_num)
                        all_results[pdf_file]["snippets"].append(result["snippets"][i])
                        all_results[pdf_file]["keywords"].append(keyword)
                elif result and "error" in result:
                    files_with_errors += 1
                    self.error_signal.emit(result)
                    # 如果發生錯誤，跳過該文件的其他關鍵詞搜索
                    all_results.pop(pdf_file, None)
                    break
            
            # 如果找到任何匹配，發送結果
            if pdf_file in all_results and all_results[pdf_file]["pages"]:
                files_with_matches += 1
                self.result_signal.emit(all_results[pdf_file])
            
            files_processed += 1
            progress = int(files_processed / total_files * 100)
            self.progress_signal.emit(progress)
        
        elapsed_time = time.time() - start_time
        summary = {
            "total_time": elapsed_time,
            "total_files": total_files,
            "match_files": files_with_matches,
            "error_files": files_with_errors
        }
        self.finished_signal.emit(summary)
    
    def stop(self):
        self.is_running = False


def search_pdf_with_fitz(pdf_path, search_text):
    """使用PyMuPDF搜索PDF文件中的文本"""
    try:
        # 打開PDF文件
        doc = fitz.open(pdf_path)
        found_pages = []
        found_snippets = []
        
        # 遍歷每一頁
        for page_num in range(len(doc)):
            try:
                # 獲取頁面
                page = doc[page_num]
                
                # 提取文本
                text = page.get_text()
                
                # 搜索文本（不區分大小寫）
                if text and search_text.lower() in text.lower():
                    found_pages.append(page_num + 1)
                    
                    # 提取包含關鍵字的上下文
                    text_lower = text.lower()
                    search_text_lower = search_text.lower()
                    start_pos = text_lower.find(search_text_lower)
                    
                    # 提取關鍵字前後各50個字符作為上下文
                    context_start = max(0, start_pos - 50)
                    context_end = min(len(text), start_pos + len(search_text) + 50)
                    
                    # 如果上下文開始不是句子開頭，尋找前一個空格
                    if context_start > 0:
                        space_before = text.rfind(" ", 0, context_start)
                        if space_before != -1:
                            context_start = space_before + 1
                    
                    # 如果上下文結束不是句子結尾，尋找下一個空格
                    if context_end < len(text):
                        space_after = text.find(" ", context_end)
                        if space_after != -1:
                            context_end = space_after
                    
                    snippet = text[context_start:context_end].strip()
                    # 將關鍵詞用 [...] 標記 (簡單替換，實際UI中會用顏色標記)
                    snippet_highlight = snippet.replace(search_text, f"[{search_text}]")
                    found_snippets.append(snippet_highlight)
                    
            except Exception as e:
                print(f"處理頁面 {page_num+1} 時出錯: {str(e)}")
                continue
        
        # 關閉文檔
        doc.close()
        
        # 返回結果
        if found_pages:
            return {
                "file_path": pdf_path,
                "file_name": os.path.basename(pdf_path),
                "pages": found_pages,
                "snippets": found_snippets
            }
        else:
            return None
            
    except Exception as e:
        error_msg = f"錯誤處理文件 {os.path.basename(pdf_path)}: {str(e)}"
        print(error_msg)
        return {
            "file_path": pdf_path,
            "file_name": os.path.basename(pdf_path),
            "error": error_msg
        }
    
class MergeCSVExcelDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("CSV and Excel File Merger")
        self.setMinimumSize(600, 400)
        self.setup_ui()

    def setup_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(20)

        # Input and Output section
        io_group = QGroupBox("Input and Output")
        io_layout = QFormLayout()

        self.input_dir_edit = QLineEdit()
        input_dir_button = QPushButton("Browse")
        input_dir_button.clicked.connect(self.select_input_directory)
        input_dir_layout = QHBoxLayout()
        input_dir_layout.addWidget(self.input_dir_edit)
        input_dir_layout.addWidget(input_dir_button)
        io_layout.addRow("Input Directory:", input_dir_layout)

        self.output_file_edit = QLineEdit()
        output_file_button = QPushButton("Browse")
        output_file_button.clicked.connect(self.select_output_file)
        output_file_layout = QHBoxLayout()
        output_file_layout.addWidget(self.output_file_edit)
        output_file_layout.addWidget(output_file_button)
        io_layout.addRow("Output File:", output_file_layout)

        io_group.setLayout(io_layout)
        main_layout.addWidget(io_group)

        # Options section
        options_group = QGroupBox("Merge Options")
        options_layout = QFormLayout()

        self.skip_rows_edit = QLineEdit()
        self.skip_rows_edit.setText("2")
        options_layout.addRow("Skip Rows:", self.skip_rows_edit)

        self.exclude_keywords_edit = QLineEdit()
        self.exclude_keywords_edit.setPlaceholderText("e.g. Total,Sum")
        options_layout.addRow("Exclude Rows (comma-separated keywords):", self.exclude_keywords_edit)

        check_range_layout = QHBoxLayout()
        self.check_first_cell_radio = QRadioButton("Check First Cell Only")
        self.check_entire_row_radio = QRadioButton("Check Entire Row")
        self.check_first_cell_radio.setChecked(True)
        check_range_group = QButtonGroup(self)
        check_range_group.addButton(self.check_first_cell_radio)
        check_range_group.addButton(self.check_entire_row_radio)
        check_range_layout.addWidget(self.check_first_cell_radio)
        check_range_layout.addWidget(self.check_entire_row_radio)
        options_layout.addRow("Exclusion Check:", check_range_layout)

        # Add encoding options
        self.encoding_combo = QComboBox()
        self.encoding_combo.addItems(['auto', 'utf-8', 'big5', 'iso-8859-1', 'windows-1252', 'ascii'])
        options_layout.addRow("CSV Encoding:", self.encoding_combo)

        options_group.setLayout(options_layout)
        main_layout.addWidget(options_group)

        # File format section
        format_group = QGroupBox("File Formats")
        format_layout = QHBoxLayout()
        format_layout.setAlignment(Qt.AlignLeft)
        self.csv_checkbox = QCheckBox(".csv")
        self.csv_checkbox.setChecked(True)
        self.xls_checkbox = QCheckBox(".xls")
        self.xlsx_checkbox = QCheckBox(".xlsx")
        format_layout.addWidget(self.csv_checkbox)
        format_layout.addWidget(self.xls_checkbox)
        format_layout.addWidget(self.xlsx_checkbox)
        format_group.setLayout(format_layout)
        main_layout.addWidget(format_group)

        # Merge button
        self.merge_button = QPushButton("Merge Files")
        self.merge_button.clicked.connect(self.run_merge)
        self.merge_button.setStyleSheet("font-weight: bold; padding: 10px;")
        main_layout.addWidget(self.merge_button)

    def select_input_directory(self):
        directory = QFileDialog.getExistingDirectory(self, "Select Input Directory")
        if directory:
            self.input_dir_edit.setText(directory)

    def select_output_file(self):
        file_name, _ = QFileDialog.getSaveFileName(self, "Save Output File", "", "Excel files (*.xlsx);;CSV files (*.csv)")
        if file_name:
            self.output_file_edit.setText(file_name)

    def run_merge(self):
        input_dir = self.input_dir_edit.text()
        output_file = self.output_file_edit.text()
        
        if not input_dir or not output_file:
            QMessageBox.warning(self, "Input Error", "Please specify both input directory and output file.")
            return

        try:
            skip_rows = int(self.skip_rows_edit.text())
        except ValueError:
            QMessageBox.warning(self, "Input Error", "Invalid value for 'Skip Rows'. Please enter a number.")
            return

        exclude_keywords = [keyword.strip() for keyword in self.exclude_keywords_edit.text().split(',') if keyword.strip()]
        check_first_cell_only = self.check_first_cell_radio.isChecked()
        file_formats = []
        if self.csv_checkbox.isChecked():
            file_formats.append('.csv')
        if self.xls_checkbox.isChecked():
            file_formats.append('.xls')
        if self.xlsx_checkbox.isChecked():
            file_formats.append('.xlsx')

        if not file_formats:
            QMessageBox.warning(self, "Input Error", "Please select at least one file format.")
            return

        try:
            self.merge_files(input_dir, output_file, skip_rows, exclude_keywords, check_first_cell_only, file_formats)
            QMessageBox.information(self, "Success", f"Files merged successfully!\nOutput file: {output_file}")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"An error occurred during merging:\n{str(e)}")

    def merge_files(self, input_dir, output_file, skip_rows, exclude_keywords, check_first_cell_only, file_formats):
        all_data = []
        header_written = False
        selected_encoding = self.encoding_combo.currentText()

        for root, dirs, files in os.walk(input_dir):
            for file in files:
                if any(file.endswith(fmt) for fmt in file_formats):
                    filepath = os.path.join(root, file)
                    try:
                        if file.endswith('.csv'):
                            encoding = self.detect_encoding(filepath) if selected_encoding == 'auto' else selected_encoding
                            with open(filepath, 'r', encoding=encoding, newline='') as f:
                                content = f.read()
                            content = content.replace('\r\n', '\n')  # 統一換行符
                            df = pd.read_csv(io.StringIO(content), skiprows=skip_rows)
                        elif file.endswith(('.xls', '.xlsx')):
                            df = pd.read_excel(filepath, skiprows=skip_rows)
                        else:
                            continue  # 跳過不支持的文件格式

                        if not header_written:
                            all_data.append(df.columns)
                            header_written = True

                        for _, row in df.iterrows():
                            if not exclude_keywords or (check_first_cell_only and not any(str(row.iloc[0]).strip().startswith(keyword) for keyword in exclude_keywords)) or (not check_first_cell_only and not any(any(str(cell).strip().startswith(keyword) for keyword in exclude_keywords) for cell in row)):
                                all_data.append(row)

                    except Exception as e:
                        print(f"處理文件 {filepath} 時出錯: {str(e)}")

        if all_data:
            output_df = pd.DataFrame(all_data)
            if output_file.endswith('.csv'):
                output_df.to_csv(output_file, index=False, header=False, encoding='utf-8-sig', line_terminator='\r\n')
            else:
                output_df.to_excel(output_file, index=False, header=False, engine='openpyxl')
        else:
            raise Exception("未找到有效的數據文件")

    def detect_encoding(self, filepath):
        with open(filepath, 'rb') as file:
            raw = file.read()
            result = chardet.detect(raw)
            return result['encoding']

class BatchPrintDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Batch Print PDF Files")
        self.setMinimumSize(1200, 600)
        self.file_paths = []  # 用於存儲添加的PDF文件路徑
        self.current_sort_column = -1
        self.sort_order = Qt.AscendingOrder
        self.initUI()

    def initUI(self):
        layout = QHBoxLayout(self)

        # 左側文件選擇區域
        left_layout = QVBoxLayout()

        # 創建 DropArea 並添加到佈局
        self.drop_area = DropArea()
        self.drop_area.fileDropped.connect(self.files_dropped)
        left_layout.addWidget(self.drop_area)

        # 創建表格部件
        self.table_widget = QTableWidget()
        self.table_widget.setColumnCount(6)  # 修改列數為6
        self.table_widget.setHorizontalHeaderLabels(["項目", "文件名", "頁數", "打印進度", "打印日期", "打印時間"])
        self.table_widget.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table_widget.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table_widget.verticalHeader().setVisible(False)
        self.table_widget.horizontalHeader().sectionClicked.connect(self.sort_files)  # 允許點擊表頭進行排序
        left_layout.addWidget(self.table_widget)

        file_btn_layout = QHBoxLayout()
        self.btn_add = QPushButton("載入PDF")
        self.btn_add.clicked.connect(self.add_pdf)
        file_btn_layout.addWidget(self.btn_add)

        self.btn_remove = QPushButton("移除PDF")
        self.btn_remove.clicked.connect(self.remove_pdf)
        file_btn_layout.addWidget(self.btn_remove)

        # 添加"上移"和"下移"按鈕
        self.btn_up = QPushButton("上移")
        self.btn_up.clicked.connect(self.move_up)
        file_btn_layout.addWidget(self.btn_up)

        self.btn_down = QPushButton("下移")
        self.btn_down.clicked.connect(self.move_down)
        file_btn_layout.addWidget(self.btn_down)

        left_layout.addLayout(file_btn_layout)

        self.total_label = QLabel("PDF文件總數：0")
        left_layout.addWidget(self.total_label)

        layout.addLayout(left_layout, 4)  # 左邊區域佔4

        # 右側打印選項區域
        right_layout = QVBoxLayout()
        
        # 打印機選擇
        self.setupPrinterSelection(right_layout)

        # 紙張大小選項
        self.setupPaperSize(right_layout)

        # 打印方向選項
        self.setupOrientation(right_layout)

        # 偏移量選項
        self.setupOffsets(right_layout)

        # 打印份數選項
        self.setupCopies(right_layout)

        # 雙面打印選項
        self.setupDuplex(right_layout)

        # 全頁打印選項（無邊距）
        self.setupFitToMargin(right_layout)

        # 打印按鈕
        self.btn_print = QPushButton("Batch Print")
        self.btn_print.clicked.connect(self.start_batch_printing)
        right_layout.addWidget(self.btn_print)

        layout.addLayout(right_layout, 1)  # 右邊區域佔1

        # 打印進度顯示
        self.progress_text = QTextEdit()
        self.progress_text.setReadOnly(True)
        layout.addWidget(self.progress_text)

    def setupPrinterSelection(self, layout):
        group_box = QGroupBox("Printer Selection")
        group_layout = QVBoxLayout()
        printer_label = QLabel("Select Printer:")
        self.printer_combo = QComboBox()
        group_layout.addWidget(printer_label)
        group_layout.addWidget(self.printer_combo)

        # 添加打印机偏好设置按钮
        self.printer_preferences_button = QPushButton("Printer Preferences")
        self.printer_preferences_button.clicked.connect(self.open_printer_preferences)
        group_layout.addWidget(self.printer_preferences_button)

        group_box.setLayout(group_layout)
        layout.addWidget(group_box)
        self.populate_printers()

    def populate_printers(self):
        available_printers = QPrinterInfo.availablePrinters()
        default_printer = QPrinterInfo.defaultPrinter()

        for printer in available_printers:
            self.printer_combo.addItem(printer.printerName())
        
        index = self.printer_combo.findText(default_printer.printerName())
        if index >= 0:
            self.printer_combo.setCurrentIndex(index)

    def setupPaperSize(self, layout):
        group_box = QGroupBox("Paper Options")
        group_layout = QVBoxLayout()
        paper_size_label = QLabel("Paper Size:")
        self.paper_size_combo = QComboBox()
        self.paper_size_combo.addItems(["A4", "A3", "A5", "Letter"])
        group_layout.addWidget(paper_size_label)
        group_layout.addWidget(self.paper_size_combo)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupOrientation(self, layout):
        group_box = QGroupBox("Orientation")
        group_layout = QVBoxLayout()
        orientation_label = QLabel("Orientation:")
        self.portrait_radio = QRadioButton("Portrait")
        self.landscape_radio = QRadioButton("Landscape")
        self.auto_orientation_radio = QRadioButton("Auto")  # 新增自动选项
        self.auto_orientation_radio.setChecked(True)  # 默认选中自动
        orientation_layout = QHBoxLayout()
        orientation_layout.addWidget(self.portrait_radio)
        orientation_layout.addWidget(self.landscape_radio)
        orientation_layout.addWidget(self.auto_orientation_radio)
        group_layout.addWidget(orientation_label)
        group_layout.addLayout(orientation_layout)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupOffsets(self, layout):
        group_box = QGroupBox("Margin Shift")
        group_layout = QVBoxLayout()
        offsets_label = QLabel("Margin Shift (mm):")
        offsets_layout = QGridLayout()
        self.top_offset_spin = QDoubleSpinBox()
        self.bottom_offset_spin = QDoubleSpinBox()
        self.left_offset_spin = QDoubleSpinBox()
        self.right_offset_spin = QDoubleSpinBox()

        self.top_offset_spin.setMinimum(-9999)
        self.bottom_offset_spin.setMinimum(-9999)
        self.left_offset_spin.setMinimum(-9999)
        self.right_offset_spin.setMinimum(-9999)

        offsets_layout.addWidget(QLabel("Top:"), 0, 0)
        offsets_layout.addWidget(self.top_offset_spin, 0, 1)
        offsets_layout.addWidget(QLabel("Bottom:"), 1, 0)
        offsets_layout.addWidget(self.bottom_offset_spin, 1, 1)
        offsets_layout.addWidget(QLabel("Left:"), 0, 2)
        offsets_layout.addWidget(self.left_offset_spin, 0, 3)
        offsets_layout.addWidget(QLabel("Right:"), 1, 2)
        offsets_layout.addWidget(self.right_offset_spin, 1, 3)
        group_layout.addWidget(offsets_label)
        group_layout.addLayout(offsets_layout)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupFitToMargin(self, layout):
        self.fit_to_margin_checkbox = QCheckBox("Fit to Printer Margin")
        layout.addWidget(self.fit_to_margin_checkbox)

    def setupCopies(self, layout):
        group_box = QGroupBox("Number of Copies")
        group_layout = QVBoxLayout()
        copies_label = QLabel("Copies:")
        self.copies_spin = QSpinBox()
        self.copies_spin.setMinimum(1)
        group_layout.addWidget(copies_label)
        group_layout.addWidget(self.copies_spin)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupDuplex(self, layout):
        self.duplex_checkbox = QCheckBox("Enable Duplex Printing")
        layout.addWidget(self.duplex_checkbox)

    def open_printer_preferences(self):
        selected_printer_name = self.printer_combo.currentText()
        if selected_printer_name:
            printer = QPrinter()
            printer.setPrinterName(selected_printer_name)
            dialog = QPrintDialog(printer, self)
            dialog.exec_()

    def files_dropped(self, file_paths):
        for path in file_paths:
            if path not in self.file_paths:
                self.file_paths.append(path)
                self.add_table_row(path)
        self.update_total_label()

    def add_pdf(self):
        file_paths, _ = QFileDialog.getOpenFileNames(self, "選擇PDF文件", "", "PDF files (*.pdf)")
        for path in file_paths:
            if path not in self.file_paths:
                self.file_paths.append(path)
                self.add_table_row(path)
        self.update_total_label()

    def add_table_row(self, path):
        file_name = os.path.basename(path)
        with fitz.open(path) as doc:
            page_count = doc.page_count

        row_count = self.table_widget.rowCount()
        self.table_widget.insertRow(row_count)

        self.table_widget.setItem(row_count, 0, QTableWidgetItem(str(row_count + 1)))
        self.table_widget.setItem(row_count, 1, QTableWidgetItem(file_name))
        self.table_widget.setItem(row_count, 2, QTableWidgetItem(str(page_count)))
        self.table_widget.setItem(row_count, 3, QTableWidgetItem("等待中"))
        self.table_widget.setItem(row_count, 4, QTableWidgetItem(""))  # 打印日期
        self.table_widget.setItem(row_count, 5, QTableWidgetItem(""))  # 打印時間

    def remove_pdf(self):
        selected_rows = self.table_widget.selectionModel().selectedRows()
        for model_index in reversed(selected_rows):
            row = model_index.row()
            file_path = self.file_paths[row]
            self.file_paths.remove(file_path)
            self.table_widget.removeRow(row)
        self.update_item_numbers()
        self.update_total_label()

    def move_up(self):
        current_row = self.table_widget.currentRow()
        if current_row > 0:
            self.file_paths.insert(current_row - 1, self.file_paths.pop(current_row))
            self.table_widget.insertRow(current_row - 1)
            for col in range(self.table_widget.columnCount()):
                self.table_widget.setItem(current_row - 1, col, self.table_widget.takeItem(current_row + 1, col))
            self.table_widget.removeRow(current_row + 1)
            self.table_widget.selectRow(current_row - 1)
            self.update_item_numbers()

    def move_down(self):
        current_row = self.table_widget.currentRow()
        if current_row < self.table_widget.rowCount() - 1:
            self.file_paths.insert(current_row + 1, self.file_paths.pop(current_row))
            self.table_widget.insertRow(current_row + 2)
            for col in range(self.table_widget.columnCount()):
                self.table_widget.setItem(current_row + 2, col, self.table_widget.takeItem(current_row, col))
            self.table_widget.removeRow(current_row)
            self.table_widget.selectRow(current_row + 1)
            self.update_item_numbers()

    def update_item_numbers(self):
        for row in range(self.table_widget.rowCount()):
            self.table_widget.setItem(row, 0, QTableWidgetItem(str(row + 1)))

    def update_total_label(self):
        total_files = len(self.file_paths)
        self.total_label.setText(f"PDF文件總數：{total_files}")

    def sort_files(self, column):
        if column == self.current_sort_column:
            self.sort_order = Qt.DescendingOrder if self.sort_order == Qt.AscendingOrder else Qt.AscendingOrder
        else:
            self.current_sort_column = column
            self.sort_order = Qt.AscendingOrder

        try:
            if column == 1:  # 文件名
                self.file_paths.sort(key=lambda path: os.path.basename(path), reverse=self.sort_order == Qt.DescendingOrder)
            elif column == 2:  # 頁數
                self.file_paths.sort(key=lambda path: fitz.open(path).page_count, reverse=self.sort_order == Qt.DescendingOrder)
            elif column == 3:  # 打印進度
                self.file_paths.sort(key=lambda path: self.table_widget.item(self.file_paths.index(path), 3).text(), reverse=self.sort_order == Qt.DescendingOrder)
        except Exception as e:
            self.show_error(f"排序錯誤: {e}")

        self.update_table()

    def update_table(self):
        self.table_widget.setRowCount(0)
        for path in self.file_paths:
            self.add_table_row(path)
        self.update_item_numbers()

    def start_batch_printing(self):
        if not self.file_paths:
            QMessageBox.warning(self, "警告", "請先添加要打印的PDF文件。")
            return

        selected_printer_name = self.printer_combo.currentText()
        print(f"Selected printer: {selected_printer_name}")  # Print the selected printer name

        paper_size = self.paper_size_combo.currentText()
        left_offset = self.left_offset_spin.value()
        top_offset = self.top_offset_spin.value()
        right_offset = self.right_offset_spin.value() + 5  # Hardcoded right margin
        bottom_offset = self.bottom_offset_spin.value() + 5  # Hardcoded bottom margin
        fit_to_margin = self.fit_to_margin_checkbox.isChecked()
        copies = self.copies_spin.value()
        duplex = self.duplex_checkbox.isChecked()

        progress_dialog = QProgressDialog("Batch Printing...", "Cancel", 0, len(self.file_paths), self)
        progress_dialog.setWindowTitle("Print Progress")
        progress_dialog.setWindowModality(Qt.WindowModal)
        progress_dialog.setMinimumDuration(0)

        print_log = []  # 用於記錄打印日誌

        for i, file_path in enumerate(self.file_paths):
            progress_dialog.setValue(i)
            progress_dialog.setLabelText(f"Printing {os.path.basename(file_path)} ({i + 1}/{len(self.file_paths)})")
            if progress_dialog.wasCanceled():
                break

            printer = QPrinter(QPrinter.HighResolution)
            printer.setPrinterName(selected_printer_name)  # Set the selected printer
            
            # Set the document name to the original PDF file name
            doc_name = os.path.basename(file_path)
            printer.setDocName(doc_name)

            if paper_size == "A4":
                printer.setPageSize(QPrinter.A4)
            elif paper_size == "A3":
                printer.setPageSize(QPrinter.A3)
            elif paper_size == "A5":
                printer.setPageSize(QPrinter.A5)
            elif paper_size == "Letter":
                printer.setPageSize(QPrinter.Letter)
            else:
                self.show_error(f"不支持的纸张尺寸: {paper_size}")
                return

            # Set orientation based on PDF
            self.set_orientation_based_on_pdf(printer, file_path)

            printer.setCopyCount(copies)

            if duplex:
                printer.setDuplex(QPrinter.DuplexAuto)
            else:
                printer.setDuplex(QPrinter.DuplexNone)

            self.print_direct(printer, left_offset, top_offset, right_offset, bottom_offset, fit_to_margin, file_path, i)

            # 記錄打印日期和時間
            print_date = datetime.now().strftime("%Y-%m-%d")
            print_time = datetime.now().strftime("%H:%M:%S")

            self.table_widget.setItem(i, 4, QTableWidgetItem(print_date))
            self.table_widget.setItem(i, 5, QTableWidgetItem(print_time))

            # 添加到打印日誌
            print_log.append({
                "項目": i + 1,
                "文件名": os.path.basename(file_path),
                "頁數": fitz.open(file_path).page_count,
                "打印進度": "完成",
                "打印日期": print_date,
                "打印完成時間": print_time
            })

        progress_dialog.setValue(len(self.file_paths))
        QMessageBox.information(self, "成功", "所有PDF文件已成功列印！")

    def set_orientation_based_on_pdf(self, printer, pdf_path):
        if self.auto_orientation_radio.isChecked():
            doc = fitz.open(pdf_path)
            first_page = doc[0]
            width, height = first_page.rect.width, first_page.rect.height
            doc.close()
            
            if width > height:
                printer.setOrientation(QPrinter.Landscape)
            else:
                printer.setOrientation(QPrinter.Portrait)
        elif self.landscape_radio.isChecked():
            printer.setOrientation(QPrinter.Landscape)
        else:
            printer.setOrientation(QPrinter.Portrait)

    def print_direct(self, printer, left_offset, top_offset, right_offset, bottom_offset, fit_to_margin, file_path, row_index):
        painter = QPainter(printer)
        try:
            doc = fitz.open(file_path)
            paper_rect = printer.paperRect()
            page_rect = printer.pageRect()

            for page_num in range(doc.page_count):
                page = doc.load_page(page_num)
                zoom = 5
                mat = fitz.Matrix(zoom, zoom)
                pixmap = page.get_pixmap(matrix=mat)
                qimage = QImage(pixmap.samples, pixmap.width, pixmap.height, pixmap.stride, QImage.Format_RGB888)

                if fit_to_margin:
                    target_rect = page_rect
                else:
                    target_rect = paper_rect

                scale_x = target_rect.width() / qimage.width()
                scale_y = target_rect.height() / qimage.height()
                scale = min(scale_x, scale_y)
                scaled_width = int(qimage.width() * scale)
                scaled_height = int(qimage.height() * scale)
                scaled_image = qimage.scaled(scaled_width, scaled_height, Qt.KeepAspectRatio, Qt.SmoothTransformation)

                image_rect = QRect(0, 0, scaled_image.width(), scaled_image.height())
                image_rect.moveCenter(target_rect.center())

                offset_x = int(left_offset - right_offset) * printer.logicalDpiX() / 25.4
                offset_y = int(top_offset - bottom_offset) * printer.logicalDpiY() / 25.4
                painter.translate(offset_x, offset_y)
                painter.drawImage(image_rect, scaled_image)

                # 更新打印進度
                self.progress_text.append(f"正在打印 {os.path.basename(file_path)} 的第 {page_num + 1} 頁，共 {doc.page_count} 頁")
                self.table_widget.setItem(row_index, 3, QTableWidgetItem(f"打印第 {page_num + 1} 頁，共 {doc.page_count} 頁"))

                # 確保界面即時更新
                QApplication.processEvents()

                if page_num < doc.page_count - 1:
                    printer.newPage()

                painter.resetTransform()

            doc.close()
            self.table_widget.setItem(row_index, 3, QTableWidgetItem("完成"))
        except Exception as e:
            self.show_error(f"Print error: {e}")
            self.table_widget.setItem(row_index, 3, QTableWidgetItem("錯誤"))
        finally:
            painter.end()

    def show_error(self, message):
        QMessageBox.critical(self, "錯誤", message)

    def show_warning(self, message):
        QMessageBox.warning(self, "警告", message)

    def show_info(self, message):
        QMessageBox.information(self, "信息", message)

class PDFOverlayDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("PDF Overlay")
        self.setMinimumSize(500, 300)
        self.setup_ui()
        self.parent_window = parent

    def setup_ui(self):
        layout = QVBoxLayout()
        layout.setSpacing(15)
        layout.setContentsMargins(20, 20, 20, 20)

        # Title Label
        title_label = QLabel("PDF Form Overlay Settings")
        title_label.setAlignment(Qt.AlignCenter)
        title_label.setStyleSheet("""
            font-weight: bold;
            font-size: 18px;
            color: #333;
            margin-bottom: 10px;
        """)
        layout.addWidget(title_label)

        # Form container with shadow effect
        form_container = QFrame()
        form_container.setFrameShape(QFrame.StyledPanel)
        form_container.setStyleSheet("""
            QFrame {
                background-color: white;
                border-radius: 8px;
                border: 1px solid #ddd;
            }
        """)
        
        shadow = QGraphicsDropShadowEffect()
        shadow.setBlurRadius(15)
        shadow.setColor(QColor(0, 0, 0, 30))
        shadow.setOffset(0, 3)
        form_container.setGraphicsEffect(shadow)
        
        form_layout = QVBoxLayout(form_container)
        form_layout.setSpacing(15)
        form_layout.setContentsMargins(20, 20, 20, 20)

        # Template file selection
        template_label = QLabel("Template PDF")
        template_label.setStyleSheet("font-weight: bold; color: #555;")
        
        template_field_layout = QHBoxLayout()
        template_field_layout.setSpacing(10)
        
        self.template_path_edit = QLineEdit()
        self.template_path_edit.setReadOnly(True)
        self.template_path_edit.setStyleSheet("""
            QLineEdit {
                border: 1px solid #ddd;
                border-radius: 4px;
                padding: 8px;
                background-color: #f9f9f9;
            }
        """)
        
        browse_template_button = QPushButton("Browse")
        browse_template_button.setCursor(Qt.PointingHandCursor)
        browse_template_button.setStyleSheet("""
            QPushButton {
                background-color: #f0f0f0;
                border: 1px solid #ddd;
                border-radius: 4px;
                padding: 8px 15px;
                color: #333;
            }
            QPushButton:hover {
                background-color: #e0e0e0;
            }
            QPushButton:pressed {
                background-color: #d0d0d0;
            }
        """)
        browse_template_button.clicked.connect(self.browse_template_pdf)
        
        template_field_layout.addWidget(self.template_path_edit)
        template_field_layout.addWidget(browse_template_button)
        
        form_layout.addWidget(template_label)
        form_layout.addLayout(template_field_layout)

        # Target folder selection
        target_folder_label = QLabel("Target Folder")
        target_folder_label.setStyleSheet("font-weight: bold; color: #555;")
        
        target_folder_field_layout = QHBoxLayout()
        target_folder_field_layout.setSpacing(10)
        
        self.target_folder_edit = QLineEdit()
        self.target_folder_edit.setReadOnly(True)
        self.target_folder_edit.setStyleSheet("""
            QLineEdit {
                border: 1px solid #ddd;
                border-radius: 4px;
                padding: 8px;
                background-color: #f9f9f9;
            }
        """)
        
        browse_target_folder_button = QPushButton("Browse")
        browse_target_folder_button.setCursor(Qt.PointingHandCursor)
        browse_target_folder_button.setStyleSheet("""
            QPushButton {
                background-color: #f0f0f0;
                border: 1px solid #ddd;
                border-radius: 4px;
                padding: 8px 15px;
                color: #333;
            }
            QPushButton:hover {
                background-color: #e0e0e0;
            }
            QPushButton:pressed {
                background-color: #d0d0d0;
            }
        """)
        browse_target_folder_button.clicked.connect(self.browse_target_folder)
        
        target_folder_field_layout.addWidget(self.target_folder_edit)
        target_folder_field_layout.addWidget(browse_target_folder_button)
        
        form_layout.addWidget(target_folder_label)
        form_layout.addLayout(target_folder_field_layout)

        # Output folder selection
        output_folder_label = QLabel("Output Folder")
        output_folder_label.setStyleSheet("font-weight: bold; color: #555;")
        
        output_folder_field_layout = QHBoxLayout()
        output_folder_field_layout.setSpacing(10)
        
        self.output_folder_edit = QLineEdit()
        self.output_folder_edit.setReadOnly(True)
        self.output_folder_edit.setStyleSheet("""
            QLineEdit {
                border: 1px solid #ddd;
                border-radius: 4px;
                padding: 8px;
                background-color: #f9f9f9;
            }
        """)
        
        browse_output_folder_button = QPushButton("Browse")
        browse_output_folder_button.setCursor(Qt.PointingHandCursor)
        browse_output_folder_button.setStyleSheet("""
            QPushButton {
                background-color: #f0f0f0;
                border: 1px solid #ddd;
                border-radius: 4px;
                padding: 8px 15px;
                color: #333;
            }
            QPushButton:hover {
                background-color: #e0e0e0;
            }
            QPushButton:pressed {
                background-color: #d0d0d0;
            }
        """)
        browse_output_folder_button.clicked.connect(self.browse_output_folder)
        
        output_folder_field_layout.addWidget(self.output_folder_edit)
        output_folder_field_layout.addWidget(browse_output_folder_button)
        
        form_layout.addWidget(output_folder_label)
        form_layout.addLayout(output_folder_field_layout)

        layout.addWidget(form_container)
        
        # Spacer
        layout.addSpacing(10)

        # 按钮布局
        button_layout = QHBoxLayout()
        button_layout.setSpacing(15)
        
        # Close button
        self.close_button = QPushButton("Cancel")
        self.close_button.setCursor(Qt.PointingHandCursor)
        self.close_button.setStyleSheet("""
            QPushButton {
                background-color: #f0f0f0;
                border: 1px solid #ddd;
                border-radius: 4px;
                padding: 10px 20px;
                color: #333;
                font-weight: bold;
                min-width: 100px;
            }
            QPushButton:hover {
                background-color: #e0e0e0;
            }
            QPushButton:pressed {
                background-color: #d0d0d0;
            }
        """)
        self.close_button.clicked.connect(self.reject)
        
        # Execute button
        self.execute_button = QPushButton("Execute")
        self.execute_button.setCursor(Qt.PointingHandCursor)
        self.execute_button.setStyleSheet("""
            QPushButton {
                background-color: #2196F3;
                border: none;
                border-radius: 4px;
                padding: 10px 20px;
                color: white;
                font-weight: bold;
                min-width: 100px;
            }
            QPushButton:hover {
                background-color: #1976D2;
            }
            QPushButton:pressed {
                background-color: #0D47A1;
            }
        """)
        self.execute_button.clicked.connect(self.execute_overlay)
        
        button_layout.addStretch()
        button_layout.addWidget(self.close_button)
        button_layout.addWidget(self.execute_button)
        
        layout.addLayout(button_layout)

        self.setLayout(layout)

    def browse_template_pdf(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select Template PDF", "", "PDF files (*.pdf)")
        if file_path:
            self.template_path_edit.setText(file_path)

    def browse_target_folder(self):
        folder_path = QFileDialog.getExistingDirectory(self, "Select Target Folder")
        if folder_path:
            self.target_folder_edit.setText(folder_path)

    def browse_output_folder(self):
        folder_path = QFileDialog.getExistingDirectory(self, "Select Output Folder")
        if folder_path:
            self.output_folder_edit.setText(folder_path)

    def get_paths(self):
        return self.template_path_edit.text(), self.target_folder_edit.text(), self.output_folder_edit.text()
    
    def execute_overlay(self):
        template_path, target_folder, output_folder = self.get_paths()
        
        # 检查所有路径是否都已设置
        if not template_path or not target_folder or not output_folder:
            msg_box = QMessageBox(self)
            msg_box.setIcon(QMessageBox.Warning)
            msg_box.setWindowTitle("Missing Information")
            msg_box.setText("Please ensure all paths are specified before executing.")
            msg_box.setStyleSheet("""
                QMessageBox {
                    background-color: white;
                    color: #333;
                }
                QPushButton {
                    background-color: #f0f0f0;
                    border: 1px solid #ddd;
                    border-radius: 4px;
                    padding: 5px 15px;
                    color: #333;
                }
                QPushButton:hover {
                    background-color: #e0e0e0;
                }
            """)
            msg_box.exec_()
            return
        
        # 调用父窗口的overlay_pdf方法进行处理
        if self.parent_window:
            self.parent_window.overlay_pdf(template_path, target_folder, output_folder)


class BatchConvertTXTtoPDFDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Batch Convert TXT to PDF")
        self.setMinimumSize(600, 400)

        layout = QVBoxLayout()

        # Add button to select txt files
        self.btn_select_txt_files = QPushButton("Select TXT Files", self)
        self.btn_select_txt_files.clicked.connect(self.select_txt_files)
        layout.addWidget(self.btn_select_txt_files)

        # Table widget to show selected files
        self.txt_files_table = QTableWidget(self)
        self.txt_files_table.setColumnCount(3)
        self.txt_files_table.setHorizontalHeaderLabels(["File Name", "File Path", "Remove"])
        layout.addWidget(self.txt_files_table)

        # Options to convert
        self.convert_option_group = QButtonGroup(self)
        self.radio_convert_single = QRadioButton("Convert to a Single PDF", self)
        self.radio_convert_multiple = QRadioButton("Convert to Multiple PDFs", self)
        self.radio_convert_multiple.setChecked(True)
        self.convert_option_group.addButton(self.radio_convert_single)
        self.convert_option_group.addButton(self.radio_convert_multiple)

        option_layout = QHBoxLayout()
        option_layout.addWidget(self.radio_convert_single)
        option_layout.addWidget(self.radio_convert_multiple)
        layout.addLayout(option_layout)

        # Output file name input for single PDF conversion
        self.output_filename_label = QLabel("Output File Name (for single PDF):", self)
        self.output_filename_input = QLineEdit(self)
        layout.addWidget(self.output_filename_label)
        layout.addWidget(self.output_filename_input)
        self.output_filename_label.hide()
        self.output_filename_input.hide()

        self.radio_convert_single.toggled.connect(self.toggle_output_filename_input)

        # Add button to start conversion
        self.btn_convert = QPushButton("Convert", self)
        self.btn_convert.clicked.connect(self.convert_to_pdf)
        layout.addWidget(self.btn_convert)

        self.setLayout(layout)

    def select_txt_files(self):
        file_paths, _ = QFileDialog.getOpenFileNames(self, "Select TXT Files", "", "Text files (*.txt)")
        if file_paths:
            for file_path in file_paths:
                file_name = os.path.basename(file_path)
                row_position = self.txt_files_table.rowCount()
                self.txt_files_table.insertRow(row_position)
                self.txt_files_table.setItem(row_position, 0, QTableWidgetItem(file_name))
                self.txt_files_table.setItem(row_position, 1, QTableWidgetItem(file_path))
                
                # Add remove button in the last column
                remove_button = QPushButton("Remove")
                remove_button.clicked.connect(lambda _, row=row_position: self.remove_file(row))
                self.txt_files_table.setCellWidget(row_position, 2, remove_button)

            # Resize columns to fit content
            self.txt_files_table.resizeColumnsToContents()

    def remove_file(self, row):
        self.txt_files_table.removeRow(row)
        # Resize columns to fit content after removing a row
        self.txt_files_table.resizeColumnsToContents()

    def toggle_output_filename_input(self):
        if self.radio_convert_single.isChecked():
            self.output_filename_label.show()
            self.output_filename_input.show()
        else:
            self.output_filename_label.hide()
            self.output_filename_input.hide()

    def convert_to_pdf(self):
        txt_files = [self.txt_files_table.item(row, 1).text() for row in range(self.txt_files_table.rowCount())]
        if not txt_files:
            QMessageBox.warning(self, "Warning", "No TXT files selected.")
            return

        save_dir = QFileDialog.getExistingDirectory(self, "Select Directory to Save PDFs")
        if not save_dir:
            QMessageBox.warning(self, "Warning", "No directory selected.")
            return

        if self.radio_convert_single.isChecked():
            self.convert_to_single_pdf(txt_files, save_dir)
        else:
            self.convert_to_multiple_pdfs(txt_files, save_dir)

    def convert_to_single_pdf(self, txt_files, save_dir):
        try:
            combined_text = ""
            for txt_file in txt_files:
                with open(txt_file, 'r', encoding='utf-8') as f:
                    combined_text += f.read() + "\n\n"  # Add some spacing between the contents of different files
            pdf_filename = self.output_filename_input.text().strip()
            if not pdf_filename:
                QMessageBox.warning(self, "Warning", "Output file name cannot be empty.")
                return

            pdf_path = os.path.join(save_dir, f"{pdf_filename}.pdf")

            document = QTextDocument()
            document.setPlainText(combined_text)

            printer = QPrinter(QPrinter.HighResolution)
            printer.setOutputFormat(QPrinter.PdfFormat)
            printer.setOutputFileName(pdf_path)

            document.print_(printer)
            QMessageBox.information(self, "Conversion Successful", f"Successfully converted to {pdf_path}")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to convert to single PDF. Error: {str(e)}")

    def convert_to_multiple_pdfs(self, txt_files, save_dir):
        try:
            for txt_file in txt_files:
                with open(txt_file, 'r', encoding='utf-8') as f:
                    text = f.read()

                file_name = os.path.basename(txt_file)
                pdf_filename = os.path.splitext(file_name)[0] + '.pdf'
                pdf_path = os.path.join(save_dir, pdf_filename)

                document = QTextDocument()
                document.setPlainText(text)

                printer = QPrinter(QPrinter.HighResolution)
                printer.setOutputFormat(QPrinter.PdfFormat)
                printer.setOutputFileName(pdf_path)

                document.print_(printer)
            QMessageBox.information(self, "Conversion Successful", "Successfully converted all TXT files to PDFs.")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to convert to multiple PDFs. Error: {str(e)}")



class ConvertThread(QThread):
    progress_updated = pyqtSignal(int)
    file_converted = pyqtSignal(str)
    conversion_completed = pyqtSignal(int, int)

    def __init__(self, files, source_folder, target_folder):
        super().__init__()
        self.files = files
        self.source_folder = source_folder
        self.target_folder = target_folder

    def run(self):
        converted_count = 0
        for i, file_name in enumerate(self.files):
            source_path = os.path.join(self.source_folder, file_name)
            output_path = os.path.join(self.target_folder, os.path.splitext(file_name)[0] + '.pdf')
            try:
                self.convert_file(source_path, output_path)
                converted_count += 1
                self.file_converted.emit(file_name)
            except Exception as e:
                QMessageBox.warning(None, 'Conversion Error', str(e))
            self.progress_updated.emit(i + 1)
        self.conversion_completed.emit(converted_count, len(self.files))

    def convert_file(self, path, output_path):
        try:
            normalized_path = os.path.normpath(path)
            if not os.path.exists(normalized_path):
                raise FileNotFoundError(f"File not found: {normalized_path}")

            file_type = normalized_path.split('.')[-1]
            app = None
            if file_type in ['xlsx', 'xls']:
                app = comtypes.client.CreateObject('Excel.Application')
                doc = app.Workbooks.Open(normalized_path)
                doc.ExportAsFixedFormat(0, output_path)
            elif file_type in ['docx', 'doc']:
                app = comtypes.client.CreateObject('Word.Application')
                doc = app.Documents.Open(normalized_path)
                doc.SaveAs(output_path, FileFormat=17)
            elif file_type in ['pptx', 'ppt']:
                app = comtypes.client.CreateObject('PowerPoint.Application')
                doc = app.Presentations.Open(normalized_path)
                doc.SaveAs(output_path, FileFormat=32)  # PDF format
            if app:
                doc.Close()
                app.Quit()
        except Exception as e:
            raise Exception(f"Error converting file: {normalized_path}. {str(e)}")

class BatchConvertDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle('Batch Convert Office Files to PDF')
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)

        # Source and Target Folders
        folder_group = QGroupBox('Folders')
        folder_layout = QVBoxLayout()

        self.source_button = QPushButton('Select Source Folder')
        self.source_button.clicked.connect(self.select_source_folder)
        folder_layout.addWidget(self.source_button)

        self.source_label = QLabel()
        folder_layout.addWidget(self.source_label)

        self.target_button = QPushButton('Select Target Folder')
        self.target_button.clicked.connect(self.select_target_folder)
        folder_layout.addWidget(self.target_button)

        self.target_label = QLabel()
        folder_layout.addWidget(self.target_label)

        folder_group.setLayout(folder_layout)
        layout.addWidget(folder_group)

        # File Types
        file_type_group = QGroupBox('File Types')
        file_type_layout = QHBoxLayout()

        self.word_checkbox = QCheckBox('Word (.doc, .docx)')
        self.word_checkbox.setChecked(True)
        file_type_layout.addWidget(self.word_checkbox)
        
        self.excel_checkbox = QCheckBox('Excel (.xls, .xlsx)')
        self.excel_checkbox.setChecked(True)
        file_type_layout.addWidget(self.excel_checkbox)
        
        self.powerpoint_checkbox = QCheckBox('PowerPoint (.ppt, .pptx)')
        self.powerpoint_checkbox.setChecked(True)
        file_type_layout.addWidget(self.powerpoint_checkbox)

        file_type_group.setLayout(file_type_layout)
        layout.addWidget(file_type_group)

        # Overwrite Option
        self.overwrite_checkbox = QCheckBox('Overwrite existing PDF files')
        layout.addWidget(self.overwrite_checkbox)

        # Conversion Buttons
        button_layout = QHBoxLayout()

        self.convert_button = QPushButton('Convert')
        self.convert_button.clicked.connect(self.start_conversion)
        button_layout.addWidget(self.convert_button)

        self.cancel_button = QPushButton('Cancel')
        self.cancel_button.clicked.connect(self.cancel_conversion)
        button_layout.addWidget(self.cancel_button)

        layout.addLayout(button_layout)

        # Progress and Status
        progress_group = QGroupBox('Progress and Status')
        progress_layout = QVBoxLayout()

        self.progress_bar = QProgressBar(self)
        progress_layout.addWidget(self.progress_bar)

        self.status_label = QLabel('Conversion Status:')
        progress_layout.addWidget(self.status_label)

        self.status_text_edit = QTextEdit()
        self.status_text_edit.setReadOnly(True)
        progress_layout.addWidget(self.status_text_edit)

        progress_group.setLayout(progress_layout)
        layout.addWidget(progress_group)

        # Open Target Folder Button
        self.open_folder_button = QPushButton('Open Target Folder')
        self.open_folder_button.clicked.connect(self.open_target_folder)
        self.open_folder_button.setEnabled(False)
        layout.addWidget(self.open_folder_button)

    def select_source_folder(self):
        self.source_folder = QFileDialog.getExistingDirectory(self, "Select Source Folder")
        self.source_label.setText(f"Source Folder: {self.source_folder}")

    def select_target_folder(self):
        self.target_folder = QFileDialog.getExistingDirectory(self, "Select Target Folder")
        self.target_label.setText(f"Target Folder: {self.target_folder}")

    def start_conversion(self):
        if not hasattr(self, 'source_folder') or not hasattr(self, 'target_folder'):
            QMessageBox.warning(self, 'Error', 'Please select both source and target folders.')
            return

        selected_extensions = []
        if self.word_checkbox.isChecked():
            selected_extensions.extend(['.doc', '.docx'])
        if self.excel_checkbox.isChecked():
            selected_extensions.extend(['.xls', '.xlsx'])
        if self.powerpoint_checkbox.isChecked():
            selected_extensions.extend(['.ppt', '.pptx'])

        files = [f for f in os.listdir(self.source_folder) if f.lower().endswith(tuple(selected_extensions))]
        total_files = len(files)
        self.progress_bar.setMaximum(total_files)
        self.progress_bar.setValue(0)
        self.status_text_edit.clear()
        self.convert_thread = ConvertThread(files, self.source_folder, self.target_folder)
        self.convert_thread.progress_updated.connect(self.update_progress)
        self.convert_thread.file_converted.connect(self.update_current_file)
        self.convert_thread.conversion_completed.connect(self.conversion_completed)
        self.convert_thread.start()

        self.convert_button.setEnabled(False)
        self.cancel_button.setEnabled(True)

    def update_progress(self, value):
        self.progress_bar.setValue(value)

    def update_current_file(self, file_name):
        self.status_text_edit.append(f"Processing: {file_name}")

    def conversion_completed(self, converted_count, total_files):
        self.status_text_edit.append(f"Conversion Complete. {converted_count}/{total_files} files converted.")
        self.convert_button.setEnabled(True)
        self.cancel_button.setEnabled(False)
        self.open_folder_button.setEnabled(True)

    def cancel_conversion(self):
        self.convert_thread.terminate()
        self.convert_button.setEnabled(True)
        self.cancel_button.setEnabled(False)
        self.status_text_edit.append("Conversion Canceled.")

    def open_target_folder(self):
        os.startfile(self.target_folder)

class DecryptDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("移除 PDF 密碼")
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout()

        # 密碼輸入
        self.password_input = QLineEdit()
        self.password_input.setEchoMode(QLineEdit.Password)
        layout.addWidget(QLabel("輸入 PDF 密碼:"))
        layout.addWidget(self.password_input)

        # 按鈕
        button_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        button_box.accepted.connect(self.accept)
        button_box.rejected.connect(self.reject)
        layout.addWidget(button_box)

        self.setLayout(layout)

    def get_password(self):
        return self.password_input.text()

class EncryptDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("加密 PDF")
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout()

        # 密碼輸入
        self.password_input = QLineEdit()
        self.password_input.setEchoMode(QLineEdit.Password)
        layout.addWidget(QLabel("輸入密碼:"))
        layout.addWidget(self.password_input)

        # 確認密碼輸入
        self.confirm_password_input = QLineEdit()
        self.confirm_password_input.setEchoMode(QLineEdit.Password)
        layout.addWidget(QLabel("確認密碼:"))
        layout.addWidget(self.confirm_password_input)

        # 按鈕
        button_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        button_box.accepted.connect(self.accept)
        button_box.rejected.connect(self.reject)
        layout.addWidget(button_box)

        self.setLayout(layout)

    def get_password(self):
        return self.password_input.text(), self.confirm_password_input.text()


class READMEDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("README")
        self.setGeometry(100, 100, 600, 400)
        
        layout = QVBoxLayout()
        
        # 將你的README內容作為字符串賦值給readme_content變量
        readme_content = """
README

PDFDocuEdit Pro v0.910b (28/09/2024) 
Program developer : Andy Leung
------------------------------------------------------

PDFDocuEdit Pro is a powerful PDF editing and management tool designed to streamline your workflow and boost productivity. With its user-friendly interface and rich feature set, it enables you to effortlessly view, edit, and manipulate PDF documents.

## Key Features:

### Viewing and Navigation
- Seamless browsing of PDF documents
- Zoom in and out functionality
- Jump to specific pages
- View document properties

### Editing and Annotation
- Rotate pages
- Insert new pages from other PDFs
- Extract specific pages
- Delete unwanted pages

### Merging and Splitting
- Combine multiple PDF files into a single document
- Split large PDFs into smaller, more manageable files

### Text Extraction
- Extract text from selected areas of a PDF
- Export extracted text to Excel for further analysis or editing

### Search and Highlight
- Quickly locate specific text within documents
- Highlight important information

### Page Ordering
- Reorder pages within PDF documents to create custom structures

### Conversion Functions
- Convert PDF documents to editable Word files
- Batch convert MS Office documents to PDF
- Batch convert TXT files to PDF

### Security Features
- Encrypt PDF files to protect sensitive information
- Decrypt PDF files

### Page Count Statistics
- Generate comprehensive reports with page counts and file details for all PDFs in a folder

### Form Overlay
- Overlay PDF forms onto other PDF documents

### Batch Printing
- Print multiple PDF files in batch, with customizable print settings

### File Merging
- Merge multiple CSV or Excel files

## System Requirements:
- Microsoft Windows 10 or 11

## Installation:
1. Download the `PDFDocuEdit Pro.zip` file
2. Extract the contents of the zip file
3. Run the `PDFDocuEdit Pro.exe` file to launch the application

## Getting Started:
1. Open a PDF document using the "Open" button or by dragging and dropping files onto the application window
2. Explore the various features and tools available in the menus and toolbar
3. Utilize the intuitive interface to edit, manipulate, and manage your PDF documents

## Support:
For any questions or technical support, please contact Andy Leung @ andy846@gmail.com

We hope PDFDocuEdit Pro helps you work more efficiently and effectively with your PDF documents. Enjoy using it!

        """
        
        # 創建一個文本框來顯示README內容
        self.text_edit = QTextEdit()
        self.text_edit.setText(readme_content)
        self.text_edit.setReadOnly(True)  # 設置為只讀
        
        layout.addWidget(self.text_edit)
        
        self.setLayout(layout)

class SplitOptionsWidget(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.initUI()

    def initUI(self):
        layout = QVBoxLayout()

        # 创建一个按钮组，用于管理互斥的单选按钮
        self.radio_button_group = QButtonGroup()

        # 選項 1：在預定義頁面分割
        group_box_1 = QGroupBox("在預定義頁面分割")
        group_layout_1 = QVBoxLayout()
        self.radio_every_page = QRadioButton("每頁")
        self.radio_even_pages = QRadioButton("偶數頁")
        self.radio_odd_pages = QRadioButton("奇數頁")
        group_layout_1.addWidget(self.radio_every_page)
        group_layout_1.addWidget(self.radio_even_pages)
        group_layout_1.addWidget(self.radio_odd_pages)
        group_box_1.setLayout(group_layout_1)
        layout.addWidget(group_box_1)
        layout.addStretch(1)

        # 添加到按钮组
        self.radio_button_group.addButton(self.radio_every_page)
        self.radio_button_group.addButton(self.radio_even_pages)
        self.radio_button_group.addButton(self.radio_odd_pages)

        # 選項 2：指定頁面分割
        group_box_2 = QGroupBox("指定頁面分割")
        group_layout_2 = QVBoxLayout()
        self.radio_custom_pages = QRadioButton("自定義頁面")
        self.custom_pages_edit = QLineEdit()
        self.custom_pages_edit.setPlaceholderText("例如：1,3,5-7")
        self.custom_pages_edit.setEnabled(False)
        self.radio_custom_pages.toggled.connect(self.custom_pages_edit.setEnabled)
        group_layout_2.addWidget(self.radio_custom_pages)
        group_layout_2.addWidget(self.custom_pages_edit)
        group_box_2.setLayout(group_layout_2)
        layout.addWidget(group_box_2)
        layout.addStretch(1)

        # 添加到按钮组
        self.radio_button_group.addButton(self.radio_custom_pages)

        # 選項 3：每 N 頁分割
        group_box_3 = QGroupBox("每 N 頁分割")
        group_layout_3 = QVBoxLayout()
        self.radio_every_n_pages = QRadioButton("每 N 頁")
        self.every_n_pages_spin = QSpinBox()
        self.every_n_pages_spin.setMinimum(1)
        self.every_n_pages_spin.setMaximum(10000)
        self.every_n_pages_spin.setEnabled(False)
        self.radio_every_n_pages.toggled.connect(self.every_n_pages_spin.setEnabled)
        group_layout_3.addWidget(self.radio_every_n_pages)
        group_layout_3.addWidget(self.every_n_pages_spin)
        group_box_3.setLayout(group_layout_3)
        layout.addWidget(group_box_3)
        layout.addStretch(1)

        # 添加到按钮组
        self.radio_button_group.addButton(self.radio_every_n_pages)

        # 创建GroupBox
        save_location_group = QGroupBox("保存文件设置")
        group_layout = QVBoxLayout()  # 创建GroupBox的垂直布局

        # 第一行：保存文件位置
        save_location_layout = QVBoxLayout()  # 创建垂直布局
        self.save_location_label = QLabel("未選擇位置")  # 创建标签
        self.save_location_label.setWordWrap(True)  # 设置标签文本自动换行
        self.choose_location_button = QPushButton("選擇保存位置")  # 创建按钮
        self.choose_location_button.clicked.connect(self.openFileDialog)  # 链接事件
        save_location_layout.addWidget(self.save_location_label)
        save_location_layout.addWidget(self.choose_location_button)
        group_layout.addLayout(save_location_layout)  # 添加到GroupBox的布局

        # 第二行：文件名前缀输入
        file_prefix_layout = QHBoxLayout()  # 创建水平布局
        self.file_prefix_label = QLabel("输入文件名前缀:")
        self.file_prefix_input = QLineEdit()  # 文本输入框
        file_prefix_layout.addWidget(self.file_prefix_label)
        file_prefix_layout.addWidget(self.file_prefix_input)
        group_layout.addLayout(file_prefix_layout)  # 添加到GroupBox的布局

        # 将布局设置给GroupBox
        save_location_group.setLayout(group_layout)

        # 将GroupBox添加到主布局
        layout.addWidget(save_location_group)


        layout.addStretch(8)  # 添加彈性空間

        # 分割按鈕
        self.split_button = QPushButton("分割")
        layout.addWidget(self.split_button)

        self.setLayout(layout)

    def openFileDialog(self):
        # 打开一个文件夹选择对话框
        directory = QFileDialog.getExistingDirectory(self, "選擇保存位置")
        
        # 如果用户选择了一个目录，则更新标签
        if directory:
            self.save_location_label.setText(directory)

    def get_split_options(self):
        if self.radio_every_page.isChecked():
            return "every_page"
        elif self.radio_even_pages.isChecked():
            return "even_pages"
        elif self.radio_odd_pages.isChecked():
            return "odd_pages"
        elif self.radio_custom_pages.isChecked():
            return "custom_pages", self.custom_pages_edit.text()
        elif self.radio_every_n_pages.isChecked():
            return "every_n_pages", self.every_n_pages_spin.value()
        return None
class LoadFontsThread(QThread):
    update_signal = pyqtSignal(list)

    def __init__(self, pdf_doc):
        super().__init__()
        self.pdf_doc = pdf_doc

    def run(self):
        font_dict = {}
        for page in self.pdf_doc:
            if hasattr(page, 'get_fonts'):
                fonts = page.get_fonts(full=True)
                for font in fonts:
                    font_flags, font_name, font_family, font_full_name, font_type, font_encoding, font_file = font
                    # 尝试进行 UTF-8 解码
                    try:
                        font_full_name = font_full_name.encode('latin1').decode('utf-8')
                    except (UnicodeEncodeError, UnicodeDecodeError):
                        pass
                    embedded = "Yes" if font_flags & 4 else "No"
                    subset = "Yes" if font_name.startswith('X') else "No"
                    font_key = (font_full_name, font_type, font_family, embedded, subset)
                    
                    if font_key in font_dict:
                        font_dict[font_key] += 1
                    else:
                        font_dict[font_key] = 1

        font_list = [
            {"name": name, "type": type_, "family": family, "embedded": embedded, "subset": subset, "count": count}
            for (name, type_, family, embedded, subset), count in font_dict.items()
        ]
        self.update_signal.emit(font_list)

class PDFInfoDialog(QDialog):
    def __init__(self, pdf_doc, parent=None):
        super().__init__(parent)
        self.pdf_doc = pdf_doc
        self.setWindowTitle("PDF内容")
        self.setMinimumSize(600, 380)
        self.layout = QVBoxLayout(self)

        self.tab_widget = QTabWidget(self)
        self.tab_widget.currentChanged.connect(self.tab_changed)
        
        self.create_general_info_tab()
        self.layout.addWidget(self.tab_widget)

        self.fonts_thread = LoadFontsThread(pdf_doc)
        self.fonts_thread.update_signal.connect(self.append_font_info)

    def create_general_info_tab(self):
        # 創建一個QWidget作為標籤頁的主控件
        widget = QWidget()

        # 創建一個可滾動區域
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)

        # 創建布局管理器
        layout = QVBoxLayout()

        # 檔案屬性分組
        group_box_file = QGroupBox("檔案屬性")
        layout_file = QVBoxLayout()
        label_path = QLabel(f"文件路徑: {self.pdf_doc.name}")
        label_name = QLabel(f"文件名: {os.path.basename(self.pdf_doc.name)}")
        label_size = QLabel(f"文件大小: {self.format_file_size(os.path.getsize(self.pdf_doc.name))}")
        layout_file.addWidget(label_path)
        layout_file.addWidget(label_name)
        layout_file.addWidget(label_size)
        group_box_file.setLayout(layout_file)

        # 頁面資訊分組
        group_box_page = QGroupBox("頁面資訊")
        layout_page = QVBoxLayout()
        page_size = self.pdf_doc[0].rect
        width = round(page_size.width * 0.3528, 2)  # 將點數轉換為毫米
        height = round(page_size.height * 0.3528, 2)
        label_page_size = QLabel(f"頁面尺寸: {width}mm × {height}mm")
        label_page_count = QLabel(f"頁數: {self.pdf_doc.page_count}")
        layout_page.addWidget(label_page_size)
        layout_page.addWidget(label_page_count)
        group_box_page.setLayout(layout_page)

        # 文件元數據分組
        group_box_meta = QGroupBox("文件元數據")
        layout_meta = QVBoxLayout()
        info = self.pdf_doc.metadata
        label_creation_date = QLabel(f"創建日期: {info.get('CreationDate', '未知')}")
        label_mod_date = QLabel(f"修改日期: {info.get('ModDate', '未知')}")
        label_author = QLabel(f"作者: {info.get('Author', '未知')}")
        label_encrypted = QLabel(f"加密: {'是' if self.pdf_doc.is_encrypted else '否'}")
        label_producer = QLabel(f"生成器: {info.get('producer', '未知')}")
        layout_meta.addWidget(label_creation_date)
        layout_meta.addWidget(label_mod_date)
        layout_meta.addWidget(label_author)
        layout_meta.addWidget(label_encrypted)
        layout_meta.addWidget(label_producer)
        group_box_meta.setLayout(layout_meta)

        # 將分組框添加到布局
        layout.addWidget(group_box_file)
        layout.addWidget(group_box_page)
        layout.addWidget(group_box_meta)

        # 設置小部件的布局
        scroll_area.setWidget(widget)
        widget.setLayout(layout)

        # 在標籤頁控件中添加新的標籤頁
        self.tab_widget.addTab(scroll_area, "一般資訊")

        self.fonts_images_tab = QWidget()
        self.tab_widget.addTab(self.fonts_images_tab, "字體訊息")

    def tab_changed(self, index):
        if index == 1 and self.fonts_images_tab.layout() is None:
            self.setup_fonts_images_tab()
            self.fonts_thread.start()

    def setup_fonts_images_tab(self):
        # 創建一個QWidget作為標籤頁的主控件
        widget = QWidget()

        # 創建總布局管理器
        layout = QVBoxLayout(widget)  # 將布局直接綁定到widget上

        # 字體信息表格
        self.fonts_table = QTableWidget()
        self.fonts_table.setColumnCount(6)
        self.fonts_table.setHorizontalHeaderLabels(["名稱", "類型", "字型", "嵌入", "子集", "使用次數"])
        self.fonts_table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.fonts_table)

        # 设置表格字体为支持中文的字体，例如 SimSun 或 Microsoft YaHei
        font = QFont("Microsoft YaHei", 10)
        self.fonts_table.setFont(font)

        # 創建一個可滾動區域並設置其內容部件
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_area.setWidget(widget)  # 設置scroll_area的內容部件

        # 在標籤頁控件中設置scroll_area為tab的主內容
        self.fonts_images_tab.setLayout(QVBoxLayout())  # 確保tab有自己的布局
        self.fonts_images_tab.layout().addWidget(scroll_area)  # 將scroll_area添加到tab的布局中

    def append_font_info(self, font_list):
        self.fonts_table.setRowCount(len(font_list))
        for row, font in enumerate(font_list):
            self.fonts_table.setItem(row, 0, QTableWidgetItem(font["name"]))
            self.fonts_table.setItem(row, 1, QTableWidgetItem(font["type"]))
            self.fonts_table.setItem(row, 2, QTableWidgetItem(font["family"]))
            self.fonts_table.setItem(row, 3, QTableWidgetItem(font["embedded"]))
            self.fonts_table.setItem(row, 4, QTableWidgetItem(font["subset"]))
            self.fonts_table.setItem(row, 5, QTableWidgetItem(str(font["count"])))

    def format_file_size(self, size_bytes):
        if size_bytes == 0:
            return "0B"
        size_name = ("B", "KB", "MB", "GB", "TB", "PB", "EB", "ZB", "YB")
        i = int(math.floor(math.log(size_bytes, 1024)))
        p = math.pow(1024, i)
        s = round(size_bytes / p, 2)
        return f"{s} {size_name[i]}"
    

class SortWidget(QDockWidget):
    sort_applied = pyqtSignal(list)

    def __init__(self, parent=None):
        super().__init__("Page Sorting", parent)
        self.setAllowedAreas(Qt.LeftDockWidgetArea | Qt.RightDockWidgetArea)

        widget = QWidget(self)
        layout = QVBoxLayout(widget)

        self.page_list = QListWidget()
        self.page_list.itemDoubleClicked.connect(self.jump_to_page)
        self.page_list.setSelectionMode(QAbstractItemView.ExtendedSelection)
        layout.addWidget(self.page_list)

        button_layout = QHBoxLayout()
        self.move_up_button = QPushButton("Up")
        self.move_up_button.clicked.connect(self.move_item_up)
        button_layout.addWidget(self.move_up_button)

        self.move_down_button = QPushButton("Down")
        self.move_down_button.clicked.connect(self.move_item_down)
        button_layout.addWidget(self.move_down_button)

        self.apply_button = QPushButton("Apply")
        self.apply_button.clicked.connect(self.apply_sort)
        button_layout.addWidget(self.apply_button)

        layout.addLayout(button_layout)
        self.setWidget(widget)

    def update_page_list(self, page_count):
        self.page_list.clear()
        for i in range(page_count):
            self.page_list.addItem(f"Page {i+1}")

    def move_item_up(self):
        row = self.page_list.currentRow()
        if row > 0:
            item = self.page_list.takeItem(row)
            self.page_list.insertItem(row - 1, item)
            self.page_list.setCurrentRow(row - 1)

    def move_item_down(self):
        row = self.page_list.currentRow()
        if row < self.page_list.count() - 1:
            item = self.page_list.takeItem(row)
            self.page_list.insertItem(row + 1, item)
            self.page_list.setCurrentRow(row + 1)

    def apply_sort(self):
        order = []
        for i in range(self.page_list.count()):
            item_text = self.page_list.item(i).text()
            page_num = int(item_text.split()[1]) - 1
            order.append(page_num)
        self.sort_applied.emit(order)

    def jump_to_page(self, item):
        page_num = int(item.text().split()[1]) - 1
        self.parent().jump_to_page(page_num)

class DropArea(QWidget):
    # 自定義信號，用於當文件被拖放時發送文件路徑
    fileDropped = pyqtSignal(list)

    def __init__(self, parent=None):
        super(DropArea, self).__init__(parent)
        self.setAcceptDrops(True)
        self.setFixedHeight(100)  # 設置 DropArea 的高度
        # 設置背景色
        self.setAutoFillBackground(True)
        palette = self.palette()
        palette.setColor(QPalette.Window, QColor('medium sea green'))
        self.setPalette(palette)
        self.text = "Drop PDF Files Here"


    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        file_paths = [url.toLocalFile() for url in event.mimeData().urls() if url.toString().lower().endswith('.pdf')]
        self.fileDropped.emit(file_paths)  # 發送信號，附帶文件路徑

    def paintEvent(self, event):
        painter = QPainter(self)
        pen = QPen(Qt.black, 2, Qt.DashLine)
        painter.setPen(pen)
        rect = self.rect().adjusted(1, 1, -1, -1)  # 調整矩形以適應邊框寬度
        painter.drawRect(rect)

        # 設置文本的字體和顏色
        painter.setPen(QColor('black'))
        painter.setFont(QFont('Arial', 10, QFont.Bold))
        
        # 獲取控件中心點並根據文本大小調整位置
        text_rect = QRect(rect.left(), rect.top(), rect.width(), rect.height())
        painter.drawText(text_rect, Qt.AlignCenter, self.text)


class SearchWidget(QDockWidget):
    search_requested = pyqtSignal(str)
    jump_to_page_requested = pyqtSignal(int, bool)  # 添加參數指示是否需要高亮

    def __init__(self, parent=None):
        super().__init__("Searching Tools", parent)
        self.setAllowedAreas(Qt.LeftDockWidgetArea | Qt.RightDockWidgetArea)

        widget = QWidget()
        layout = QVBoxLayout(widget)

        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("Enter search text")
        self.search_edit.returnPressed.connect(self.search)
        layout.addWidget(self.search_edit)

        self.search_button = QPushButton("Search")
        self.search_button.clicked.connect(self.search)
        layout.addWidget(self.search_button)

        # 添加高亮選項的勾選框，默認不勾選
        self.highlight_checkbox = QCheckBox("Highlight search results")
        self.highlight_checkbox.setChecked(False)  # 默認不勾選
        layout.addWidget(self.highlight_checkbox)

        self.result_list = QListWidget()
        self.result_list.itemDoubleClicked.connect(self.jump_to_result)
        self.result_list.setSelectionMode(QAbstractItemView.ExtendedSelection)
        layout.addWidget(self.result_list)

        self.result_count_label = QLabel()
        self.result_count_label.setStyleSheet("color: white;")
        layout.addWidget(self.result_count_label)

        self.extract_ALL_button = QPushButton("Extract All Pages")
        self.extract_ALL_button.clicked.connect(self.extract_ALL_pages)
        layout.addWidget(self.extract_ALL_button)

        self.delete_ALL_button = QPushButton("Delete All Pages")
        self.delete_ALL_button.clicked.connect(self.delete_ALL_pages)
        layout.addWidget(self.delete_ALL_button)

        self.clear_button = QPushButton("Clear Result")
        self.clear_button.clicked.connect(self.clear_search)
        layout.addWidget(self.clear_button)

        self.setWidget(widget)

    def showEvent(self, event):
        super().showEvent(event)
        # 使用 QTimer.singleShot 確保焦點設置在 UI 完全顯示後執行
        QTimer.singleShot(0, self.search_edit.setFocus)

    def search(self):
        self.result_list.clear()
        text = self.search_edit.text()
        if text:
            self.search_requested.emit(text)
        else:
            self.result_count_label.setText("No input for search.")

    def update_results(self, results):
        self.result_list.clear()
        for page_num, _ in results:
            self.result_list.addItem(f"Page {page_num + 1}")
        self.result_count_label.setText(f"Found {len(results)} results")

        # 更新頁面，根據高亮選項決定是否顯示高亮
        if self.parent():
            # 先清除當前所有高亮
            self.parent().clear_highlights()
            
            # 如果勾選了高亮選項，則重新顯示當前頁面以應用高亮
            if self.highlight_checkbox.isChecked():
                self.parent().show_page(self.parent().current_page)

    def clear_result_list(self):
        """清空搜索結果列表和相關顯示內容。"""
        try:
            self.result_list.clear()
            self.result_count_label.setText("")
        except Exception as e:
            print(f"清除結果列表時出錯: {e}")

    def extract_pages(self):
        try:
            pages_to_extract = [int(item.text().split()[1]) - 1 for item in self.result_list.selectedItems()]
            if pages_to_extract and self.parent():
                self.parent().extract_pages_by_number(pages_to_extract)
        except Exception as e:
            print(f"提取頁面時發生錯誤: {e}")

    def delete_pages(self):
        try:
            pages_to_delete = [int(item.text().split()[1]) - 1 for item in self.result_list.selectedItems()]
            if pages_to_delete and self.parent():
                self.parent().delete_pages_by_number(pages_to_delete)
        except Exception as e:
            print(f"刪除頁面時發生錯誤: {e}")

    def jump_to_result(self, item):
        try:
            page_num = int(item.text().split()[1]) - 1
            if self.parent():
                # 如果父窗口有 jump_to_page 方法，使用該方法
                if hasattr(self.parent(), 'jump_to_page_with_highlight'):
                    # 傳遞頁碼和是否高亮的標誌
                    self.parent().jump_to_page_with_highlight(page_num, self.highlight_checkbox.isChecked())
                else:
                    # 向後兼容：如果沒有新方法，使用舊方法
                    self.parent().jump_to_page(page_num)
                    
                    # 如果不需要高亮，則手動清除高亮
                    if not self.highlight_checkbox.isChecked() and hasattr(self.parent(), 'clear_highlights'):
                        self.parent().clear_highlights()
                        self.parent().show_page(page_num)  # 重新顯示頁面
        except Exception as e:
            print(f"跳轉到結果頁面時發生錯誤: {e}")

    def clear_search(self):
        try:
            self.search_edit.clear()
            self.result_list.clear()
            self.result_count_label.setText("")
            
            # 先檢查父對象是否存在，然後再調用其方法
            if self.parent():
                if hasattr(self.parent(), 'clear_highlights'):
                    self.parent().clear_highlights()
                # 確保 search_results 存在後再清除
                if hasattr(self.parent(), 'search_results'):
                    self.parent().search_results.clear()
                
                # 重新顯示當前頁面以刷新視圖
                if hasattr(self.parent(), 'show_page') and hasattr(self.parent(), 'current_page'):
                    self.parent().show_page(self.parent().current_page)
        except Exception as e:
            print(f"清除搜索時發生錯誤: {e}")

    def contextMenuEvent(self, event):
        selected_items = self.result_list.selectedItems()
        if selected_items:
            menu = QMenu(self)
            menu.setStyleSheet("""
                QMenu {
                    background-color: #272727; /* 深灰色背景 */
                    color: white; /* 白色文字 */
                    border: 1px solid #666666; /* 浅灰色边框 */
                    border-radius: 10px; /* 菜单边框圆角 */
                }
                QMenu::item {
                    padding: 5px 20px; /* 菜单项内边距 */
                    margin: 2px 2px; /* 菜单项外边距 */
                }
                QMenu::item:selected {
                    background-color: #02F78E; /* 选中时的背景色 */
                    color: black; /* 选中时的文字颜色 */
                    border-radius: 5px; /* 选中项的边框圆角 */
                    padding: 5px 20px; /* 选中项的内边距 */
                }
            """)
            extract_action = menu.addAction("Extract Pages")
            delete_action = menu.addAction("Delete Pages")
            action = menu.exec_(event.globalPos())
            
            if action == extract_action:
                self.extract_pages()
            elif action == delete_action:
                self.delete_pages()

    def extract_ALL_pages(self):
        try:
            # 獲取所有搜索結果頁面編號
            all_items = []
            for i in range(self.result_list.count()):
                all_items.append(self.result_list.item(i))
                
            pages_to_extract = [int(item.text().split()[1]) - 1 for item in all_items]
            
            if pages_to_extract and self.parent():
                self.parent().extract_pages_by_number(pages_to_extract)
            else:
                if self.parent():
                    QMessageBox.information(self, "Information", "No pages to extract.")
        except Exception as e:
            print(f"提取所有頁面時發生錯誤: {e}")

    def delete_ALL_pages(self):
        try:
            # 獲取所有搜索結果頁面編號
            all_items = []
            for i in range(self.result_list.count()):
                all_items.append(self.result_list.item(i))
                
            pages_to_delete = [int(item.text().split()[1]) - 1 for item in all_items]
            
            if pages_to_delete and self.parent():
                self.parent().delete_pages_by_number(pages_to_delete)
            else:
                if self.parent():
                    QMessageBox.information(self, "Information", "No pages to delete.")
        except Exception as e:
            print(f"刪除所有頁面時發生錯誤: {e}")
            
class MergePDFsDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("合併PDF文件")
        self.setMinimumSize(600, 300)
        self.file_paths = []  # 用於存儲添加的PDF文件路徑
        self.sort_column = 0
        self.sort_order = Qt.AscendingOrder
        self.initUI()

    def initUI(self):
        layout = QVBoxLayout(self)

        # 創建 DropArea 並添加到佈局
        self.drop_area = DropArea()
        self.drop_area.fileDropped.connect(self.files_dropped)
        layout.addWidget(self.drop_area)

        # 創建表格部件
        self.table_widget = QTableWidget()
        self.table_widget.setColumnCount(5)
        self.table_widget.setHorizontalHeaderLabels(["項目", "文件名", "頁數", "大小 (MB)", "修改日期"])
        self.table_widget.horizontalHeader().sectionClicked.connect(self.sort_files)
        self.table_widget.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table_widget.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table_widget.verticalHeader().setVisible(False)
        layout.addWidget(self.table_widget)

        btn_layout = QHBoxLayout()
        self.btn_add = QPushButton("載入PDF")
        self.btn_add.clicked.connect(self.add_pdf)
        btn_layout.addWidget(self.btn_add)

        self.btn_remove = QPushButton("移除PDF")
        self.btn_remove.clicked.connect(self.remove_pdf)
        btn_layout.addWidget(self.btn_remove)

        self.btn_up = QPushButton("上移")
        self.btn_up.clicked.connect(self.move_up)
        btn_layout.addWidget(self.btn_up)

        self.btn_down = QPushButton("下移")
        self.btn_down.clicked.connect(self.move_down)
        btn_layout.addWidget(self.btn_down)

        layout.addLayout(btn_layout)

        # 創建標籤以顯示PDF文件總數
        self.total_label = QLabel("PDF文件總數：0")
        layout.addWidget(self.total_label)

        self.btn_merge = QPushButton("Merge PDF")
        self.btn_merge.clicked.connect(self.merge_pdfs)
        layout.addWidget(self.btn_merge)

    def files_dropped(self, file_paths):
        for path in file_paths:
            if path not in self.file_paths:
                self.file_paths.append(path)
                self.add_table_row(path)
        self.update_total_label()

    def add_pdf(self):
        file_paths, _ = QFileDialog.getOpenFileNames(self, "選擇PDF文件", "", "PDF files (*.pdf)")
        for path in file_paths:
            if path not in self.file_paths:
                self.file_paths.append(path)
                self.add_table_row(path)
        self.update_total_label()

    def add_table_row(self, path):
        file_name = os.path.basename(path)
        file_size = os.path.getsize(path)
        file_size_mb = file_size / (1024 * 1024)  # 轉換為MB
        with fitz.open(path) as doc:
            page_count = doc.page_count

        modification_time = os.path.getmtime(path)
        modification_date = datetime.fromtimestamp(modification_time).strftime("%Y-%m-%d %H:%M:%S")

        row_count = self.table_widget.rowCount()
        self.table_widget.insertRow(row_count)

        self.table_widget.setItem(row_count, 0, QTableWidgetItem(str(row_count + 1)))
        self.table_widget.setItem(row_count, 1, QTableWidgetItem(file_name))
        self.table_widget.setItem(row_count, 2, QTableWidgetItem(str(page_count)))
        self.table_widget.setItem(row_count, 3, QTableWidgetItem(f"{file_size_mb:.2f}"))
        self.table_widget.setItem(row_count, 4, QTableWidgetItem(modification_date))

    def remove_pdf(self):
        selected_rows = self.table_widget.selectionModel().selectedRows()
        for model_index in reversed(selected_rows):
            row = model_index.row()
            file_path = self.file_paths[row]
            self.file_paths.remove(file_path)
            self.table_widget.removeRow(row)
        self.update_item_numbers()
        self.update_total_label()

    def move_up(self):
        current_row = self.table_widget.currentRow()
        if current_row > 0:
            self.file_paths.insert(current_row - 1, self.file_paths.pop(current_row))
            self.table_widget.insertRow(current_row - 1)
            for col in range(self.table_widget.columnCount()):
                self.table_widget.setItem(current_row - 1, col, self.table_widget.takeItem(current_row + 1, col))
            self.table_widget.removeRow(current_row + 1)
            self.table_widget.selectRow(current_row - 1)
            self.update_item_numbers()

    def move_down(self):
        current_row = self.table_widget.currentRow()
        if current_row < self.table_widget.rowCount() - 1:
            self.file_paths.insert(current_row + 1, self.file_paths.pop(current_row))
            self.table_widget.insertRow(current_row + 2)
            for col in range(self.table_widget.columnCount()):
                self.table_widget.setItem(current_row + 2, col, self.table_widget.takeItem(current_row, col))
            self.table_widget.removeRow(current_row)
            self.table_widget.selectRow(current_row + 1)
            self.update_item_numbers()

    def sort_files(self, column):
        if self.sort_column == column:
            self.sort_order = Qt.AscendingOrder if self.sort_order == Qt.DescendingOrder else Qt.DescendingOrder
        else:
            self.sort_column = column
            self.sort_order = Qt.AscendingOrder

        if column == 1:
            self.file_paths.sort(key=os.path.basename, reverse=self.sort_order == Qt.DescendingOrder)
        elif column == 2:
            self.file_paths.sort(key=lambda path: fitz.open(path).page_count, reverse=self.sort_order == Qt.DescendingOrder)
        elif column == 3:
            self.file_paths.sort(key=os.path.getsize, reverse=self.sort_order == Qt.DescendingOrder)
        elif column == 4:
            self.file_paths.sort(key=os.path.getmtime, reverse=self.sort_order == Qt.DescendingOrder)

        self.table_widget.sortItems(column, self.sort_order)

        self.table_widget.clearContents()
        self.table_widget.setRowCount(0)
        for path in self.file_paths:
            self.add_table_row(path)
            self.update_sort_indicator()
            self.update_item_numbers()

    def update_sort_indicator(self):
        if self.sort_order == Qt.AscendingOrder:
            self.table_widget.horizontalHeader().setSortIndicator(self.sort_column, Qt.AscendingOrder)
        else:
            self.table_widget.horizontalHeader().setSortIndicator(self.sort_column, Qt.DescendingOrder)

    def update_total_label(self):
        total_files = len(self.file_paths)
        self.total_label.setText(f"PDF文件總數：{total_files}")

    def update_item_numbers(self):
        for row in range(self.table_widget.rowCount()):
            self.table_widget.item(row, 0).setText(str(row + 1))

    def merge_pdfs(self):
        if self.file_paths:
            merged_doc = fitz.open()
            for path in self.file_paths:
                with fitz.open(path) as doc:
                    merged_doc.insert_pdf(doc)
            save_path, _ = QFileDialog.getSaveFileName(self, "保存合併後的PDF文件", "", "PDF files (*.pdf)")
            if save_path:
                merged_doc.save(save_path)
                QMessageBox.information(self, "合併完成", "PDF文件合併完成！")
        else:
            QMessageBox.warning(self, "無文件", "請先添加要合併的PDF文件。")

class ExtractPagesDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Extract Page Options")
        self.setMinimumSize(400, 200)
  
        layout = QVBoxLayout()

        # Create button group for exclusive options
        self.button_group = QButtonGroup()

        # Predefined options group
        predefined_group = QGroupBox("預定義選項")
        predefined_layout = QVBoxLayout()
        self.radio_odd_pages = QRadioButton("提取單數頁面")
        self.radio_even_pages = QRadioButton("提取雙數頁面")
        self.button_group.addButton(self.radio_odd_pages)
        self.button_group.addButton(self.radio_even_pages)
        predefined_layout.addWidget(self.radio_odd_pages)
        predefined_layout.addWidget(self.radio_even_pages)
        predefined_group.setLayout(predefined_layout)
        layout.addWidget(predefined_group)

        # Multiple pages option group
        multiple_group = QGroupBox("提取倍數頁面")
        multiple_layout = QVBoxLayout()
        self.radio_multiple_pages = QRadioButton("提取[N]為倍數的頁面")
        self.multiple_pages_edit = QLineEdit()
        self.multiple_pages_edit.setPlaceholderText("輸入倍數")
        self.multiple_pages_edit.setEnabled(False)
        self.radio_multiple_pages.toggled.connect(self.multiple_pages_edit.setEnabled)
        self.button_group.addButton(self.radio_multiple_pages)
        multiple_layout.addWidget(self.radio_multiple_pages)
        multiple_layout.addWidget(self.multiple_pages_edit)
        multiple_group.setLayout(multiple_layout)
        layout.addWidget(multiple_group)

        # Custom pages option group
        custom_group = QGroupBox("自定義頁面")
        custom_layout = QVBoxLayout()
        self.radio_custom_pages = QRadioButton("自定義頁面")
        self.custom_pages_edit = QLineEdit()
        self.custom_pages_edit.setPlaceholderText("e.g., 1,3,5-7")
        self.custom_pages_edit.setEnabled(False)
        self.radio_custom_pages.toggled.connect(self.custom_pages_edit.setEnabled)
        self.button_group.addButton(self.radio_custom_pages)
        custom_layout.addWidget(self.radio_custom_pages)
        custom_layout.addWidget(self.custom_pages_edit)
        custom_group.setLayout(custom_layout)
        layout.addWidget(custom_group)

        # Complex option group
        complex_group = QGroupBox("複雜選項")
        complex_layout = QVBoxLayout()
        self.radio_complex_option = QRadioButton("從第 N 頁開始，每 X 頁提取 Z 頁")
        self.n_edit = QLineEdit()
        self.n_edit.setPlaceholderText("N: 起始頁")
        self.x_edit = QLineEdit()
        self.x_edit.setPlaceholderText("X: 間隔頁數")
        self.z_edit = QLineEdit()
        self.z_edit.setPlaceholderText("Z: 提取頁數")
        self.n_edit.setEnabled(False)
        self.x_edit.setEnabled(False)
        self.z_edit.setEnabled(False)
        self.radio_complex_option.toggled.connect(self.toggle_complex_option_enabled)
        self.button_group.addButton(self.radio_complex_option)
        complex_layout.addWidget(self.radio_complex_option)
        complex_layout.addWidget(self.n_edit)
        complex_layout.addWidget(self.x_edit)
        complex_layout.addWidget(self.z_edit)
        complex_group.setLayout(complex_layout)
        layout.addWidget(complex_group)

        # Confirm button
        confirm_button = QPushButton("確認")
        confirm_button.clicked.connect(self.accept)
        layout.addWidget(confirm_button, alignment=Qt.AlignRight)

        self.setLayout(layout)

    def toggle_complex_option_enabled(self, checked):
        self.n_edit.setEnabled(checked)
        self.x_edit.setEnabled(checked)
        self.z_edit.setEnabled(checked)

    def get_extract_options(self):
        if self.radio_odd_pages.isChecked():
            return "odd"
        elif self.radio_even_pages.isChecked():
            return "even"
        elif self.radio_multiple_pages.isChecked() and self.multiple_pages_edit.text().isdigit():
            return "multiple", int(self.multiple_pages_edit.text())
        elif self.radio_custom_pages.isChecked():
            return "custom", self.custom_pages_edit.text()
        elif self.radio_complex_option.isChecked():
            try:
                n = int(self.n_edit.text()) - 1  # Convert to zero-indexed
                x = int(self.x_edit.text())
                z = int(self.z_edit.text())
                return "complex", n, x, z
            except ValueError:
                QMessageBox.warning(self, "Input Error", "Please enter valid integers for N, X, and Z.")
                return None
        return None

class PDFConverter(QThread):
    progress_updated = pyqtSignal(int)
    
    def __init__(self, pdf_path, docx_path):
        super().__init__()
        self.pdf_path = pdf_path
        self.docx_path = docx_path

    def run(self):
        try:
            cv = Converter(self.pdf_path)
            cv.convert(self.docx_path, start=0, end=None)
            cv.close()
        except Exception as e:
            print(f"轉換失敗：{e}")

class DeletePagesDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("刪除頁面選項")
        self.setMinimumSize(400, 350)
        
        layout = QVBoxLayout()

        # Create button group for exclusive options
        self.button_group = QButtonGroup()

        # Predefined options group
        predefined_group = QGroupBox("預定義選項")
        predefined_layout = QVBoxLayout()
        self.radio_odd_pages = QRadioButton("刪除單數頁碼")
        self.radio_even_pages = QRadioButton("刪除雙數頁碼")
        self.button_group.addButton(self.radio_odd_pages)
        self.button_group.addButton(self.radio_even_pages)
        predefined_layout.addWidget(self.radio_odd_pages)
        predefined_layout.addWidget(self.radio_even_pages)
        predefined_group.setLayout(predefined_layout)
        layout.addWidget(predefined_group)

        # Multiple pages option group
        multiple_group = QGroupBox("刪除倍數頁碼")
        multiple_layout = QVBoxLayout()
        self.radio_multiple_pages = QRadioButton("刪除[N]倍數頁碼")
        self.multiple_pages_edit = QLineEdit()
        self.multiple_pages_edit.setPlaceholderText("輸入倍數")
        self.multiple_pages_edit.setEnabled(False)
        self.radio_multiple_pages.toggled.connect(self.multiple_pages_edit.setEnabled)
        self.button_group.addButton(self.radio_multiple_pages)
        multiple_layout.addWidget(self.radio_multiple_pages)
        multiple_layout.addWidget(self.multiple_pages_edit)
        multiple_group.setLayout(multiple_layout)
        layout.addWidget(multiple_group)

        # Custom pages option group
        custom_group = QGroupBox("自定義頁碼")
        custom_layout = QVBoxLayout()
        self.radio_custom_pages = QRadioButton("自定義頁碼")
        self.custom_pages_edit = QLineEdit()
        self.custom_pages_edit.setPlaceholderText("例如：1,3,5-7")
        self.custom_pages_edit.setEnabled(False)
        self.radio_custom_pages.toggled.connect(self.custom_pages_edit.setEnabled)
        self.button_group.addButton(self.radio_custom_pages)
        custom_layout.addWidget(self.radio_custom_pages)
        custom_layout.addWidget(self.custom_pages_edit)
        custom_group.setLayout(custom_layout)
        layout.addWidget(custom_group)

        # Complex option group
        complex_group = QGroupBox("複雜選項")
        complex_layout = QVBoxLayout()
        self.radio_complex_option = QRadioButton("從第 N 頁開始，每 X 頁刪除 Z 頁")
        self.n_edit = QLineEdit()
        self.n_edit.setPlaceholderText("N: 起始頁")
        self.x_edit = QLineEdit()
        self.x_edit.setPlaceholderText("X: 間隔頁數")
        self.z_edit = QLineEdit()
        self.z_edit.setPlaceholderText("Z: 刪除頁數")
        self.n_edit.setEnabled(False)
        self.x_edit.setEnabled(False)
        self.z_edit.setEnabled(False)
        self.radio_complex_option.toggled.connect(self.toggle_complex_option_enabled)
        self.button_group.addButton(self.radio_complex_option)
        complex_layout.addWidget(self.radio_complex_option)
        complex_layout.addWidget(self.n_edit)
        complex_layout.addWidget(self.x_edit)
        complex_layout.addWidget(self.z_edit)
        complex_group.setLayout(complex_layout)
        layout.addWidget(complex_group)
        
        # 確認按鈕
        confirm_button = QPushButton("確認")
        confirm_button.clicked.connect(self.accept)
        layout.addWidget(confirm_button, alignment=Qt.AlignRight)

        self.setLayout(layout)

    def toggle_complex_option_enabled(self, checked):
        self.n_edit.setEnabled(checked)
        self.x_edit.setEnabled(checked)
        self.z_edit.setEnabled(checked)

    def get_delete_options(self):
        if self.radio_odd_pages.isChecked():
            return "odd"
        elif self.radio_even_pages.isChecked():
            return "even"
        elif self.radio_multiple_pages.isChecked() and self.multiple_pages_edit.text().isdigit():
            return "multiple", int(self.multiple_pages_edit.text())
        elif self.radio_custom_pages.isChecked():
            return "custom", self.custom_pages_edit.text()
        elif self.radio_complex_option.isChecked():
            try:
                n = int(self.n_edit.text()) - 1  # Convert to zero-indexed
                x = int(self.x_edit.text())
                z = int(self.z_edit.text())
                return "complex", n, x, z
            except ValueError:
                QMessageBox.warning(self, "Input Error", "Please enter valid integers for N, X, and Z.")
                return None
        return None

class TextExtractorDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("提取PDF文本")
        self.setMinimumSize(600, 500)
        self.initUI()

    def mm_to_points(self, mm):
        return mm * 2.83464567  # 每毫米等於2.83464567點

    def initUI(self):

        layout = QVBoxLayout()

        # PDF文件路徑輸入框
        pdf_path_layout = QHBoxLayout()
        pdf_path_label = QLabel("PDF文件:")
        self.pdf_path_edit = QLineEdit()
        pdf_path_button = QPushButton("瀏覽...")
        pdf_path_button.clicked.connect(self.select_pdf_file)
        pdf_path_layout.addWidget(pdf_path_label)
        pdf_path_layout.addWidget(self.pdf_path_edit)
        pdf_path_layout.addWidget(pdf_path_button)
        layout.addLayout(pdf_path_layout)

        # 座標輸入框
        coords_layout = QGridLayout()
        coords_layout.addWidget(QLabel("x1:"), 0, 0)
        self.x1_edit = QLineEdit()
        coords_layout.addWidget(self.x1_edit, 0, 1)
        coords_layout.addWidget(QLabel("y1:"), 0, 2)
        self.y1_edit = QLineEdit()
        coords_layout.addWidget(self.y1_edit, 0, 3)
        coords_layout.addWidget(QLabel("x2:"), 1, 0)
        self.x2_edit = QLineEdit()
        coords_layout.addWidget(self.x2_edit, 1, 1)
        coords_layout.addWidget(QLabel("y2:"), 1, 2)
        self.y2_edit = QLineEdit()
        coords_layout.addWidget(self.y2_edit, 1, 3)
        layout.addLayout(coords_layout)

        # 頁碼輸入框
        page_layout = QHBoxLayout()
        page_label = QLabel("預覽頁碼:")
        page_label.setAlignment(Qt.AlignRight)
        page_layout.addWidget(page_label)

        self.prev_page_button = QPushButton("<")  # 添加“<”按钮
        self.prev_page_button.clicked.connect(self.show_prev_page)
        self.prev_page_button.setFixedWidth(30)
        page_layout.addWidget(self.prev_page_button)

        self.page_edit = QLineEdit()
        self.page_edit.setFixedWidth(50)
        self.page_edit.setAlignment(Qt.AlignCenter)  # 設置文本對齊方式為置中
        self.page_edit.setText("1")  # 設置預設頁碼為1
        page_layout.addWidget(self.page_edit)

        self.next_page_button = QPushButton(">")  # 添加“>”按钮
        self.next_page_button.clicked.connect(self.show_next_page)
        self.next_page_button.setFixedWidth(30)
        page_layout.addWidget(self.next_page_button)
        layout.addLayout(page_layout)

        # 預覽按鈕
        self.preview_button = QPushButton("預覽")
        self.preview_button.clicked.connect(self.update_preview)
        layout.addWidget(self.preview_button)

        # 預覽區域
        self.preview_label = QLabel()
        self.preview_label.setAlignment(Qt.AlignCenter)
        self.scroll_area = QScrollArea()
        self.scroll_area.setMinimumSize(400, 200)
        self.scroll_area.setWidgetResizable(True)  # 允許內容根據視窗大小調整
        self.scroll_area.setWidget(self.preview_label)  # 將預覽標籤添加到滾動區域
        
        layout.addWidget(self.scroll_area)  # 在佈局中添加滾動區域而不是僅僅添加預覽標籤
        

        # 提取按鈕
        self.extract_button = QPushButton("Extract to Excel")
        self.extract_button.clicked.connect(self.extract_to_excel)
        layout.addWidget(self.extract_button)

        self.setLayout(layout)


    def show_prev_page(self):
        try:
            current_page = int(self.page_edit.text())
            if current_page > 1:
                self.page_edit.setText(str(current_page - 1))
                self.update_preview()  # 更新预览
        except ValueError:
            pass  # 忽略无效输入

    def show_next_page(self):
        try:
            current_page = int(self.page_edit.text())
            pdf_path = self.pdf_path_edit.text()
            if pdf_path and current_page < fitz.open(pdf_path).page_count:
                self.page_edit.setText(str(current_page + 1))
                self.update_preview()  # 更新预览
        except ValueError:
            pass  # 忽略无效输入

    def select_pdf_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "選擇PDF文件", "", "PDF files (*.pdf)")
        if file_path:
            self.pdf_path_edit.setText(file_path)

    def update_preview(self):
        # 獲取選定的行
        selected_items = self.result_table.selectedItems()
        if not selected_items:
            return
        
        # 獲取所選行
        row = selected_items[0].row()
        
        # 更新按鈕狀態
        self.open_selected_button.setEnabled(True)
        
        # 獲取對應的結果
        if row < len(self.search_results):
            result = self.search_results[row]
            
            # 構建預覽內容 - 適配深色模式
            preview_text = f"<div style='color: white;'>"
            preview_text += f"<h3 style='color: #4a86e8;'>文件: {result['file_name']}</h3>"
            preview_text += f"<p><b>路徑:</b> {result['file_path']}</p>"
            preview_text += f"<p><b>匹配頁碼:</b> {', '.join(str(p) for p in result['pages'])}</p>"
            preview_text += f"<p><b>匹配數量:</b> {len(result['pages'])}</p>"
            preview_text += "<h4>匹配內容:</h4><hr style='border: 1px solid #444444;'>"
            
            # 添加每個匹配片段，突出顯示關鍵詞
            for i, snippet in enumerate(result["snippets"]):
                # 獲取對應的關鍵詞（如果有的話）
                keyword = result.get("keywords", [""])[i] if i < len(result.get("keywords", [])) else ""
                
                # HTML格式高亮顯示關鍵詞
                highlighted_snippet = snippet
                if keyword:
                    highlighted_snippet = snippet.replace(
                        f"[{keyword}]", 
                        f"<span style='background-color: #ffd700; color: black; font-weight: bold;'>{keyword}</span>"
                    )
                
                preview_text += f"<p><b style='color: #4a86e8;'>第 {result['pages'][i]} 頁 "
                if keyword:
                    preview_text += f"(關鍵詞: {keyword})"
                preview_text += f":</b><br>{highlighted_snippet}</p>"
                
                if i < len(result["snippets"]) - 1:
                    preview_text += "<hr style='border: 1px solid #444444;'>"
            
            preview_text += "</div>"
            
            # 設置預覽文本 (使用HTML格式)
            self.preview_text.setHtml(preview_text)

    def extract_to_excel(self):
        pdf_path = self.pdf_path_edit.text()
        coords = (
            self.mm_to_points(float(self.x1_edit.text())),
            self.mm_to_points(float(self.y1_edit.text())),
            self.mm_to_points(float(self.x2_edit.text())),
            self.mm_to_points(float(self.y2_edit.text()))
        )

        try:
            doc = fitz.open(pdf_path)
            excel_path = os.path.splitext(pdf_path)[0] + "_extracted.xlsx"
            workbook = Workbook()
            sheet = workbook.active

            for page_num in range(doc.page_count):
                page = doc[page_num]
                text_blocks = page.get_text("blocks", clip=coords)
                sorted_blocks = sorted(text_blocks, key=lambda b: b[1])  # 根据y1坐标排序
                
                col_counter = 1  # 初始列计数器
                
                for block in sorted_blocks:
                    block_text = block[4].strip()  # 直接取得文本部分
                    lines = block_text.splitlines()  # 按行分割文本
                    
                    for line in lines:
                        if line.strip():  # 只处理非空行
                            cell_id = '{}{}'.format(get_column_letter(col_counter), page_num + 1)
                            sheet[cell_id] = line
                            col_counter += 1  # 更新列计数器至下一个字母

                # 为下一页重置列计数器，如果每页独立
                col_counter = 1

            workbook.save(excel_path)
            QMessageBox.information(self, "成功", f"提取的文本已保存到：{excel_path}")
            doc.close()
        except Exception as e:
            QMessageBox.critical(self, "錯誤", f"提取文本失敗：{e}")

class InsertPagesDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("插入PDF頁面")
        self.setMinimumSize(400, 200)
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout()
        self.tab_widget = QTabWidget()
        layout.addWidget(self.tab_widget)

        self.single_tab = QWidget()
        self.setup_single_tab()
        self.tab_widget.addTab(self.single_tab, "SINGLE")

        self.repeat_tab = QWidget()
        self.setup_repeat_tab()
        self.tab_widget.addTab(self.repeat_tab, "REPEAT")

        self.setLayout(layout)

    def setup_single_tab(self):
        layout = QVBoxLayout()

        self.pdf_path_label = QLabel("選擇PDF文件:")
        self.pdf_path_edit_single = QLineEdit()
        self.pdf_path_edit_single.setReadOnly(True)
        self.browse_button_single = QPushButton("瀏覽...")
        self.browse_button_single.clicked.connect(self.browse_pdf)

        browse_layout = QVBoxLayout()
        browse_layout.addWidget(self.pdf_path_label)
        browse_layout.addWidget(self.pdf_path_edit_single)
        browse_layout.addWidget(self.browse_button_single)
        layout.addLayout(browse_layout)

        self.page_range_label = QLabel("插入頁碼範圍 (例如: 1,3,5-7):")
        self.page_range_edit_single = QLineEdit()

        pagerange_layout = QVBoxLayout()
        pagerange_layout.addWidget(self.page_range_label)
        pagerange_layout.addWidget(self.page_range_edit_single)
        layout.addLayout(pagerange_layout)

        self.insert_pos_label = QLabel("插入位置:")
        self.insert_pos_combo = QComboBox()
        self.insert_pos_combo.addItems(["開頭", "結尾", "指定位置"])
        self.insert_page_edit = QLineEdit()
        self.insert_page_edit.setPlaceholderText("輸入頁碼")
        self.insert_page_edit.setEnabled(False)
        self.insert_pos_combo.currentIndexChanged.connect(self.toggle_page_input)

        insert_layout = QHBoxLayout()
        insert_layout.addWidget(self.insert_pos_label)
        insert_layout.addWidget(self.insert_pos_combo)
        insert_layout.addWidget(self.insert_page_edit)
        layout.addLayout(insert_layout)

        self.button_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        self.button_box.accepted.connect(self.accept)
        self.button_box.rejected.connect(self.reject)
        layout.addWidget(self.button_box)

        self.single_tab.setLayout(layout)

    def setup_repeat_tab(self):
        layout = QVBoxLayout(self.repeat_tab)

        self.pdf_path_label = QLabel("選擇PDF文件:")
        self.pdf_path_edit_repeat = QLineEdit()
        browse_button = QPushButton("瀏覽...")
        browse_button.clicked.connect(self.browse_pdf)

        page_label = QLabel("頁碼 (例如: 1,3,5-7):")
        self.page_range_edit_repeat = QLineEdit()

        interval_label = QLabel("每 N 頁插入一次:")
        self.interval_spin = QSpinBox()
        self.interval_spin.setMinimum(1)
        self.interval_spin.setMaximum(99999)  # 设置 QSpinBox 的最大值为 99999

        execute_button = QPushButton("執行")
        execute_button.clicked.connect(self.perform_insertion)

        layout.addWidget(self.pdf_path_label)
        layout.addWidget(self.pdf_path_edit_repeat)
        layout.addWidget(browse_button)
        layout.addWidget(page_label)
        layout.addWidget(self.page_range_edit_repeat)
        layout.addWidget(interval_label)
        layout.addWidget(self.interval_spin)
        layout.addWidget(execute_button)

    def browse_pdf(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "選擇要插入的PDF文件", "", "PDF files (*.pdf)")
        if file_path:
            # Set the file path to both single and repeat tabs' QLineEdit
            self.pdf_path_edit_single.setText(file_path)
            self.pdf_path_edit_repeat.setText(file_path)

    def perform_insertion(self):
        pdf_path = self.pdf_path_edit_repeat.text()
        page_range = self.page_range_edit_repeat.text()
        interval = self.interval_spin.value()
        if hasattr(self.parent(), 'insert_pages_repeat'):
            self.parent().insert_pages_repeat(pdf_path, page_range, interval)
        else:
            QMessageBox.warning(self, "警告", "父窗口未實現insert_pages_repeat方法。")

    def toggle_page_input(self):
        self.insert_page_edit.setEnabled(self.insert_pos_combo.currentText() == "指定位置")

    def get_insert_details(self):
        return {
            'pdf_path': self.pdf_path_edit_single.text(),
            'page_range': self.page_range_edit_single.text(),
            'insert_position': self.insert_pos_combo.currentText(),
            'insert_page': self.insert_page_edit.text()
        }
    
class PrintOptionsWidget(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.initUI()
        self.populate_printers()  # Populate the printer combo box

    def initUI(self):
        layout = QVBoxLayout()
        layout.setSpacing(10)

        # Setting up Printer Selection
        self.setupPrinterSelection(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 设置纸张大小选项
        self.setupPaperSize(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 设置打印方向选项
        self.setupOrientation(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 设置偏移量选项
        self.setupOffsets(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 设置打印份数选项
        self.setupCopies(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 设置打印页数选项
        self.setupPageRange(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 添加自定义打印比例选项
        self.setupScaleFactor(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 设置双面打印选项
        self.setupDuplex(layout)
        layout.addStretch(1)  # 添加弹性空间

        # 设置全页打印选项（无边距）
        self.setupFitToMargin(layout)
        layout.addStretch(5)  # 添加弹性空间

        # 设置打印按钮
        self.setupPrintButton(layout)
        layout.addStretch(1)

        # 应用布局
        self.setLayout(layout)

    def setupPrinterSelection(self, layout):
        group_box = QGroupBox("Printer Selection")
        group_layout = QVBoxLayout()
        printer_label = QLabel("Select Printer:")
        self.printer_combo = QComboBox()
        group_layout.addWidget(printer_label)
        group_layout.addWidget(self.printer_combo)

        # 添加打印机偏好设置按钮
        self.printer_preferences_button = QPushButton("Printer Preferences")
        self.printer_preferences_button.clicked.connect(self.open_printer_preferences)
        group_layout.addWidget(self.printer_preferences_button)

        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def populate_printers(self):
        available_printers = QPrinterInfo.availablePrinters()
        default_printer = QPrinterInfo.defaultPrinter()

        # 添加打印機到下拉選單並標記預設打印機
        for printer in available_printers:
            self.printer_combo.addItem(printer.printerName())
        
        # 設置預設打印機
        index = self.printer_combo.findText(default_printer.printerName())
        if index >= 0:
            self.printer_combo.setCurrentIndex(index)

    def setupPaperSize(self, layout):
        group_box = QGroupBox("Paper Options")
        group_layout = QVBoxLayout()
        paper_size_label = QLabel("Paper Size:")
        self.paper_size_combo = QComboBox()
        self.paper_size_combo.addItems(["A4", "A3", "A5", "Letter"])
        group_layout.addWidget(paper_size_label)
        group_layout.addWidget(self.paper_size_combo)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupOrientation(self, layout):
        group_box = QGroupBox("Orientation")
        group_layout = QVBoxLayout()
        orientation_label = QLabel("Orientation:")
        self.portrait_radio = QRadioButton("Portrait")
        self.landscape_radio = QRadioButton("Landscape")
        self.auto_orientation_radio = QRadioButton("Auto")  # 新增自动选项
        self.auto_orientation_radio.setChecked(True)  # 默认选中自动
        orientation_layout = QHBoxLayout()
        orientation_layout.addWidget(self.portrait_radio)
        orientation_layout.addWidget(self.landscape_radio)
        orientation_layout.addWidget(self.auto_orientation_radio)
        group_layout.addWidget(orientation_label)
        group_layout.addLayout(orientation_layout)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def set_orientation_based_on_pdf(self, pdf_doc):
        if self.auto_orientation_radio.isChecked():
            # 获取第一页的尺寸
            first_page = pdf_doc[0]
            width, height = first_page.rect.width, first_page.rect.height
            
            if width > height:
                self.landscape_radio.setChecked(True)
            else:
                self.portrait_radio.setChecked(True)

    def get_orientation(self):
        if self.auto_orientation_radio.isChecked():
            # 如果选择了自动，返回 None，稍后再根据 PDF 尺寸决定
            return None
        elif self.landscape_radio.isChecked():
            return QPrinter.Landscape
        else:
            return QPrinter.Portrait


    def setupOffsets(self, layout):
        group_box = QGroupBox("Margin Shift")
        group_layout = QVBoxLayout()
        offsets_label = QLabel("Margin Shift (mm):")
        offsets_layout = QGridLayout()
        self.top_offset_spin = QDoubleSpinBox()
        self.bottom_offset_spin = QDoubleSpinBox()
        self.left_offset_spin = QDoubleSpinBox()
        self.right_offset_spin = QDoubleSpinBox()

        # 允许负值
        self.top_offset_spin.setMinimum(-9999)
        self.bottom_offset_spin.setMinimum(-9999)
        self.left_offset_spin.setMinimum(-9999)
        self.right_offset_spin.setMinimum(-9999)

        offsets_layout.addWidget(QLabel("Top:"), 0, 0)
        offsets_layout.addWidget(self.top_offset_spin, 0, 1)
        offsets_layout.addWidget(QLabel("Bottom:"), 1, 0)
        offsets_layout.addWidget(self.bottom_offset_spin, 1, 1)
        offsets_layout.addWidget(QLabel("Left:"), 0, 2)
        offsets_layout.addWidget(self.left_offset_spin, 0, 3)
        offsets_layout.addWidget(QLabel("Right:"), 1, 2)
        offsets_layout.addWidget(self.right_offset_spin, 1, 3)
        group_layout.addWidget(offsets_label)
        group_layout.addLayout(offsets_layout)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupCopies(self, layout):
        group_box = QGroupBox("Number of Copies")
        group_layout = QVBoxLayout()
        copies_label = QLabel("Copies:")
        self.copies_spin = QSpinBox()
        self.copies_spin.setMinimum(1)
        group_layout.addWidget(copies_label)
        group_layout.addWidget(self.copies_spin)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupFitToMargin(self, layout):
        self.fit_to_margin_checkbox = QCheckBox("Fit to Printer Margin")
        layout.addWidget(self.fit_to_margin_checkbox)

    def setupDuplex(self, layout):
        self.duplex_checkbox = QCheckBox("Enable Duplex Printing")
        layout.addWidget(self.duplex_checkbox)

    def setupPageRange(self, layout):
        group_box = QGroupBox("Page Range")
        group_layout = QVBoxLayout()
        page_range_label = QLabel("Pages (e.g., 1,3,5-7):")
        self.page_range_edit = QLineEdit()
        group_layout.addWidget(page_range_label)
        group_layout.addWidget(self.page_range_edit)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupScaleFactor(self, layout):
        group_box = QGroupBox("Scale Factor")
        group_layout = QVBoxLayout()
        scale_factor_label = QLabel("Scale (%):")
        self.scale_factor_spin = QSpinBox()
        self.scale_factor_spin.setRange(10, 400)  # Allow scaling from 10% to 400%
        self.scale_factor_spin.setValue(100)  # Default to 100%
        group_layout.addWidget(scale_factor_label)
        group_layout.addWidget(self.scale_factor_spin)
        group_box.setLayout(group_layout)
        layout.addWidget(group_box)

    def setupPrintButton(self, layout):
        self.print_button = QPushButton("Print")
        layout.addWidget(self.print_button)

    def open_printer_preferences(self):
        selected_printer_name = self.printer_combo.currentText()
        if selected_printer_name:
            printer = QPrinter()
            printer.setPrinterName(selected_printer_name)
            dialog = QPrintDialog(printer, self)
            dialog.exec_()

class PDFViewer(QMainWindow):
    ZOOM_IN_FACTOR = 1.25  # 定义放大因子
    ZOOM_OUT_FACTOR = 1 / ZOOM_IN_FACTOR  # 定义缩小因子
    def __init__(self):
        super().__init__()
        self.config_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'config.json')
        self.settings = self.load_settings()
        self.doc = None
        self.sort_widget = None
        self.current_page = 0
        self.zoom_ratio = 1.0
        self.min_zoom = 0.25
        self.max_zoom = 4.0
        self.original_pdf_path = None  # 原始PDF文件路徑
        self.temp_pdf_path = None  # 臨時PDF文件路徑
        self.setAcceptDrops(True)  # 啟用拖放功能
        self.search_widget = None
        self.search_results = []  
        self.open_windows = []  # Keep track of open windows 
        self.is_modified = False 
        self.initUI()
        self.apply_settings()


    def show_sort_widget(self):
        if not self.doc:
            self.show_warning("請先打開一個PDF文件。")
            return

        if self.sort_widget is None:
            # SortWidget 尚未创建，创建并显示
            self.sort_widget = SortWidget(self)
            self.sort_widget.sort_applied.connect(self.apply_page_sorting)
            self.addDockWidget(Qt.RightDockWidgetArea, self.sort_widget)

        # 显示 SortWidget
        self.sort_widget.update_page_list(self.doc.page_count)
        self.sort_widget.show()

    def apply_page_sorting(self, order):
        if self.doc:
            self.doc.select(order)
            self.show_page(0)  # Show the first page after sorting
            self.show_success("Page order has been applied!")
        else:
            self.show_warning("Please open a PDF file first.")
        self.is_modified = True

    def merge_pdfs(self):
        merge_dialog = MergePDFsDialog(self)
        merge_dialog.exec_()  # 显示合并对话框

    def update_file_info(self, path):
        # 設置文件信息標籤
        if self.doc:
            file_info = f"文件路徑: {path} | 頁數: {self.doc.page_count}"
            self.file_info_label.setText(file_info)
        else:
            self.file_info_label.setText("未加載文件")


    def open_text_extractor(self):
        extractor_dialog = TextExtractorDialog(self)
        extractor_dialog.exec_()

    def contextMenuEvent(self, event):
        local_pos = self.image_widget.mapFromGlobal(event.globalPos())

        # 檢查是否點擊在image_label內部
        if self.image_label.geometry().contains(local_pos):
            menu = QMenu(self)
            menu.setStyleSheet("""
                QMenu {
                    background-color: #272727; /* 深灰色背景 */
                    color: white; /* 白色文字 */
                    border: 1px solid #666666; /* 浅灰色边框 */
                    border-radius: 10px; /* 菜单边框圆角 */
                }
                QMenu::item {
                    padding: 5px 20px; /* 菜单项内边距 */
                    margin: 2px 2px; /* 菜单项外边距 */
                }
                QMenu::item:selected {
                    background-color: #02F78E; /* 选中时的背景色 */
                    color: black; /* 选中时的文字颜色 */
                    border-radius: 5px; /* 选中项的边框圆角 */
                    padding: 5px 20px; /* 选中项的内边距 */
                }
            """)

            insert_action = menu.addAction('Insert Page')
            insert_action.triggered.connect(self.insert_pages_current)

            delete_action = menu.addAction("Delete Page")
            delete_action.triggered.connect(self.delete_pages_current)

            extract_action = menu.addAction("Extract Page")
            extract_action.triggered.connect(self.extract_pages_current)

            rotate_action = menu.addAction("Rotate")
            rotate_action.triggered.connect(self.rotate_pages_current)

            Search_action = menu.addAction("Search")
            Search_action.triggered.connect(self.show_search_widget)

            pdf_info_action = menu.addAction("PDF Info")
            pdf_info_action.triggered.connect(self.show_pdf_info)

            menu.exec_(event.globalPos())

    def insert_pages_current(self):
        if not self.doc:
            self.show_warning("請先打開一個PDF文件。")
            return

        insert_pdf_path, _ = QFileDialog.getOpenFileName(self, "選擇要插入的PDF文件", "", "PDF files (*.pdf)")
        if insert_pdf_path:
            try:
                page_range, ok = QInputDialog.getText(self, "頁碼範圍", "輸入要插入的頁碼範圍 (例如: 1,3,5-7):")
                if ok:
                    pages_to_insert = self.parse_page_range(page_range)
                    with fitz.open(insert_pdf_path) as insert_doc:
                        self.doc.insert_pdf(insert_doc, from_page=pages_to_insert[0], to_page=pages_to_insert[-1], start_at=self.current_page)
                    self.show_page(self.current_page)
                    self.show_success("頁面插入成功！")
            except Exception as e:
                self.show_error(str(e))
        self.is_modified = True

    def delete_pages_current(self):
        if not self.doc:
            self.show_warning("請先打開一個PDF文件。")
            return

        try:
            self.doc.delete_page(self.current_page)
            self.current_page = max(0, self.current_page - 1)  # Adjust current page index
            self.show_page(self.current_page)
            self.show_success("頁面已刪除！")
        except Exception as e:
            self.show_error(str(e))
        self.is_modified = True

    def extract_pages_current(self):
        if not self.doc:
            self.show_warning("請先打開一個PDF文件。")
            return

        extracted_doc = fitz.open()
        extracted_doc.insert_pdf(self.doc, from_page=self.current_page, to_page=self.current_page)
        save_path, _ = QFileDialog.getSaveFileName(self, "保存提取的頁面", "", "PDF files (*.pdf)")
        if save_path:
            extracted_doc.save(save_path)
            self.show_success("頁面已提取并保存！")
        self.is_modified = True

    def rotate_pages_current(self):
        if not self.doc:
            self.show_warning("請先打開一個PDF文件。")
            return

        angle, ok = QInputDialog.getInt(self, "旋轉角度", "輸入旋轉角度 (90, 180, 270):", min=0, max=360, step=90)
        if ok:
            try:
                page = self.doc.load_page(self.current_page)
                page.set_rotation(angle)
                self.show_page(self.current_page)
                self.show_success("頁面已旋轉！")
            except Exception as e:
                self.show_error(str(e))
        self.is_modified = True    
        
    def initUI(self):
        self.setWindowTitle('PDFDocuEdit Pro - 無檔案')
        
        # 獲取屏幕大小
        screen = QDesktopWidget().screenNumber(QDesktopWidget().cursor().pos())
        screen_size = QDesktopWidget().screenGeometry(screen)
        
        # 設置窗口大小為屏幕大小的75%
        width = int(screen_size.width() * 0.75)
        height = int(screen_size.height() * 0.75)
        
        # 居中顯示
        x = (screen_size.width() - width) // 2
        y = (screen_size.height() - height) // 2
        
        self.setGeometry(x, y, width, height)

        # 設置最小大小，以確保UI元素能夠正常顯示
        self.setMinimumSize(800, 600)

        shortcut_first_page = QShortcut(QKeySequence("Ctrl+Home"), self)
        shortcut_first_page.activated.connect(self.goto_first_page)

        shortcut_last_page = QShortcut(QKeySequence("Ctrl+End"), self)
        shortcut_last_page.activated.connect(self.goto_last_page)

        shortcut_prev_page = QShortcut(QKeySequence("Page Up"), self)
        shortcut_prev_page.activated.connect(self.show_prev_page)

        shortcut_next_page = QShortcut(QKeySequence("Page Down"), self)
        shortcut_next_page.activated.connect(self.show_next_page)

        search_shortcut = QShortcut(QKeySequence("Ctrl+F"), self)
        search_shortcut.activated.connect(self.show_search_widget)

        close_shortcut = QShortcut(QKeySequence("Ctrl+W"), self)
        close_shortcut.activated.connect(self.close)

        zoom_in_shortcut = QShortcut(QKeySequence("Ctrl++"), self)
        zoom_in_shortcut.activated.connect(self.zoom_in)

        zoom_out_shortcut = QShortcut(QKeySequence("Ctrl+-"), self)
        zoom_out_shortcut.activated.connect(self.zoom_out)

        # 旋轉 PDF 向左（逆時針）
        rotate_left_shortcut = QShortcut(QKeySequence("Ctrl+{"), self)
        rotate_left_shortcut.activated.connect(self.rotate_pdf_left)

        # 旋轉 PDF 向右（順時針）
        rotate_right_shortcut = QShortcut(QKeySequence("Ctrl+}"), self)
        rotate_right_shortcut.activated.connect(self.rotate_pdf_right)

        # Create shortcuts for left and right arrow keys
        self.shortcut_left = QShortcut(QKeySequence(Qt.Key_Left), self)
        self.shortcut_right = QShortcut(QKeySequence(Qt.Key_Right), self)

        # Connect the shortcuts to their respective handlers
        self.shortcut_left.activated.connect(self.show_prev_page)
        self.shortcut_right.activated.connect(self.show_next_page)

        # Create actions
        open_action = QAction('&Open', self)
        open_action.triggered.connect(self.open_pdf)
        open_action.setShortcut(QKeySequence("Ctrl+O"))

        # Add the new action for opening a new file
        open_new_action = QAction('Open a File in New Window', self)
        open_new_action.triggered.connect(self.open_new_pdf)  # Connect to a new method
        open_new_action.setShortcut(QKeySequence("Ctrl+N"))
        
        save_action = QAction('&Save', self)
        save_action.triggered.connect(self.save_file)
        save_action.setShortcut(QKeySequence("Ctrl+S"))
        
        save_as_action = QAction('Save &As', self)
        save_as_action.triggered.connect(self.save_as_file)

        encrypt_action = QAction('&Encrypt PDF', self)
        encrypt_action.triggered.connect(self.on_encrypt_pdf_triggered)

        decrypt_action = QAction('&Decrypt PDF', self)
        decrypt_action.triggered.connect(self.on_decrypt_pdf_triggered)

        print_action = QAction('&Print', self)
        print_action.triggered.connect(self.show_print_options)
        print_action.setShortcut(QKeySequence("Ctrl+P"))

        exit_action = QAction("&Exit", self)
        exit_action.triggered.connect(self.close)  # 连接到关闭窗口的函数

        # Create the window management action
        window_management_action = QAction('Window Management', self)
        window_management_action.triggered.connect(self.show_window_list)

        # Create a menu for window management
        window_menu = QMenu("&Window", self)
        window_menu.addAction(window_management_action)

        # Create a button for the window menu
        window_button = QToolButton(self)
        window_button.setText("&Window")
        window_button.setMenu(window_menu)
        window_button.setPopupMode(QToolButton.InstantPopup)

        # Create Print Options Dock Widget
        self.print_options_dock = QDockWidget("Print Options", self)
        self.print_options_widget = PrintOptionsWidget()
        self.print_options_dock.setWidget(self.print_options_widget)
        self.print_options_dock.setVisible(False)  # Set visibility to Hidden

        # Connect Print Button
        self.print_options_widget.print_button.clicked.connect(self.start_printing)

        # Create Split Options Dock Widget
        self.split_options_dock = QDockWidget("Split Options", self)
        self.split_options_widget = SplitOptionsWidget()
        self.split_options_dock.setWidget(self.split_options_widget)
        self.split_options_dock.setVisible(False)

        # Connect Split Button
        self.split_options_widget.split_button.clicked.connect(self.start_splitting)


        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        layout = QVBoxLayout(central_widget)
       

        file_menu = QMenu("&File", self)
        file_menu.addAction(open_action)
        file_menu.addAction(open_new_action)  # Add the new action to the file menu
        file_menu.addAction(save_action)
        file_menu.addAction(save_as_action)
        file_menu.addSeparator()  # 添加分隔线
        file_menu.addAction(encrypt_action)
        file_menu.addAction(decrypt_action)
        file_menu.addSeparator()  # 添加分隔线
        file_menu.addAction(print_action)
        file_menu.addSeparator()  # 添加分隔线
        file_menu.addAction(exit_action)  # 添加 Exit 操作到菜单

        file_button = QToolButton(self)
        file_button.setText("&File")
        file_button.setMenu(file_menu)
        file_button.setPopupMode(QToolButton.InstantPopup)  # 设置为点击时立即弹出菜单
    
        toolbar = QToolBar("Main", self)
        toolbar.addWidget(file_button)
        self.addToolBar(toolbar)


        # 創建滾動區域
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)

        # 創建一個 widget 來容納所有內容
        content_widget = QWidget()
        content_layout = QVBoxLayout(content_widget)

        # 創建左側按鈕區域的 Widget 和 Layout
        left_button_widget = QWidget()
        self.left_button_layout = QVBoxLayout(left_button_widget)

        # 創建 QGroupBox 分組
        page_operations_group = QGroupBox("Page Operations")
        page_operations_layout = QVBoxLayout()
        page_operations_group.setLayout(page_operations_layout)

        self.btn_search_pages = self.setup_button('Search Pages', self.show_search_widget, resource_path('App_icon/Search.png'))
        self.btn_insert_pages = self.setup_button('Insert Pages', self.insert_pages, resource_path('App_icon/insert.png'))
        self.btn_delete_pages = self.setup_button('Delete Pages', self.delete_pages, resource_path('App_icon/delete.png'))
        self.btn_extract_pages = self.setup_button('Extract Pages', self.extract_pages, resource_path('App_icon/Extract-page.png'))
        self.btn_sort_pages = self.setup_button('Ordering Pages', self.show_sort_widget, resource_path('App_icon/sorting.png'))
        self.btn_split = self.setup_button('Split PDF', self.show_split_options, resource_path('App_icon/split.png'))

        page_operations_layout.addWidget(self.btn_search_pages)
        page_operations_layout.addWidget(self.btn_insert_pages)
        page_operations_layout.addWidget(self.btn_delete_pages)
        page_operations_layout.addWidget(self.btn_extract_pages)
        page_operations_layout.addWidget(self.btn_sort_pages)
        page_operations_layout.addWidget(self.btn_split)

        self.left_button_layout.addWidget(page_operations_group)

        conversion_group = QGroupBox("Conversion")
        conversion_layout = QVBoxLayout()
        conversion_group.setLayout(conversion_layout)

        self.btn_convert = self.setup_button('Pdf to Word', self.convert_to_word, resource_path('App_icon/PDF-to-Word.png'))
        self.btn_batch_convert = self.setup_button('Batch Convert MS Office to PDF', self.open_batch_convert_dialog, resource_path('App_icon/doc_to_PDF.png'))
        self.btn_batch_convert_txt = self.setup_button('Batch Convert TXT to PDF', self.open_batch_TXT_convert_dialog, resource_path('App_icon/txt_to_PDF.png'))

        conversion_layout.addWidget(self.btn_convert)
        conversion_layout.addWidget(self.btn_batch_convert)
        conversion_layout.addWidget(self.btn_batch_convert_txt)

        self.left_button_layout.addWidget(conversion_group)

        utilities_group = QGroupBox("Utilities")
        utilities_layout = QVBoxLayout()
        utilities_group.setLayout(utilities_layout)

        self.btn_page_count = self.setup_button('Statistical Page Count', self.count_pdf_pages, resource_path('App_icon/PDF-report.png'))
        self.btn_extract_text = self.setup_button('Extract Text By Position', self.open_text_extractor, resource_path('App_icon/Extract-text.png'))
        self.btn_merge_pdfs = self.setup_button('Merge Multiple PDF Files', self.merge_pdfs, resource_path('App_icon/Merge-PDF.png'))
        self.btn_overlay_pdf = self.setup_button('PDF Form Overlay', self.open_overlay_dialog, resource_path('App_icon/overlay.png'))
        self.btn_batch_print = self.setup_button('Batch Print PDF Files', self.open_batch_print_dialog, resource_path('App_icon/batch_print.png'))
        self.btn_deep_search = self.setup_button('Deep Search PDFs', self.open_deep_search_dialog, resource_path('App_icon/deep_search.png'))
        self.btn_merge_csv_excel = self.setup_button('Merge CSV/Excel Files', self.open_merge_csv_excel_dialog, resource_path('App_icon/merge_csv_excel.png'))

        utilities_layout.addWidget(self.btn_page_count)
        utilities_layout.addWidget(self.btn_extract_text)
        utilities_layout.addWidget(self.btn_merge_pdfs)
        utilities_layout.addWidget(self.btn_overlay_pdf)
        utilities_layout.addWidget(self.btn_batch_print)
        utilities_layout.addWidget(self.btn_deep_search)
        utilities_layout.addWidget(self.btn_merge_csv_excel)

        self.left_button_layout.addWidget(utilities_group)

        security_group = QGroupBox("Security")
        security_layout = QVBoxLayout()
        security_group.setLayout(security_layout)

        self.btn_encrypt_pdf = self.setup_button('Encrypt PDF', self.on_encrypt_pdf_triggered, resource_path('App_icon/encrypt.png'))
        self.btn_decrypt_pdf = self.setup_button('Decrypt PDF', self.on_decrypt_pdf_triggered, resource_path('App_icon/decrypt.png'))

        security_layout.addWidget(self.btn_encrypt_pdf)
        security_layout.addWidget(self.btn_decrypt_pdf)

        self.left_button_layout.addWidget(security_group)
        self.left_button_layout.addStretch(10)

        # 將 left_button_widget 添加到 content layout
        content_layout.addWidget(left_button_widget)

        # 將 content widget 設置為滾動區域的 widget
        scroll_area.setWidget(content_widget)

        # 創建並設置 Dock Widget
        self.menu_dock = QDockWidget("Menu", self)
        self.menu_dock.setWidget(scroll_area)  # 將滾動區域設置為 dock 的 widget
        self.menu_dock.setAllowedAreas(Qt.LeftDockWidgetArea | Qt.RightDockWidgetArea)
        self.addDockWidget(Qt.LeftDockWidgetArea, self.menu_dock)
            
        # 設置 Dock Widget 為關閉狀態
        self.menu_dock.hide()

        insert_pages_action = QAction('Insert page', self)
        insert_pages_action.triggered.connect(self.btn_insert_pages.click)

        delete_pages_action = QAction('Delete page', self)
        delete_pages_action.triggered.connect(self.btn_delete_pages.click)

        extract_pages_action = QAction('Extract page', self)
        extract_pages_action.triggered.connect(self.btn_extract_pages.click)

        Split_PDF_action = QAction('Split PDF', self)
        Split_PDF_action.triggered.connect(self.btn_split.click)

        sort_pages_action = QAction('Ordering page sequence', self)
        sort_pages_action.triggered.connect(self.btn_sort_pages.click)

        Search_pages_action = QAction('Search', self)
        Search_pages_action.triggered.connect(self.btn_search_pages.click)

        main_menu_icon = QIcon(resource_path('App_icon/Main_menu.png'))  # 替換為你的圖標文件路徑
        main_menu_action = QAction(main_menu_icon, 'Main Menu', self)
        main_menu_action.triggered.connect(self.toggle_menu_dock)


        # 创建 Edit 按钮和菜单
        edit_menu = QMenu("&Edit", self)
        edit_menu.addAction(Search_pages_action)
        edit_menu.addSeparator()  # 添加分隔线
        edit_menu.addAction(insert_pages_action)
        edit_menu.addAction(delete_pages_action)
        edit_menu.addAction(extract_pages_action)
        edit_menu.addAction(Split_PDF_action)
        edit_menu.addAction(sort_pages_action)


        edit_button = QToolButton(self)
        edit_button.setText("&Edit")
        edit_button.setMenu(edit_menu)
        edit_button.setPopupMode(QToolButton.InstantPopup)

        convert_to_word_action = QAction('Pdf to Word', self)
        convert_to_word_action.triggered.connect(self.convert_to_word)

        Batch_convert_MSOFFICE_to_PDF_action = QAction('Batch Convert MSOFFICE Document to PDF', self)
        Batch_convert_MSOFFICE_to_PDF_action.triggered.connect(self.open_batch_convert_dialog)

        Batch_convert_TXT_to_PDF_action = QAction('Batch Convert TXT to PDF', self)
        Batch_convert_TXT_to_PDF_action.triggered.connect(self.open_batch_TXT_convert_dialog)

        count_pages_action = QAction('Count PDF pages in a folder', self)
        count_pages_action.triggered.connect(self.count_pdf_pages)

        extract_text_action = QAction('Extract text from PDF by selecting position', self)
        extract_text_action.triggered.connect(self.open_text_extractor)

        merge_pdfs_action = QAction('Merge multiple PDF files', self)
        merge_pdfs_action.triggered.connect(self.merge_pdfs)

        PDF_Form_Overlay_action = QAction('PDF_Form_Overlay', self)
        PDF_Form_Overlay_action.triggered.connect(self.open_overlay_dialog)

        PDF_Batch_Print_action = QAction('Batch Print PDF', self)
        PDF_Batch_Print_action.triggered.connect(self.open_batch_print_dialog)

        deep_search_action = QAction('Deep Search PDF Files', self)
        deep_search_action.triggered.connect(self.open_deep_search_dialog)

        merge_csv_excel_action = QAction('Merge CSV/Excel Files', self)
        merge_csv_excel_action.triggered.connect(self.open_merge_csv_excel_dialog)


        # 2. Create "Utilities" menu
        utilities_menu = QMenu("&Utilities", self)

        # 3. Add QActions to the menu
        utilities_menu.addAction(convert_to_word_action)
        utilities_menu.addSeparator()  # Add separator
        utilities_menu.addAction(Batch_convert_MSOFFICE_to_PDF_action)
        utilities_menu.addAction(Batch_convert_TXT_to_PDF_action)
        utilities_menu.addSeparator()  # Add separator
        utilities_menu.addAction(count_pages_action)
        utilities_menu.addSeparator()  # Add separator
        utilities_menu.addAction(extract_text_action)
        utilities_menu.addAction(merge_pdfs_action)
        utilities_menu.addAction(PDF_Form_Overlay_action)
        utilities_menu.addSeparator()  # Add separator
        utilities_menu.addAction(PDF_Batch_Print_action)
        utilities_menu.addSeparator()  # Add separator
        utilities_menu.addAction(deep_search_action)
        utilities_menu.addSeparator()  # Add separator
        utilities_menu.addAction(merge_csv_excel_action)

        # 4. Create "Utilities" button
        utilities_button = QToolButton(self)
        utilities_button.setText("&Utilities")
        utilities_button.setMenu(utilities_menu)
        utilities_button.setPopupMode(QToolButton.InstantPopup)

        # Add Edit button to toolbar
        toolbar.addWidget(edit_button)

        # 5. Add button to toolbar
        toolbar.addWidget(utilities_button)


        file_menu.setStyleSheet("""
            QMenu {
                background-color: #272727; /* 深灰色背景 */
                color: white; /* 白色文字 */
                border: 1px solid #666666; /* 浅灰色边框 */
                border-radius: 10px; /* 菜单边框圆角 */
            }
            QMenu::item {
                padding: 5px 20px; /* 菜单项内边距 */
                margin: 2px 2px; /* 菜单项外边距 */
            }
            QMenu::item:selected {
                background-color: #02F78E; /* 选中时的背景色 */
                color: black; /* 选中时的文字颜色 */
                border-radius: 5px; /* 选中项的边框圆角 */
                padding: 5px 20px; /* 选中项的内边距 */
            }
        """)

        # 为 edit_menu 设置相同的样式
        edit_menu.setStyleSheet("""
            QMenu {
                background-color: #272727; /* 深灰色背景 */
                color: white; /* 白色文字 */
                border: 1px solid #666666; /* 浅灰色边框 */
                border-radius: 10px; /* 菜单边框圆角 */
            }
            QMenu::item {
                padding: 5px 20px; /* 菜单项内边距 */
                margin: 2px 2px; /* 菜单项外边距 */
            }
            QMenu::item:selected {
                background-color: #02F78E; /* 选中时的背景色 */
                color: black; /* 选中时的文字颜色 */
                border-radius: 5px; /* 选中项的边框圆角 */
                padding: 5px 20px; /* 选中项的内边距 */
            }
        """)

        # 设置工具栏样式
        toolbar.setStyleSheet("""
            QToolBar {
                background-color: #272727; /* 深灰色背景 */
                border: none;
            }
            QToolButton {
                color: white; /* 白色文字 */
                background-color: transparent; /* 透明背景 */
                border: none;
                font-family: Arial Black;
                font-weight: bold;     
            }
            QToolButton:hover {
                background-color: #02F78E; /* 选中时变暗 */
                color: black;
            }
            QToolButton::menu-indicator {
                image: none; /* 隐藏默认的下拉箭头 */
            }
        """)

        # 6. 设置菜单样式
        utilities_menu.setStyleSheet("""
            QMenu {
                background-color: #272727; /* 深灰色背景 */
                color: white; /* 白色文字 */
                border: 1px solid #666666; /* 浅灰色边框 */
                border-radius: 10px; /* 菜单边框圆角 */
            }
            QMenu::item {
                padding: 5px 20px; /* 菜单项内边距 */
                margin: 2px 2px; /* 菜单项外边距 */
            }
            QMenu::item:selected {
                background-color: #02F78E; /* 选中时的背景色 */
                color: black; /* 选中时的文字颜色 */
                border-radius: 5px; /* 选中项的边框圆角 */
                padding: 5px 20px; /* 选中项的内边距 */
            }
        """)

        # Add the button to the toolbar
        toolbar.addWidget(window_button)
        window_menu.setStyleSheet("""
            QMenu {
                background-color: #272727; /* 深灰色背景 */
                color: white; /* 白色文字 */
                border: 1px solid #666666; /* 浅灰色边框 */
                border-radius: 10px; /* 菜单边框圆角 */
            }
            QMenu::item {
                padding: 5px 20px; /* 菜单项内边距 */
                margin: 2px 2px; /* 菜单项外边距 */
            }
            QMenu::item:selected {
                background-color: #02F78E; /* 选中时的背景色 */
                color: black; /* 选中时的文字颜色 */
                border-radius: 5px; /* 选中项的边框圆角 */
                padding: 5px 20px; /* 选中项的内边距 */
            }
        """)


        self.toolbar = self.addToolBar('Main Toolbar')
        self.toolbar.addAction(main_menu_action)
        self.addToolBar(Qt.LeftToolBarArea, self.toolbar)


        self.add_toolbar_spacer()  # 添加间隔

        # 添加Open按钮
        open_action = QAction(QIcon(resource_path('App_icon/file.png')), 'Open', self)
        open_action.triggered.connect(self.open_pdf)
        self.toolbar.addAction(open_action)

        self.add_toolbar_spacer()  # 添加间隔

        # 添加Save按钮
        save_action = QAction(QIcon(resource_path('App_icon/Save.png')), 'Save', self)
        save_action.triggered.connect(self.save_file)
        self.toolbar.addAction(save_action)

        self.add_toolbar_spacer()  # 添加间隔

        # 添加PDF Info按钮
        pdf_info_action = QAction(QIcon(resource_path('App_icon/info.png')), 'PDF Info', self)
        pdf_info_action.triggered.connect(self.show_pdf_info)
        self.toolbar.addAction(pdf_info_action)

        self.add_toolbar_spacer()  # 添加间隔

        # 添加Print按钮
        print_action = QAction(QIcon(resource_path('App_icon/print.png')), 'PRINT', self)
        print_action.triggered.connect(self.show_print_options)
        self.toolbar.addAction(print_action)

        spacer = QWidget()
        spacer.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.toolbar.addWidget(spacer)

        # 添加README按钮
        readme_action = QAction(QIcon(resource_path('App_icon/readme.png')), '', self)
        readme_action.setToolTip('README')
        readme_action.triggered.connect(self.show_readme)
        self.toolbar.addAction(readme_action)


        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        self.progress_bar.setMinimum(0)
        self.progress_bar.setMaximum(0)
        layout.addWidget(self.progress_bar)

        progress_layout = QHBoxLayout()
        progress_layout.addStretch(1)
        progress_layout.addWidget(self.progress_bar)
        progress_layout.addStretch(1)
        
        # 設置進度條的固定大小
        self.progress_bar.setFixedWidth(500)
        self.progress_bar.setFixedHeight(15)
        
        layout.addLayout(progress_layout)


        # Display area for PDF pages
        self.image_widget = QWidget()
        self.image_widget.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        self.image_label = QLabel(self.image_widget)
        self.image_label.setAlignment(Qt.AlignCenter)

        image_layout = QVBoxLayout(self.image_widget)
        image_layout.addWidget(self.image_label)

        self.scroll_area = QScrollArea()
        self.scroll_area.setWidgetResizable(True)  # Set to true to allow widget resizing
        self.scroll_area.setWidget(self.image_widget)  # Add image_widget to the scroll area
        self.scroll_area.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.scroll_area.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.scroll_area.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.scroll_area.setStyleSheet("background-color: #1c1c1c")

        layout.addWidget(self.scroll_area)  # Add the scroll area to the main layout instead of the image_widget
        self.image_widget.wheelEvent = self.wheel_event
        
        # 創建底部信息區域
        bottom_info_widget = QWidget()
        bottom_info_layout = QHBoxLayout(bottom_info_widget)
        bottom_info_layout.setContentsMargins(10, 5, 10, 5)  # 設置內邊距

        # 設置底部信息區域的背景顏色
        bottom_info_widget.setStyleSheet("""
            QWidget {
                background-color: #222222;  /* 深灰色背景 */
                border: none;
                border-radius: 5px; 
            }
        """)

        # 文件信息標籤
        self.file_info_label = QLabel(self)
        self.file_info_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        bottom_info_layout.addWidget(self.file_info_label)
        self.file_info_label.setStyleSheet("""
            QLabel {
                color: white;
            }
        """)

        # 添加彈性空間
        bottom_info_layout.addStretch()

        # 顯示頁面尺寸的標籤
        self.page_size_label = QLabel(self)
        self.page_size_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        bottom_info_layout.addWidget(self.page_size_label)
        self.page_size_label.setStyleSheet("""
            QLabel {
                color: white;
                margin-right: 10px;
            }
        """)

        # 顯示當前頁碼
        self.page_label = QLabel(self)
        self.page_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        bottom_info_layout.addWidget(self.page_label)
        self.page_label.setStyleSheet("""
            QLabel {
                color: white;
            }
        """)

        # 將底部信息區域添加到主佈局的底部
        layout.addWidget(bottom_info_widget)                  
                                      
        # 按鈕和輸入框的佈局
        button_widget = QWidget()
        button_layout = QHBoxLayout(button_widget)
        button_layout.setContentsMargins(10, 5, 10, 5)  # 設置內邊距

        button_widget.setStyleSheet("""
            QWidget {
                background-color: #222222;
                border-radius: 5px; 
            }
        """)

        self.btn_rotate_pages = QPushButton('ROTATE', self)
        self.btn_rotate_pages.clicked.connect(self.rotate_pages)
        button_layout.addWidget(self.btn_rotate_pages)
        self.btn_rotate_pages.setStyleSheet("""
            QPushButton {
                background-color: #004C99;  
                color: white;
                border: none;
                border-radius: 5px;
                padding: 8px;
                font-weight: bold;
                min-width: 30px;
            }
            QPushButton:hover {
                background-color: #0073E7; 
            }
            QPushButton:pressed {
                background-color: #003366;  
            }
        """)

        self.rotate_pages_edit = QLineEdit(self)
        self.rotate_pages_edit.setFixedWidth(80)
        self.rotate_pages_edit.setPlaceholderText("輸入頁碼")
        button_layout.addWidget(self.rotate_pages_edit)
        self.rotate_pages_edit.setStyleSheet("""
            QLineEdit {
                background-color: #FFFFFF;
                color: #000000;
                border: 1px solid #CCCCCC;
                border-radius: 5px;
                padding: 5px;
            }
        """)

        self.rotate_angle_combo = QComboBox(self)
        self.rotate_angle_combo.addItems(['0','90', '180', '270'])
        button_layout.addWidget(self.rotate_angle_combo)
        self.rotate_angle_combo.setStyleSheet("""
            QComboBox {
                background-color: #171717;
                color: #FFFFFF;
                border: 1px solid #CCCCCC;
                border-radius: 5px;
                padding: 5px;
            }
            QComboBox::drop-down {
                subcontrol-origin: padding;
                subcontrol-position: top right;
                width: 15px;
                border-left: 1px solid darkgray;
            }
            QComboBox QAbstractItemView {
                border: 2px solid darkgray;
                selection-background-color: #6a6ea9;
                selection-color: darkgreen;
            }
        """)

        button_layout.addStretch()

        self.btn_prev = QPushButton('<<', self)
        self.btn_prev.clicked.connect(self.show_prev_page)
        button_layout.addWidget(self.btn_prev)

        self.page_num_edit = QLineEdit(self)
        self.page_num_edit.setFixedWidth(50)
        self.page_num_edit.setPlaceholderText("輸入頁碼")
        button_layout.addWidget(self.page_num_edit)
        self.page_num_edit.setStyleSheet("""
            QLineEdit {
                background-color: #FFFFFF;
                color: #000000;
                border: 1px solid #CCCCCC;
                border-radius: 5px;
                padding: 5px;
            }
        """)

        self.btn_goto = QPushButton('GO', self)
        self.btn_goto.clicked.connect(self.goto_page)
        button_layout.addWidget(self.btn_goto)
        self.btn_goto.setStyleSheet("""
            QPushButton {
                background-color: #004C99;  
                color: white;
                border: none;
                border-radius: 5px;
                padding: 8px;
                font-weight: bold;
                min-width: 25px;
            }
            QPushButton:hover {
                background-color: #0073E7; 
            }
            QPushButton:pressed {
                background-color: #003366;  
            }
        """)

        self.btn_next = QPushButton('>>', self)
        self.btn_next.clicked.connect(self.show_next_page)
        button_layout.addWidget(self.btn_next)

        button_layout.addStretch()

        self.btn_zoom_out = QPushButton('-', self)
        self.btn_zoom_out.clicked.connect(self.zoom_out)
        button_layout.addWidget(self.btn_zoom_out)
        self.btn_zoom_out.setStyleSheet("""
            QPushButton {
                background-color: #606060;  
                color: white;
                border: none;
                border-radius: 5px;
                padding: 8px;
                min-width: 20px;
            }
            QPushButton:hover {
                background-color: #808080; 
            }
            QPushButton:pressed {
                background-color: #404040;  
            }
        """)

        self.zoom_label = QLabel("100%", self)
        self.zoom_label.setAlignment(Qt.AlignCenter)
        button_layout.addWidget(self.zoom_label)
        self.zoom_label.setStyleSheet("""
            QLabel {                          
                color: white;
            }
        """)

        self.btn_zoom_in = QPushButton('+', self)
        self.btn_zoom_in.clicked.connect(self.zoom_in)
        button_layout.addWidget(self.btn_zoom_in)
        self.btn_zoom_in.setStyleSheet("""
            QPushButton {
                background-color: #606060;  
                color: white;
                border: none;
                border-radius: 5px;
                padding: 8px;
                min-width: 20px;
            }
            QPushButton:hover {
                background-color: #808080; 
            }
            QPushButton:pressed {
                background-color: #404040;  
            }
        """)

        layout.addWidget(button_widget)

        self.show()


    def setup_button(self, text, callback, icon_path, normal_color='#343434', hover_color='#02F78E'):
        button = QPushButton(text, self)
        button.clicked.connect(callback)
        self.left_button_layout.addWidget(button)
        button.setStyleSheet(self.button_style(normal_color, hover_color))

        # 設置圖示
        icon = QIcon(icon_path)
        button.setIcon(icon)
        button.setIconSize(button.size())

        return button


    def button_style(self, normal_color, hover_color):
        """返回按钮的通用样式"""
        return f"""
            QPushButton {{
                background-color: {normal_color};
                color: white;
                border: none;
                border-radius: 5px;
                padding: 5px;
                font-family: Calibri;
                font-weight: bold;
                min-width: 120px;
                text-align: left;
                padding-left: 10px;
            }}
            QPushButton:hover {{
                background-color: {hover_color};
                color: black;
            }}
            QPushButton:pressed {{
                background-color: #FF6347;
            }}
        """
    

    def load_settings(self):
        if os.path.exists(self.config_file):
            try:
                with open(self.config_file, 'r') as f:
                    return json.load(f)
            except json.JSONDecodeError:
                print("Config file is corrupted. Using default settings.")
        return {}  # 返回空字典作為默認設定

    def save_settings(self):
        settings = {
            'window_size': [self.width(), self.height()],
            'window_position': [self.x(), self.y()],
            'zoom_ratio': self.zoom_ratio,
            'last_directory': getattr(self, 'last_directory', '')
        }
        # 移除值為 None 的項
        settings = {k: v for k, v in settings.items() if v is not None}
        
        # 確保使用絕對路徑保存設置文件
        with open(self.config_file, 'w') as f:
            json.dump(settings, f)


    def apply_settings(self):
        if 'window_size' in self.settings:
            self.resize(*self.settings['window_size'])
        if 'window_position' in self.settings:
            self.move(*self.settings['window_position'])
        if 'zoom_ratio' in self.settings:
            self.zoom_ratio = self.settings['zoom_ratio']
        if 'last_directory' in self.settings:
            self.last_directory = self.settings['last_directory']


    def open_pdf(self):
        initial_dir = self.last_directory if hasattr(self, 'last_directory') else ''
        path, _ = QFileDialog.getOpenFileName(self, "選擇 PDF 文件", initial_dir, "PDF files (*.pdf)")
        if path:
            try:
                # 创建一个临时文件夹来存储临时副本
                temp_dir = tempfile.mkdtemp()
                temp_pdf_path = os.path.join(temp_dir, os.path.basename(path))
                shutil.copy2(path, temp_pdf_path)  # 将原始文件复制到临时文件夹
                
                self.load_pdf(temp_pdf_path)  # 加载临时副本进行编辑
                self.page_num_edit.setText(str(self.current_page + 1)) 
                
                # 存储原始文件路径和暂存文件路径
                self.original_pdf_path = path
                self.temp_pdf_path = temp_pdf_path
                
                # 更新UI显示文件信息
                file_info = f"文件路徑: {self.original_pdf_path} | 頁數: {self.doc.page_count}"
                self.file_info_label.setText(file_info)
                self.setWindowTitle(f'PDFDocuEdit Pro - {os.path.basename(path)}')

                # 保存最后打开的目录
                self.last_directory = os.path.dirname(path)
            except Exception as e:
                self.show_error(f"無法打開文件: {e}")

    def load_pdf(self, path):
        try:
            # Always create a temporary copy
            temp_dir = tempfile.mkdtemp()
            temp_pdf_path = os.path.join(temp_dir, os.path.basename(path))
            shutil.copy2(path, temp_pdf_path)

            # Attempt to open the PDF
            doc = fitz.open(temp_pdf_path)

            # Check if the document is encrypted
            if doc.is_encrypted:
                # Attempt to decrypt the document
                password_given = False
                while not password_given:
                    password, ok = QInputDialog.getText(self, "PDF加密", "此PDF文件已被加密。請輸入密碼：", QLineEdit.Password)
                    if ok and password:
                        if doc.authenticate(password):
                            password_given = True
                        else:
                            self.show_error("提供的密碼不正確。請再試一次。")
                    else:
                        raise Exception("無法打開加密的PDF文件，因為未提供密碼。")

            self.doc = doc
            self.current_page = 0
            self.zoom_ratio = 1.0
            self.show_page(self.current_page)
            self.update_file_info(path)  # Update file info in the UI
            self.setWindowTitle(f'PDFDocuEdit Pro - {os.path.basename(path)}')

            # Store original and temporary paths
            self.original_pdf_path = path
            self.temp_pdf_path = temp_pdf_path

            # 更新最后打开的文件路径
            self.last_file = path

        except Exception as e:
            self.show_error(f"加載 PDF 時發生錯誤：{e}")

        if self.sort_widget:
            self.sort_widget.update_page_list(self.doc.page_count)

    def show_page(self, num):
        if not self.doc or num < 0 or num >= self.doc.page_count:
            return

        try:
            page = self.doc.load_page(num)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(self.zoom_ratio, self.zoom_ratio))
            qimage = QImage(pixmap.samples, pixmap.width, pixmap.height, pixmap.stride, QImage.Format_RGB888)
            self.image_label.setPixmap(QPixmap.fromImage(qimage))
            self.page_label.setText(f"第 {num + 1} 頁,共 {self.doc.page_count} 頁")

            # 獲取頁面尺寸
            page_rect = page.rect
            page_width_pts = page_rect.width
            page_height_pts = page_rect.height

            # 將點轉換為毫米並移除小數點後的數字
            page_width_mm = int(page_width_pts / 72 * 25.4)
            page_height_mm = int(page_height_pts / 72 * 25.4)

            # 更新頁碼和頁面尺寸標籤
            self.page_label.setText(f"第 {num + 1} 頁,共 {self.doc.page_count} 頁")
            self.page_size_label.setText(f"(尺寸: {page_width_mm} x {page_height_mm} mm)")

            # 檢查搜索小部件是否存在並且勾選了高亮選項
            should_highlight = (hasattr(self, 'search_widget') and 
                            self.search_widget and 
                            self.search_widget.highlight_checkbox.isChecked())

            # 只有在需要高亮的情況下才執行高亮邏輯
            if should_highlight and self.search_results:
                for page_num, instances in self.search_results:
                    if page_num == num:
                        for inst in instances:
                            try:
                                highlight = page.add_highlight_annot(inst)
                                highlight.update()
                            except Exception as e:
                                print(f"添加高亮時出錯: {e}")

            # 添加以下代码来更新 QLineEdit 显示
            self.page_num_edit.setText(str(num + 1))
        except Exception as e:
            print(f"顯示頁面時出錯: {e}")

    def extract_pixmap(self, image_obj):
        byte_array = image_obj.stream.get_rawdata()
        image = QImage.fromData(byte_array)
        pixmap = QPixmap.fromImage(image)
        return pixmap
    

    def rotate_pages(self):
        if not self.doc:
            QMessageBox.warning(self, "錯誤", "沒有打開的PDF文件。")
            return

        page_range = self.rotate_pages_edit.text().strip()
        pages_to_rotate = self.parse_page_range(page_range)
        
        if pages_to_rotate is None:
            return  # parse_page_range 已經顯示了錯誤訊息

        try:
            angle = int(self.rotate_angle_combo.currentText())
            for page_num in pages_to_rotate:
                if 0 <= page_num < self.doc.page_count:
                    page = self.doc.load_page(page_num)
                    page.set_rotation(angle)
                else:
                    QMessageBox.warning(self, "錯誤", f"頁碼 {page_num + 1} 超出範圍。")
                    return

            self.show_page(self.current_page)
            QMessageBox.information(self, "成功", "頁面已成功旋轉！")
            self.is_modified = True
        except Exception as e:
            QMessageBox.critical(self, "錯誤", f"旋轉頁面時發生錯誤：{str(e)}")

    def toggle_menu_dock(self):
        if self.menu_dock.isVisible():
            self.menu_dock.hide()
        else:
            self.menu_dock.show()

    def show_prev_page(self):
        if self.doc and self.current_page > 0:
            self.current_page -= 1
            self.show_page(self.current_page)

    def show_next_page(self):
        if self.doc and self.current_page < len(self.doc) - 1:
            self.current_page += 1
            self.show_page(self.current_page)

    def zoom_in(self):
        if self.doc:
            self.zoom_ratio *= PDFViewer.ZOOM_IN_FACTOR
            self.zoom_ratio = min(self.zoom_ratio, 4.0)  # 限制最大縮放比例為400%
            self.show_page(self.current_page)
            self.update_zoom_label()

    def zoom_out(self):
        if self.doc:
            self.zoom_ratio *= PDFViewer.ZOOM_OUT_FACTOR
            self.zoom_ratio = max(self.zoom_ratio, 0.20)  # 限制最小縮放比例為20%
            self.show_page(self.current_page)
            self.update_zoom_label()

    def update_zoom_label(self):
        zoom_percent = int(self.zoom_ratio * 100)
        self.zoom_label.setText(f"{zoom_percent}%")

    def goto_page(self):
        if self.doc:
            try:
                page_num = int(self.page_num_edit.text())
                if 1 <= page_num <= self.doc.page_count:
                    self.current_page = page_num - 1
                    self.show_page(self.current_page)
                else:
                    self.show_error("頁碼超出範圍！")
            except ValueError:
                self.show_error("請輸入有效的頁碼！")

        # 更新 QLineEdit 顯示當前頁碼
        self.page_num_edit.setText(str(self.current_page + 1))

    def delete_pages(self):
        if not self.doc:
            QMessageBox.warning(self, "警告", "請先打開一個PDF文件")
            return
        
        dialog = DeletePagesDialog(self)
        if dialog.exec_() == QDialog.Accepted:
            delete_option = dialog.get_delete_options()
            
            if not delete_option:
                return  # 如果解析失敗或用戶取消，直接返回
            
            pages_to_delete = []
            if delete_option == "odd":
                pages_to_delete = [i for i in range(self.doc.page_count) if (i + 1) % 2 != 0]
            elif delete_option == "even":
                pages_to_delete = [i for i in range(self.doc.page_count) if (i + 1) % 2 == 0]
            elif delete_option and delete_option[0] == "multiple":
                multiple = delete_option[1]
                pages_to_delete = [i for i in range(self.doc.page_count) if (i + 1) % multiple == 0]
            elif delete_option and delete_option[0] == "custom":
                pages_to_delete = self.parse_page_range(delete_option[1])
                if pages_to_delete is None:  # 若解析返回 None，表示有錯誤發生
                    return  # 直接返回，不進行刪除操作
            elif delete_option and delete_option[0] == "complex":
                n, x, z = delete_option[1], delete_option[2], delete_option[3]
                current_page = n
                while current_page < self.doc.page_count:
                    end_page = min(current_page + z, self.doc.page_count)
                    pages_to_delete.extend(range(current_page, end_page))
                    current_page += x + z  # Skip x pages after deleting z pages
            
            for page_num in sorted(pages_to_delete, reverse=True):
                self.doc.delete_page(page_num)
            
            # 更新顯示但不立即保存文件
            self.current_page = max(0, min(self.current_page, self.doc.page_count - 1))
            self.show_page(self.current_page)
            QMessageBox.information(self, "成功", "所選頁面已被標記為刪除")
            self.is_modified = True

    def merge_pdfs(self):
        merge_dialog = MergePDFsDialog(self)
        if merge_dialog.exec_() == QDialog.Accepted:
            merged_doc = fitz.open()
            for path in merge_dialog.file_paths:
                with fitz.open(path) as doc:
                    merged_doc.insert_pdf(doc)
            save_path, _ = QFileDialog.getSaveFileName(self, "保存合併後的PDF文件", "", "PDF files (*.pdf)")
            if save_path:
                merged_doc.save(save_path)
                QMessageBox.information(self, "成功", "PDF文件已成功合併！")
                self.load_pdf(save_path)  # 重新加载合併後的文檔

    def extract_pages_by_number(self, page_numbers):
        if self.doc:
            extracted_doc = fitz.open()
            for page_num in sorted(page_numbers):
                page = self.doc.load_page(page_num)
                extracted_doc.insert_pdf(self.doc, from_page=page_num, to_page=page_num)
            save_path, _ = QFileDialog.getSaveFileName(self, "Save Extracted Pages", "", "PDF files (*.pdf)")
            if save_path:
                extracted_doc.save(save_path)
                QMessageBox.information(self, "成功", "提取的頁面已保存到：" + save_path)
        else:
            QMessageBox.warning(self, "警告", "請先打開一個PDF文件。")
        self.is_modified = True

    def delete_pages_by_number(self, page_numbers):
        try:
            if self.doc:
                for page_num in sorted(page_numbers, reverse=True):
                    self.doc.delete_page(page_num)
                # 检查当前页码是否仍然有效
                if self.current_page >= self.doc.page_count:
                    self.current_page = max(0, self.doc.page_count - 1)
                # 清除搜索结果
                self.search_results.clear()
                # 清空结果列表
                if self.search_widget:
                    self.search_widget.clear_result_list() 
                self.show_page(self.current_page)
                QMessageBox.information(self, "成功", "所選頁面已被標記為刪除")
                self.is_modified = True
            else:
                QMessageBox.warning(self, "警告", "請先打開一個PDF文件")
        except Exception as e:
            QMessageBox.critical(self, "错误", f"删除页面时发生错误: {e}")
            print(f"Error deleting pages: {e}")
        self.is_modified = True

    def extract_pages(self):
        if self.doc:
            dialog = ExtractPagesDialog(self)
            if dialog.exec_() == QDialog.Accepted:
                extract_option = dialog.get_extract_options()
                
                try:
                    pages_to_extract = []
                    if extract_option == "odd":
                        pages_to_extract = [i for i in range(self.doc.page_count) if (i + 1) % 2 != 0]
                    elif extract_option == "even":
                        pages_to_extract = [i for i in range(self.doc.page_count) if (i + 1) % 2 == 0]
                    elif extract_option and extract_option[0] == "multiple":
                        multiple = extract_option[1]
                        pages_to_extract = [i for i in range(self.doc.page_count) if (i + 1) % multiple == 0]
                    elif extract_option and extract_option[0] == "custom":
                        pages_to_extract = self.parse_page_range(extract_option[1])
                    elif extract_option and extract_option[0] == "complex":
                        n, x, z = extract_option[1], extract_option[2], extract_option[3]
                        current_page = n
                        while current_page < self.doc.page_count:
                            end_page = min(current_page + z, self.doc.page_count)
                            pages_to_extract.extend(range(current_page, end_page))
                            current_page += x + z  # Skip x pages after extracting z pages
                    
                    if pages_to_extract:
                        extracted_doc = fitz.open()
                        for page_num in sorted(pages_to_extract):
                            page = self.doc.load_page(page_num)
                            extracted_doc.insert_pdf(self.doc, from_page=page_num, to_page=page_num)
                        save_path, _ = QFileDialog.getSaveFileName(self, "Save Extracted Pages", "", "PDF files (*.pdf)")
                        if save_path:
                            extracted_doc.save(save_path)
                            QMessageBox.information(self, "成功", "提取的頁面已保存到：" + save_path)
                    else:
                        QMessageBox.warning(self, "警告", "沒有選擇頁面進行提取。")
                except Exception as e:
                    QMessageBox.critical(self, "錯誤", "提取頁面時發生錯誤：" + str(e))
        else:
            QMessageBox.warning(self, "警告", "請先打開一個PDF文件。")

        self.is_modified = True


    def save_file(self):
        if self.doc:
            if self.original_pdf_path:
                try:
                    self.doc.save(self.original_pdf_path)
                    QMessageBox.information(self, "Success", "PDF saved successfully!")
                    self.update_file_info(self.original_pdf_path)
                except Exception as e:
                    self.show_error(f"Error saving PDF: {e}")
            else:
                self.save_as_file()
                self.is_modified = False

    def save_as_file(self):
        if self.doc:
            save_path, _ = QFileDialog.getSaveFileName(self, "Save PDF As", "", "PDF files (*.pdf)")
            if save_path:
                try:
                    self.doc.save(save_path)
                    QMessageBox.information(self, "Success", "PDF saved successfully!")
                    self.update_file_info(save_path)
                except Exception as e:
                    self.show_error(f"Error saving PDF: {e}")

    def closeEvent(self, event):
        if self.is_modified:
            reply = QMessageBox(self)
            reply.setWindowTitle('PDFDocuEdit Pro')
            reply.setText('文件已被修改，是否保存?')
            reply.setStandardButtons(QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel)
            reply.setStyleSheet("QLabel{ color: white; }")

            user_reply = reply.exec_()

            if user_reply == QMessageBox.Yes:
                self.save_file()
                self.save_settings()
                event.accept()
            elif user_reply == QMessageBox.No:
                self.save_settings()
                event.accept()
            else:
                event.ignore()
        else:
            self.save_settings()
            event.accept()

    def handle_insert(self, pdf_path, page_range, insert_position, insert_page):
        if not self.doc or not os.path.exists(pdf_path):
            self.show_warning("請先打開一個PDF文件或檢查插入文檔的路徑。")
            return

        try:
            insert_doc = fitz.open(pdf_path)
            pages_to_insert = self.parse_page_range(page_range)
            if insert_position == "開頭":
                position = 0
            elif insert_position == "結尾":
                position = self.doc.page_count
            elif insert_position == "指定位置":
                position = int(insert_page) - 1
            else:
                self.show_warning("未知的插入位置")
                return

            self.doc.insert_pdf(insert_doc, from_page=pages_to_insert[0], to_page=pages_to_insert[-1], start_at=position)
            self.show_page(self.current_page)  # 刷新當前頁面的顯示
            self.show_success("頁面插入成功！")
        except Exception as e:
            self.show_error(f"插入頁面時發生錯誤：{e}")
        finally:
            insert_doc.close()

    def insert_pages(self):
        dialog = InsertPagesDialog(self)
        if dialog.exec_() == QDialog.Accepted:
            details = dialog.get_insert_details()
            self.handle_insert(details['pdf_path'], details['page_range'], details['insert_position'], details['insert_page'])

    def insert_pages_repeat(self, pdf_path, page_range, interval):
        if not self.doc or not os.path.exists(pdf_path):
            QMessageBox.warning(self, "警告", "請先打開一個PDF文件或檢查插入文檔的路徑。")
            return

        try:
            insert_doc = fitz.open(pdf_path)
            pages_to_insert = self.parse_page_range(page_range)
            
            # 开始插入的起始页码调整为间隔值，即从第 interval+1 页开始插入
            current_page = interval  # 调整为从第 interval+1 页开始插入

            while current_page < self.doc.page_count:
                for page_num in pages_to_insert:
                    if current_page < self.doc.page_count:
                        self.doc.insert_pdf(insert_doc, from_page=page_num, to_page=page_num, start_at=current_page)
                        current_page += 1  # 插入完毕后，移动到下一个位置
                current_page += interval  # 完成一次插入后，跳过指定的间隔数

            self.show_page(0)  # 显示文档第一页
            QMessageBox.information(self, "成功", "頁面插入成功！")
        except Exception as e:
            QMessageBox.critical(self, "錯誤", f"插入頁面時發生錯誤：{e}")
        finally:
            insert_doc.close()


    def print_pdf(self):
        if self.doc:
            self.print_options_dock.show()  # Show the print options widget

    def show_print_options(self):
        if self.doc:
            self.addDockWidget(Qt.RightDockWidgetArea, self.print_options_dock)  # Add dock widget
            self.print_options_dock.show()
        else:
            self.show_warning("請先打開一個PDF文件。")


    def start_printing(self):
        try:
            # Get options from PrintOptionsWidget
            selected_printer_name = self.print_options_widget.printer_combo.currentText()
            print(f"Selected printer: {selected_printer_name}")  # Print the selected printer name
            
            # Create QPrinter and set properties
            printer = QPrinter(QPrinter.HighResolution)
            printer.setPrinterName(selected_printer_name)  # Set the selected printer
            
            # Set the document name to the original PDF file name
            if self.original_pdf_path:
                doc_name = os.path.basename(self.original_pdf_path)
                printer.setDocName(doc_name)
            
            paper_size = self.print_options_widget.paper_size_combo.currentText()
            
            # 在设置打印机属性之前，根据 PDF 设置方向
            self.print_options_widget.set_orientation_based_on_pdf(self.doc)

            # 获取方向设置
            orientation = self.print_options_widget.get_orientation()
            if orientation is None:
                # 如果返回 None，说明选择了自动，此时已经根据 PDF 设置了正确的方向
                orientation = QPrinter.Portrait if self.print_options_widget.portrait_radio.isChecked() else QPrinter.Landscape
            
            left_offset = self.print_options_widget.left_offset_spin.value()
            top_offset = self.print_options_widget.top_offset_spin.value()
            right_offset = self.print_options_widget.right_offset_spin.value() + 5  # Hardcoded right margin
            bottom_offset = self.print_options_widget.bottom_offset_spin.value() + 5  # Hardcoded bottom margin
            fit_to_margin = self.print_options_widget.fit_to_margin_checkbox.isChecked()
            copies = self.print_options_widget.copies_spin.value()
            duplex = self.print_options_widget.duplex_checkbox.isChecked()
            page_range = self.print_options_widget.page_range_edit.text()
            scale_factor = self.print_options_widget.scale_factor_spin.value() / 100.0  # Convert to a multiplier

            # Set paper size
            if paper_size == "A4":
                printer.setPageSize(QPrinter.A4)
            elif paper_size == "A3":
                printer.setPageSize(QPrinter.A3)
            elif paper_size == "Letter":
                printer.setPageSize(QPrinter.Letter)
            elif paper_size == "A5":  # 这里改为 A5
                printer.setPageSize(QPrinter.A5)  # 使用 QPrinter.A5
            else:
                self.show_error(f"不支持的纸张尺寸: {paper_size}")
                return
            
            printer.setOrientation(orientation)
            printer.setCopyCount(copies)
            
            # Set duplex mode
            if duplex:
                printer.setDuplex(QPrinter.DuplexAuto)
            else:
                printer.setDuplex(QPrinter.DuplexNone)

            # Parse page range
            pages = self.parse_page_range(page_range) if page_range else list(range(self.doc.page_count))

            # Call print_direct with the configured printer and offsets
            self.print_direct(printer, left_offset, top_offset, right_offset, bottom_offset, fit_to_margin, pages, scale_factor)

            # Close the dock widget after printing
            self.print_options_dock.close()
        except Exception as e:
            self.show_error(f"Start printing error: {e}")

    def print_direct(self, printer, left_offset, top_offset, right_offset, bottom_offset, fit_to_margin, pages, scale_factor):
        painter = QPainter(printer)
        try:
            progress_dialog = QProgressDialog("Printing...", "Cancel", 0, len(pages), None)
            progress_dialog.setWindowTitle("Print Progress")
            progress_dialog.setWindowModality(Qt.WindowModal)
            progress_dialog.setMinimumDuration(2000)
            progress_dialog.show()

            for i, page_num in enumerate(pages):
                if progress_dialog.wasCanceled():
                    break

                progress_dialog.setValue(i)
                progress_dialog.setLabelText(f"Printing page {i + 1} of {len(pages)}")
                QApplication.processEvents()

                # 載入當前頁面
                page = self.doc.load_page(page_num)
                page_width, page_height = page.rect.width, page.rect.height
                
                # 根據頁面方向自動設置打印機方向
                is_landscape = page_width > page_height
                if self.print_options_widget.auto_orientation_radio.isChecked():
                    if is_landscape:
                        printer.setOrientation(QPrinter.Landscape)
                    else:
                        printer.setOrientation(QPrinter.Portrait)
                
                # 獲取打印機的紙張矩形和可打印區域矩形 (在設置方向後)
                paper_rect = printer.paperRect()
                page_rect = printer.pageRect()
                
                # 根據fit_to_margin的值選擇適當的矩形
                target_rect = page_rect if fit_to_margin else paper_rect
                
                # 渲染PDF頁面
                zoom = 5
                mat = fitz.Matrix(zoom, zoom)
                pixmap = page.get_pixmap(matrix=mat)
                qimage = QImage(pixmap.samples, pixmap.width, pixmap.height, pixmap.stride, QImage.Format_RGB888)
                
                # 計算縮放因子，考慮頁面方向
                scale_x = target_rect.width() / qimage.width() * scale_factor
                scale_y = target_rect.height() / qimage.height() * scale_factor
                scale = min(scale_x, scale_y)
                
                # 縮放圖像
                scaled_width = int(qimage.width() * scale)
                scaled_height = int(qimage.height() * scale)
                scaled_image = qimage.scaled(scaled_width, scaled_height, Qt.KeepAspectRatio, Qt.SmoothTransformation)
                
                # 計算居中的圖像矩形
                image_rect = QRect(0, 0, scaled_image.width(), scaled_image.height())
                image_rect.moveCenter(target_rect.center())
                
                # 保存目前的painter狀態
                painter.save()
                
                # 應用偏移量
                offset_x = int(left_offset - right_offset) * printer.logicalDpiX() / 25.4
                offset_y = int(top_offset - bottom_offset) * printer.logicalDpiY() / 25.4
                painter.translate(offset_x, offset_y)
                
                # 繪製圖像
                painter.drawImage(image_rect, scaled_image)
                
                # 恢復painter狀態
                painter.restore()
                
                if i < len(pages) - 1:
                    printer.newPage()
                    
                time.sleep(0.5)

            progress_dialog.setValue(len(pages))
        except Exception as e:
            self.show_error(f"Direct printing error: {e}")
        finally:
            painter.end()
            progress_dialog.close()


    def convert_to_word(self):
        if self.doc:
            try:
                pdf_path = self.original_pdf_path
                docx_path = os.path.splitext(pdf_path)[0] + '.docx'
                self.thread = PDFConverter(pdf_path, docx_path)
                self.thread.finished.connect(self.conversion_finished)
                self.progress_bar.setVisible(True)
                self.thread.start()
            except Exception as e:
                self.show_error(f"轉換失敗：{e}")
        else:
            self.show_warning("請先打開一個PDF文件。")


    def conversion_finished(self):
        self.progress_bar.setVisible(False)
        self.show_success("PDF轉換為Word文件成功！")

    def auto_adjust_columns_width(self, writer, df, sheet_name):
        worksheet = writer.sheets[sheet_name]
        for column in df:
            column_width = max(df[column].astype(str).map(len).max(), len(column)) + 2
            col_idx = df.columns.get_loc(column) + 1
            worksheet.column_dimensions[get_column_letter(col_idx)].width = column_width

    def show_search_widget(self):
        if not self.search_widget:
            self.search_widget = SearchWidget(self)
            self.search_widget.search_requested.connect(self.search_text)
            self.addDockWidget(Qt.RightDockWidgetArea, self.search_widget)
        self.search_widget.show()

    def clear_highlights(self):
        """清除所有高亮標記，包含錯誤處理和更安全的實現"""
        if not self.doc:
            return
            
        try:
            for page_num in range(self.doc.page_count):
                page = self.doc.load_page(page_num)
                
                # 獲取頁面上所有註釋
                annots = page.annots()
                if not annots:
                    continue
                    
                # 使用安全的方式收集所有高亮註釋
                highlight_annots = []
                for annot in annots:
                    try:
                        if annot.type[0] == 8:  # 8 代表高亮註釋
                            highlight_annots.append(annot)
                    except Exception as e:
                        print(f"處理註釋時出錯: {e}")
                        continue
                
                # 逐個刪除收集到的高亮註釋
                for annot in highlight_annots:
                    try:
                        page.delete_annot(annot)
                    except Exception as e:
                        print(f"刪除註釋時出錯: {e}")
        except Exception as e:
            print(f"清除高亮時出錯: {e}")

    def search_text(self, text):
        if not self.doc:
            return
            
        try:
            # 先安全地清除所有高亮
            self.clear_highlights()
            
            # 清除搜索結果
            self.search_results.clear()
            
            # 執行搜索
            for page_num in range(self.doc.page_count):
                page = self.doc.load_page(page_num)
                text_instances = page.search_for(text)
                if text_instances:
                    self.search_results.append((page_num, text_instances))
            
            # 更新搜索小部件的顯示
            if hasattr(self, 'search_widget') and self.search_widget:
                self.search_widget.update_results(self.search_results)
                
            if not self.search_results:
                QMessageBox.information(self, "Search Results", "No matches found.")
        except Exception as e:
            print(f"搜索文本時出錯: {e}")
            QMessageBox.warning(self, "Search Error", f"An error occurred during search: {e}")


    def jump_to_page(self, page_num):
        self.current_page = page_num
        self.show_page(self.current_page)


    def parse_page_range(self, page_range):
        if not page_range.strip():
            QMessageBox.warning(self, "錯誤", "請輸入頁碼範圍！")
            return None

        pages = []
        for part in page_range.split(','):
            part = part.strip()
            if '-' in part:
                try:
                    start, end = map(int, part.split('-'))
                    if start > end:
                        QMessageBox.warning(self, "錯誤", "起始頁碼不能大於結束頁碼！")
                        return None
                    if start < 1 or end < 1:
                        QMessageBox.warning(self, "錯誤", "頁碼不能小於1！")
                        return None
                    pages.extend(range(start - 1, end))
                except ValueError:
                    QMessageBox.warning(self, "錯誤", "無效的頁碼範圍格式！")
                    return None
            else:
                try:
                    page_num = int(part)
                    if page_num < 1:
                        QMessageBox.warning(self, "錯誤", "頁碼不能小於1！")
                        return None
                    pages.append(page_num - 1)
                except ValueError:
                    QMessageBox.warning(self, "錯誤", "無效的頁碼格式！")
                    return None

        if not pages:
            QMessageBox.warning(self, "錯誤", "未指定有效的頁碼！")
            return None

        return sorted(set(pages))

    def open_new_pdf(self):
        """Opens a new PDF file in a separate window."""
        new_viewer = PDFViewer()  # Create a new instance of PDFViewer
        new_viewer.show()  # Show the new window
        self.open_windows.append(new_viewer)  # Add the new window to the list
        
        # (Optional) Open the file dialog in the new window
        path, _ = QFileDialog.getOpenFileName(new_viewer, "選擇 PDF 文件", "", "PDF files (*.pdf)")
        if path:
            new_viewer.load_pdf(path)
        
        return new_viewer  # Return the new viewer instance

    def show_window_list(self):
        """Displays a dialog with a list of open windows."""

        # 更新窗口列表
        self.open_windows = [window for window in self.open_windows if window.isVisible()]

        if not self.open_windows:
            self.show_info("沒有開啟其他視窗。")
            return

        window_list_dialog = QDialog(self)
        window_list_dialog.setWindowTitle("開啟的視窗")

        layout = QVBoxLayout()
        window_list = QListWidget()
        window_list.clear()  # 清空列表
        for window in self.open_windows:
            window_list.addItem(window.windowTitle())
        layout.addWidget(window_list)

        button_box = QDialogButtonBox(QDialogButtonBox.Ok)
        button_box.accepted.connect(window_list_dialog.accept)
        layout.addWidget(button_box)

        window_list_dialog.setLayout(layout)

        if window_list_dialog.exec_() == QDialog.Accepted:
            selected_item = window_list.currentItem()
            if selected_item:
                selected_window_title = selected_item.text()
                for window in self.open_windows:
                    if window.windowTitle() == selected_window_title:
                        window.activateWindow()  # Bring the selected window to the front
                        break

    def count_pdf_pages(self):
        folder_path = QFileDialog.getExistingDirectory(self, "選擇PDF文件夾")
        if folder_path:
            try:
                file_data = []

                def process_file(filepath):
                    try:
                        pdf = fitz.open(filepath)
                        num_pages = pdf.page_count
                        file_size = os.path.getsize(filepath)
                        modify_date = datetime.fromtimestamp(os.path.getmtime(filepath)).strftime('%Y-%m-%d %H:%M:%S')
                        pdf.close()
                        return {
                            'Filename': os.path.basename(filepath),
                            'PageCount': num_pages,
                            'FileSize (bytes)': file_size,
                            'Path': filepath.replace('\\', '/'),
                            'Modify Date': modify_date
                        }
                    except Exception as e:
                        return None  # 处理文件时发生错误，返回 None

                # 获取所有 PDF 文件路径
                pdf_files = []
                for root, dirs, files in os.walk(folder_path):
                    for filename in files:
                        if filename.lower().endswith('.pdf'):
                            filepath = os.path.join(root, filename)
                            pdf_files.append(filepath)

                if not pdf_files:
                    self.show_info("指定文件夾及其子資料夾中未找到PDF文件。")
                    return

                total_files = len(pdf_files)

                # 创建进度对话框
                progress_dialog = QProgressDialog("Processing files...", "Cancel", 0, total_files, self)
                progress_dialog.setWindowTitle("Progress")
                progress_dialog.setWindowModality(Qt.WindowModal)
                progress_dialog.setMinimumDuration(2000)
                progress_dialog.show()

                processed_files = 0

                def update_progress(future):
                    nonlocal processed_files
                    processed_files += 1
                    progress_dialog.setValue(processed_files)
                    if future.result() is not None:
                        file_data.append(future.result())

                # 使用线程池并发处理文件
                with concurrent.futures.ThreadPoolExecutor() as executor:
                    futures = [executor.submit(process_file, filepath) for filepath in pdf_files]
                    for future in concurrent.futures.as_completed(futures):
                        if progress_dialog.wasCanceled():
                            executor.shutdown(wait=False)
                            self.show_info("操作已取消。")
                            return
                        update_progress(future)

                if file_data:
                    df = pd.DataFrame(file_data)
                    total_pages = df['PageCount'].sum()
                    total_row = pd.DataFrame([{'Filename': 'Total', 'PageCount': total_pages, 'FileSize (bytes)': '', 'Path': '', 'Modify Date': ''}])
                    df = pd.concat([df, total_row], ignore_index=True)

                    report_path = os.path.join(folder_path, 'PDF_report.xlsx')
                    with pd.ExcelWriter(report_path, engine='openpyxl') as writer:
                        df.to_excel(writer, index=False, sheet_name='Report')
                        self.auto_adjust_columns_width(writer, df, 'Report')

                    self.show_success(f"統計完成。報告已保存到：{report_path}")
                else:
                    self.show_info("指定文件夾及其子資料夾中未找到有效的PDF文件。")
            except Exception as e:
                self.show_error(f"統計PDF頁數時出錯：{e}")
        else:
            self.show_warning("請選擇一個PDF文件夾。")

    def show_split_options(self):
        if self.doc:
            if not self.split_options_dock.isVisible():
                self.addDockWidget(Qt.RightDockWidgetArea, self.split_options_dock)
            self.split_options_dock.show()
        else:
            self.show_warning("請先打開一個PDF文件。")


    def start_splitting(self):
        split_option = self.split_options_widget.get_split_options()
        save_dir = self.split_options_widget.save_location_label.text()  # 获取保存位置
        file_prefix = self.split_options_widget.file_prefix_input.text()  # 获取文件名前缀
        original_file_name = os.path.splitext(os.path.basename(self.original_pdf_path))[0]

        if self.doc:
            try:
                if save_dir == "未選擇位置":
                    self.show_warning("請選擇文件保存位置！")
                    return

                if split_option == "every_page":
                    for page_num in range(self.doc.page_count):
                        new_doc = fitz.open()
                        new_doc.insert_pdf(self.doc, from_page=page_num, to_page=page_num)
                        save_path = f"{save_dir}/{file_prefix}_{original_file_name}_page_{page_num+1}.pdf"
                        new_doc.save(save_path)
                    self.show_success("PDF split successfully!")
                    
                elif split_option == "even_pages":
                    for page_num in range(0, self.doc.page_count, 2):
                        new_doc = fitz.open()
                        new_doc.insert_pdf(self.doc, from_page=page_num, to_page=page_num)
                        save_path = f"{save_dir}/{file_prefix}_{original_file_name}_page_{page_num+1}.pdf"
                        new_doc.save(save_path)
                    self.show_success("PDF split successfully!")
                    
                elif split_option == "odd_pages":
                    for page_num in range(1, self.doc.page_count, 2):
                        new_doc = fitz.open()
                        new_doc.insert_pdf(self.doc, from_page=page_num, to_page=page_num)
                        save_path = f"{save_dir}/{file_prefix}_{original_file_name}_page_{page_num+1}.pdf"
                        new_doc.save(save_path)
                    self.show_success("PDF split successfully!")
                    
                elif split_option and split_option[0] == "custom_pages":
                    pages_to_split = self.parse_page_range(split_option[1])
                    total_pages = self.doc.page_count

                    start_page = 0
                    for i, end_page in enumerate(pages_to_split):
                        new_doc = fitz.open()
                        new_doc.insert_pdf(self.doc, from_page=start_page, to_page=end_page - 1)
                        save_path = f"{save_dir}/{file_prefix}_{original_file_name}_pages_{start_page + 1}-{end_page}.pdf"
                        new_doc.save(save_path)
                        start_page = end_page

                    if start_page < total_pages:
                        new_doc = fitz.open()
                        new_doc.insert_pdf(self.doc, from_page=start_page, to_page=total_pages - 1)
                        save_path = f"{save_dir}/{file_prefix}_{original_file_name}_pages_{start_page + 1}-{total_pages}.pdf"
                        new_doc.save(save_path)

                    self.show_success("PDF split successfully!")
                    
                elif split_option and split_option[0] == "every_n_pages":
                    n = int(split_option[1])
                    for i in range(0, self.doc.page_count, n):
                        new_doc = fitz.open()
                        new_doc.insert_pdf(self.doc, from_page=i, to_page=min(i + n - 1, self.doc.page_count - 1))
                        save_path = f"{save_dir}/{file_prefix}_{original_file_name}_pages_{i + 1}-{min(i + n, self.doc.page_count)}.pdf"
                        new_doc.save(save_path)
                    self.show_success("PDF split successfully!")
                    
            except Exception as e:
                self.show_error(f"Error splitting PDF: {e}")
        else:
            self.show_warning("Please open a PDF file first.")

    def show_pdf_info(self):
        if self.doc:
            dialog = PDFInfoDialog(self.doc, self)
            dialog.exec_()
        else:
            QMessageBox.warning(self, "警告", "請先打開一個PDF文件。")

    def show_readme(self):
        dialog = READMEDialog(self)
        dialog.exec_()


    def encrypt_pdf(self, pdf_path, password):
        try:
            doc = fitz.open(pdf_path)
            new_path = pdf_path.replace('.pdf', '_encrypted.pdf')  # 定义加密文件的新路径
            doc.save(new_path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password)
            doc.close()
            QMessageBox.information(self, "成功", "PDF文件已加密並保存於: " + new_path)
        except Exception as e:
            QMessageBox.critical(self, "錯誤", f"加密 PDF 時發生錯誤：{e}")

    def on_encrypt_pdf_triggered(self):
        filename, _ = QFileDialog.getOpenFileName(self, "選擇 PDF 文件", "", "PDF files (*.pdf)")
        if filename:
            dialog = EncryptDialog(self)
            if dialog.exec_() == QDialog.Accepted:
                password, confirm_password = dialog.get_password()
                if password and password == confirm_password:
                    self.encrypt_pdf(filename, password)
                else:
                    QMessageBox.warning(self, "錯誤", "密碼不匹配或未輸入密碼！")


    def decrypt_pdf(self, pdf_path, password):
        try:
            # 使用密码尝试打开PDF文件
            doc = fitz.open(pdf_path, filetype="pdf")
            # 检查文档是否加密并尝试使用提供的密码解密
            if doc.is_encrypted:
                if doc.authenticate(password):
                    # 如果密码正确，保存解密后的PDF到新的文件路径
                    new_path = pdf_path.replace('.pdf', '_decrypted.pdf')
                    doc.save(new_path)
                    doc.close()
                    QMessageBox.information(self, "成功", f"PDF 密碼已移除並保存為：{new_path}")
                else:
                    raise Exception("密碼不正確")
            else:
                QMessageBox.warning(self, "警告", "PDF 文件未加密。")
        except Exception as e:
            QMessageBox.critical(self, "錯誤", f"移除 PDF 密碼時發生錯誤：{e}")

    def open_batch_convert_dialog(self):
        dialog = BatchConvertDialog(self)
        dialog.exec_()  # 顯示對話框並等待用戶操作


    def on_decrypt_pdf_triggered(self):
        filename, _ = QFileDialog.getOpenFileName(self, "選擇加密的 PDF 文件", "", "PDF files (*.pdf)")
        if filename:
            dialog = DecryptDialog(self)
            if dialog.exec_() == QDialog.Accepted:
                password = dialog.get_password()
                if password:
                    self.decrypt_pdf(filename, password)
                else:
                    QMessageBox.warning(self, "警告", "未輸入密碼！")

    def add_toolbar_spacer(self, height=10):
        spacer = QWidget()
        spacer.setFixedHeight(height)
        self.toolbar.addWidget(spacer)


    def open_batch_TXT_convert_dialog(self):
        dialog = BatchConvertTXTtoPDFDialog(self)
        dialog.exec_()

    def open_batch_print_dialog(self):
        dialog = BatchPrintDialog(self)
        dialog.exec_()

    def open_overlay_dialog(self):
        dialog = PDFOverlayDialog(self)
        dialog.exec_()  # 显示对话框并允许重复使用


    def overlay_pdf(self, template_path, target_folder, output_folder):
        try:
            template_doc = fitz.open(template_path)
            template_page_count = template_doc.page_count
            
            files_to_process = [f for f in os.listdir(target_folder) if f.lower().endswith('.pdf')]
            total_files = len(files_to_process)
            
            progress_dialog = QProgressDialog("Processing PDFs...", "Cancel", 0, total_files, self)
            progress_dialog.setWindowTitle("PDF Overlay Progress")
            progress_dialog.setWindowModality(Qt.WindowModal)
            progress_dialog.setMinimumDuration(0)
            progress_dialog.show()
            
            for i, file_name in enumerate(files_to_process):
                if progress_dialog.wasCanceled():
                    break
                
                target_path = os.path.join(target_folder, file_name)
                target_doc = fitz.open(target_path)

                new_doc = fitz.open()  # 创建一个新的目标文档

                for page_num in range(target_doc.page_count):
                    target_page = target_doc.load_page(page_num)
                    
                    # 计算要使用的模板页的页码，模板页循环使用
                    template_page_num = page_num % template_page_count
                    template_page = template_doc.load_page(template_page_num)
                    
                    # 创建一个新的空白页面，大小与目标页面相同
                    new_page = new_doc.new_page(width=target_page.rect.width, height=target_page.rect.height)
                    
                    # 先绘制模板页的内容
                    new_page.show_pdf_page(new_page.rect, template_doc, pno=template_page_num)
                    
                    # 再绘制目标页的内容
                    new_page.show_pdf_page(new_page.rect, target_doc, pno=page_num)

                output_path = os.path.join(output_folder, file_name)
                new_doc.save(output_path)
                new_doc.close()

                progress_dialog.setValue(i + 1)
                progress_dialog.setLabelText(f"Processing {i + 1} of {total_files} PDFs...")

            progress_dialog.close()
            QMessageBox.information(self, "成功", f"所有PDF文件叠加完成並保存於：{output_folder}")
        
        except Exception as e:
            QMessageBox.critical(self, "錯誤", f"叠加PDF時發生錯誤：{e}")
        finally:
            template_doc.close()

    def open_merge_csv_excel_dialog(self):
        dialog = MergeCSVExcelDialog(self)
        dialog.exec_()

    def open_deep_search_dialog(self):
        # 創建非模態對話框
        dialog = DeepSearchDialog(self)
        
        # 顯示對話框但不阻塞主窗口
        dialog.show()
        
        # 確保對話框在前台
        dialog.raise_()
        dialog.activateWindow()
        
        # 將對話框保存在類變量中，避免被垃圾收集
        self.deep_search_dialog = dialog


    def jump_to_page_with_highlight(self, page_num, highlight=False):
        """跳轉到指定頁面，並可選擇是否顯示高亮"""
        try:
            # 保存當前高亮設置
            original_highlight = False
            if hasattr(self, 'search_widget') and self.search_widget:
                original_highlight = self.search_widget.highlight_checkbox.isChecked()
                
            # 如果不需要高亮，先清除所有高亮
            if not highlight:
                self.clear_highlights()
                # 臨時設置高亮選項為 False
                if hasattr(self, 'search_widget') and self.search_widget:
                    self.search_widget.highlight_checkbox.setChecked(False)
            elif hasattr(self, 'search_widget') and self.search_widget:
                # 如果需要高亮，確保高亮選項被勾選
                self.search_widget.highlight_checkbox.setChecked(True)
            
            # 設置當前頁面並顯示
            self.current_page = page_num
            self.show_page(self.current_page)
            
            # 恢復原始高亮設置
            if hasattr(self, 'search_widget') and self.search_widget:
                self.search_widget.highlight_checkbox.setChecked(original_highlight)
        except Exception as e:
            print(f"跳轉頁面時出錯: {e}")

                    
    def wheel_event(self, event):
        if event.modifiers() & Qt.ShiftModifier:  # Check if the Shift key is pressed
            if event.angleDelta().y() > 0:
                self.show_prev_page()  # Show the previous page
            else:
                self.show_next_page()  # Show the next page
        elif event.modifiers() & Qt.ControlModifier:
            if event.angleDelta().y() > 0:
                self.zoom_in()  # Zoom in
            else:
                self.zoom_out()  # Zoom out

    def rotate_pdf_left(self):
        if self.doc:
            for page_num in range(self.doc.page_count):
                page = self.doc.load_page(page_num)
                page.set_rotation((page.rotation - 90) % 360)  # 逆時針旋轉
            self.show_page(self.current_page)
            self.is_modified = True

    def rotate_pdf_right(self):
        if self.doc:
            for page_num in range(self.doc.page_count):
                page = self.doc.load_page(page_num)
                page.set_rotation((page.rotation + 90) % 360)  # 順時針旋轉
            self.show_page(self.current_page)
            self.is_modified = True
            

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        if event.mimeData().hasUrls():
            url = event.mimeData().urls()[0]
            file_path = url.toLocalFile()
            if file_path.lower().endswith('.pdf'):
                self.load_pdf(file_path)

    # 新增跳轉到第一頁的方法
    def goto_first_page(self):
        if self.doc:
            self.current_page = 0
            self.show_page(self.current_page)

    # 新增跳轉到最後一頁的方法
    def goto_last_page(self):
        if self.doc:
            self.current_page = self.doc.page_count - 1
            self.show_page(self.current_page)

    def show_error(self, message):
        msg_box = QMessageBox(QMessageBox.Critical, "錯誤", message, parent=self)
        msg_box.setStyleSheet("QLabel{color: white;}")
        msg_box.exec_()

    def show_success(self, message):
        msg_box = QMessageBox(QMessageBox.Information, "成功", message, parent=self)
        msg_box.setStyleSheet("QLabel{color: white;}")
        msg_box.exec_()

    def show_warning(self, message):
        msg_box = QMessageBox(QMessageBox.Warning, "警告", message, parent=self)
        msg_box.setStyleSheet("QLabel{color: white;}")
        msg_box.exec_()

    def show_info(self, message):
        msg_box = QMessageBox(QMessageBox.Information, "信息", message, parent=self)
        msg_box.setStyleSheet("QLabel{color: white;}")
        msg_box.exec_()


def create_scaled_splash(app, splash_img_path):
    # 獲取主屏幕大小
    screen = QDesktopWidget().screenNumber(QDesktopWidget().cursor().pos())
    screen_size = QDesktopWidget().screenGeometry(screen)
    screen_width = screen_size.width()
    screen_height = screen_size.height()

    # 載入原始圖片
    original_pixmap = QPixmap(splash_img_path)
    original_width = original_pixmap.width()
    original_height = original_pixmap.height()

    # 計算縮放比例
    if screen_width <= 1920 and screen_height <= 1080:
        # 屏幕分辨率小於等於1920x1080，使用原始大小
        scaled_pixmap = original_pixmap
    else:
        # 屏幕分辨率大於1920x1080，計算縮放比例
        scale_x = min(screen_width / 1920, 2)  # 最大放大到4K的寬度 (3840 / 1920 = 2)
        scale_y = min(screen_height / 1080, 2)  # 最大放大到4K的高度 (2160 / 1080 = 2)
        scale_factor = min(scale_x, scale_y)  # 使用較小的縮放比例以確保完全顯示

        # 縮放圖片
        scaled_pixmap = original_pixmap.scaled(
            int(original_width * scale_factor),
            int(original_height * scale_factor),
            Qt.KeepAspectRatio,
            Qt.SmoothTransformation
        )

    # 創建並顯示splash screen
    splash = QSplashScreen(scaled_pixmap)

    # 計算居中位置
    x = (screen_width - scaled_pixmap.width()) // 2
    y = (screen_height - scaled_pixmap.height()) // 2
    splash.move(x, y)

    splash.show()
    app.processEvents()

    return splash

if __name__ == '__main__':
    app = QApplication(sys.argv)
    
    # Set the application style to 'Fusion'
    QApplication.setStyle("Fusion")
    
    # Create and set the application palette
    palette = QPalette()
    palette.setColor(QPalette.Window, QColor(53, 53, 53))
    palette.setColor(QPalette.WindowText, Qt.white)
    palette.setColor(QPalette.Base, QColor(25, 25, 25))
    palette.setColor(QPalette.AlternateBase, QColor(53, 53, 53))
    palette.setColor(QPalette.ToolTipBase, Qt.white)
    palette.setColor(QPalette.ToolTipText, Qt.white)
    palette.setColor(QPalette.Text, Qt.white)
    palette.setColor(QPalette.Button, QColor(53, 53, 53))
    palette.setColor(QPalette.ButtonText, Qt.white)
    palette.setColor(QPalette.BrightText, Qt.red)
    palette.setColor(QPalette.Link, QColor(42, 130, 218))
    palette.setColor(QPalette.Highlight, QColor(42, 130, 218))
    palette.setColor(QPalette.HighlightedText, Qt.black)
    QApplication.setPalette(palette)
    
    # Set the application icon
    app_icon = QIcon('icon.ico')
    app.setWindowIcon(app_icon)
    
    # Create and display the splash screen
    splash_img_path = resource_path('splash.png')
    splash = create_scaled_splash(app, splash_img_path)

    # Create the main window
    viewer = PDFViewer()
    viewer.setStyleSheet("QLabel { color: white; }")  # Set QLabel text color globally for the viewer
    splash.finish(viewer)
    viewer.show()

    # Load a PDF file if provided as a command line argument
    if len(sys.argv) > 1:
        pdf_path = sys.argv[1]
        viewer.load_pdf(pdf_path)

    sys.exit(app.exec_())

