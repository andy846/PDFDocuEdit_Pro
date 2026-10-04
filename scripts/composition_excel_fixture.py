"""Small synthetic Excel fixtures for native/frozen acceptance; no customer data."""
from __future__ import annotations

import struct
from datetime import date
from pathlib import Path


def make_xlsx(path, count=100):
    from openpyxl import Workbook
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["Record Count"])
    summary.append([count])
    sheet = book.create_sheet("Statements")
    sheet.append(["Synthetic production data"])
    sheet.append(["Customer Name", "Account", "Balance", "Date", "Flag", "Note"])
    for i in range(count):
        sheet.append(["陳小明", i+1, 10000.25+i, date(2026, 10, 2), i % 2 == 0, None])
        sheet.cell(i+3, 2).number_format = "000000"
    book.save(path)
    book.close()
    return Path(path)


def make_xls(path, count=3):
    """BIFF8 Workbook stream in a normal OLE compound file, with exact cell formats."""
    def rec(code, payload=b""):
        return struct.pack("<HH", code, len(payload)) + payload
    def bof(kind):
        return rec(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 1997, 0, 6))
    def label(row, col, text):
        return rec(0x0204, struct.pack("<HHHHB", row, col, 0, len(text), 1) + text.encode("utf-16le"))
    def number(row, col, value, xf=0):
        return rec(0x0203, struct.pack("<HHHd", row, col, xf, value))
    def xf(fmt):
        return rec(0x00e0, struct.pack("<HH", 0, fmt) + b"\0"*16)
    end = rec(0x000a)
    name = b"Statements"
    # Custom FORMAT must precede XF references in an actual BIFF workbook.
    header = (bof(5) + rec(0x0042, struct.pack("<H", 1200)) +
              rec(0x041e, struct.pack("<HHB", 164, 6, 0)+b"000000") + xf(0)+xf(164)+xf(14))
    bound = struct.pack("<IBBB", 0, 0, 0, len(name)) + b"\0" + name
    offset = len(header)+len(rec(0x0085, bound))+len(end)
    bound = struct.pack("<IBBB", offset, 0, 0, len(name)) + b"\0" + name
    rows = bof(0x10)+rec(0x0200, struct.pack("<IIHHH", 0, count+1, 0, 3, 0))
    rows += label(0, 0, "Name")+label(0, 1, "Account")+label(0, 2, "Date")
    from openpyxl.utils.datetime import to_excel
    for i in range(count):
        rows += label(i+1, 0, "陳小明")+number(i+1, 1, i+1, 1)+number(i+1, 2, to_excel(date(2026, 10, 2)), 2)
    raw = header+rec(0x0085, bound)+end+rows+end
    # Normal FAT streams need >=4096 bytes; trailing zero padding follows BIFF EOF.
    raw = raw.ljust(max(4096, (len(raw)+511)//512*512), b"\0")
    sectors = len(raw)//512
    if sectors+2 > 128:
        raise ValueError("QA fixture exceeds the single-FAT size")
    free, eoc, fatsect = 0xffffffff, 0xfffffffe, 0xfffffffd
    fat_index = sectors+1
    header = (bytes.fromhex("d0cf11e0a1b11ae1")+b"\0"*16+
              struct.pack("<HHHHH", 0x3e, 3, 0xfffe, 9, 6)+b"\0"*6+
              struct.pack("<IIIIIIIII", 0, 1, 0, 0, 4096, eoc, 0, eoc, 0)+
              struct.pack("<109I", fat_index, *([free]*108)))
    def entry(name, kind, child, start, size):
        raw_name = (name+"\0").encode("utf-16le")
        return (raw_name.ljust(64, b"\0")+struct.pack("<HBBIII", len(raw_name), kind, 1, free, free, child)+
                b"\0"*16+struct.pack("<IQQIQ", 0, 0, 0, start, size))
    directory = (entry("Root Entry", 5, 1, eoc, 0)+entry("Workbook", 2, free, 1, len(raw))).ljust(512, b"\0")
    table = [free]*128
    table[0] = eoc
    for index in range(1, sectors+1):
        table[index] = index+1 if index < sectors else eoc
    table[fat_index] = fatsect
    Path(path).write_bytes(header+directory+raw+struct.pack("<128I", *table))
    return Path(path)
